"""FastAPI сервер для API Tracker.

Предоставляет:
- REST API для приёма логов вызовов (/api/log)
- Выдачу статистики и графиков (/api/stats)
- Поиск и фильтрацию логов (/api/logs)
- Экспорт в CSV и Excel (/api/export/excel, /api/export/csv)
- Раздачу статики дашборда (/)
"""

from __future__ import annotations

import csv
import io
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from fastapi import FastAPI, Query, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from app.db import get_logs, get_stats, init_db, insert_call
from app.identity import detect_device_name, load_user_profile
from app.pricing import DEFAULT_PRICING, calculate_cost

STATIC_DIR = Path(__file__).resolve().parent / "static"
MSK_TZ = timezone(timedelta(hours=3))


def to_msk_str(ts_str: str) -> str:
    """Конвертировать ISO-таймстемп в московское время (MSK UTC+3)."""
    if not ts_str:
        return ""
    try:
        clean = ts_str.replace("Z", "+00:00")
        dt = datetime.fromisoformat(clean)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(MSK_TZ).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ts_str.replace("T", " ")[:19]


@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    yield


app = FastAPI(title="API Tracker & Cost Monitor", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class LogEventRequest(BaseModel):
    provider: str = Field(..., description="vibecode | kie | google | openai | outsee | elevenlabs")
    model: str = Field(..., description="ID модели, например gemini-3.7-flash")
    call_type: str = Field("text", description="text | video | image | audio")
    user_name: str = Field("", description="Имя пользователя / тестировщика")
    device_name: str = Field("", description="Имя ПК")
    key_alias: str = Field("", description="Маска или имя ключа")
    prompt_tokens: int = Field(0)
    completion_tokens: int = Field(0)
    cached_tokens: int = Field(0)
    media_count: int = Field(0)
    duration_sec: float = Field(0.0)
    cost_usd: float | None = Field(None, description="Если не передано, рассчитывается автоматически")
    status_code: int = Field(200)
    error_message: str = Field("")
    project_source: str = Field("", description="chat | pipeline | script")
    metadata: dict[str, Any] = Field(default_factory=dict)
    timestamp: str | None = Field(None)


class UserProfileRequest(BaseModel):
    user_name: str = Field(..., description="Отображаемое имя пользователя")


@app.get("/api/user-profile")
async def get_profile():
    """Получить текущий профиль пользователя."""
    from app.identity import load_user_profile
    return load_user_profile()


@app.post("/api/user-profile")
async def update_profile(req: UserProfileRequest):
    """Сохранить подтвержденное имя пользователя."""
    from app.identity import save_user_profile
    return save_user_profile(req.user_name)


@app.get("/api/users")
async def list_users():
    """Получить список всех уникальных пользователей (облако + локально)."""
    from app.cloud_sync import get_cloud_distinct_users
    from app.db import get_local_distinct_users
    cloud_users = get_cloud_distinct_users()
    local_users = get_local_distinct_users()
    combined = sorted(set(cloud_users + local_users))
    return {"users": combined}


@app.post("/api/sync")
async def sync_now():
    """Принудительная отправка накопившихся локальных логов в Supabase."""
    from app.cloud_sync import push_call_to_supabase
    from app.db import get_unsynced_calls, mark_calls_synced
    pending = get_unsynced_calls(limit=100)
    synced_ids = []
    for call in pending:
        if push_call_to_supabase(call):
            synced_ids.append(call["id"])
    if synced_ids:
        mark_calls_synced(synced_ids)
    return {"status": "ok", "synced_count": len(synced_ids), "remaining": len(pending) - len(synced_ids)}


@app.post("/api/log")
async def log_call(event: LogEventRequest):
    """Записать вызов API и рассчитать стоимость."""
    from app.cloud_sync import push_call_to_supabase
    from app.identity import get_user_identity

    u_name, d_name, _ = get_user_identity()
    active_user = event.user_name.strip() or u_name
    active_device = event.device_name.strip() or d_name

    cost = event.cost_usd
    if cost is None:
        cost = calculate_cost(
            event.model,
            call_type=event.call_type,
            prompt_tokens=event.prompt_tokens,
            completion_tokens=event.completion_tokens,
            cached_tokens=event.cached_tokens,
            media_count=event.media_count,
            duration_sec=event.duration_sec,
            chars=int(event.metadata.get("chars", 0)),
        )

    # 1. Отправка в Supabase
    is_synced = False
    try:
        is_synced = push_call_to_supabase({
            "user_name": active_user,
            "device_name": active_device,
            "provider": event.provider,
            "model": event.model,
            "key_alias": event.key_alias,
            "call_type": event.call_type,
            "prompt_tokens": event.prompt_tokens,
            "completion_tokens": event.completion_tokens,
            "cached_tokens": event.cached_tokens,
            "total_tokens": event.prompt_tokens + event.completion_tokens,
            "media_count": event.media_count,
            "duration_sec": event.duration_sec,
            "cost_usd": cost,
            "status_code": event.status_code,
            "error_message": event.error_message,
            "project_source": event.project_source,
            "metadata": event.metadata,
            "timestamp": event.timestamp,
        })
    except Exception:
        pass

    # 2. Локальная запись в SQLite
    call_id = insert_call(
        provider=event.provider,
        model=event.model,
        user_name=active_user,
        device_name=active_device,
        synced=1 if is_synced else 0,
        key_alias=event.key_alias,
        call_type=event.call_type,
        prompt_tokens=event.prompt_tokens,
        completion_tokens=event.completion_tokens,
        cached_tokens=event.cached_tokens,
        media_count=event.media_count,
        duration_sec=event.duration_sec,
        cost_usd=cost,
        status_code=event.status_code,
        error_message=event.error_message,
        project_source=event.project_source,
        metadata=event.metadata,
        timestamp=event.timestamp,
    )
    return {"status": "ok", "id": call_id, "cost_usd": cost, "synced": is_synced}


@app.get("/api/stats")
async def stats(
    scope: str = Query("local", description="cloud | local"),
    user_name: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
):
    """Сводные метрики и данные для графиков (из облака Supabase или локально)."""
    if scope == "cloud":
        from app.cloud_sync import get_cloud_stats
        cloud_data = get_cloud_stats(user_name=user_name, date_from=date_from, date_to=date_to)
        if cloud_data and cloud_data.get("summary"):
            cloud_data["scope"] = "cloud"
            return cloud_data

    local_data = get_stats(user_name=user_name, date_from=date_from, date_to=date_to)
    local_data["scope"] = "local"
    return local_data


@app.get("/api/logs")
async def logs_list(
    scope: str = Query("local", description="cloud | local"),
    user_name: str | None = Query(None),
    limit: int = Query(50, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    provider: str | None = Query(None),
    model: str | None = Query(None),
    call_type: str | None = Query(None),
    status: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
    search: str | None = Query(None),
):
    """Список логов активности с фильтрацией (из облака Supabase или локально)."""
    if scope == "cloud":
        from app.cloud_sync import get_cloud_logs
        items, total = get_cloud_logs(
            limit=limit,
            offset=offset,
            user_name=user_name,
            provider=provider,
            model=model,
            call_type=call_type,
            status=status,
            date_from=date_from,
            date_to=date_to,
            search=search,
        )
        if total > 0 or not search:
            return {"items": items, "total": total, "limit": limit, "offset": offset, "scope": "cloud"}

    items, total = get_logs(
        limit=limit,
        offset=offset,
        user_name=user_name,
        provider=provider,
        model=model,
        call_type=call_type,
        status=status,
        date_from=date_from,
        date_to=date_to,
        search=search,
    )
    return {"items": items, "total": total, "limit": limit, "offset": offset, "scope": "local"}


@app.get("/api/pricing")
async def pricing_catalog():
    """Справочник цен по моделям."""
    return DEFAULT_PRICING


class KillswitchToggleRequest(BaseModel):
    action: str  # "block" | "unblock"
    user_name: str | None = None
    reason: str | None = None


@app.get("/api/system/killswitch")
async def get_killswitch_status():
    """Получить статус аварийного рубильника."""
    # 1. Проверяем облако
    try:
        from app.cloud_sync import get_cloud_killswitch_state
        st = get_cloud_killswitch_state()
        if st.get("updated_at"):
            return st
    except Exception:
        pass
    # 2. Локально
    try:
        from app.db import get_local_killswitch_state
        return get_local_killswitch_state()
    except Exception:
        pass
    return {"is_blocked": False, "updated_by": "", "updated_at": "", "reason": ""}


@app.post("/api/system/killswitch")
async def post_killswitch_toggle(req: KillswitchToggleRequest):
    """Переключить аварийный рубильник (заблокировать или возобновить)."""
    user_name = (req.user_name or "").strip()
    if not user_name:
        prof = load_user_profile()
        user_name = prof.get("user_name") or "Пользователь"

    device_name = detect_device_name()
    action_name = "KILLSWITCH_ACTIVATED" if req.action.lower() == "block" else "KILLSWITCH_DEACTIVATED"

    # 1. Запись в облако
    cloud_ok = False
    try:
        from app.cloud_sync import push_audit_log
        cloud_ok = push_audit_log(
            action=action_name,
            user_name=user_name,
            device_name=device_name,
            reason=req.reason or "",
        )
    except Exception:
        pass

    # 2. Запись в локальную базу
    desc = "Аварийный рубильник: API заморожены" if req.action.lower() == "block" else "Аварийный рубильник: работа API возобновлена"
    if req.reason:
        desc += f" ({req.reason})"
    try:
        insert_call(
            provider="SYSTEM",
            model=action_name,
            cost_usd=0.0,
            user_name=user_name,
            device_name=device_name,
            synced=1 if cloud_ok else 0,
            call_type="system",
            prompt_tokens=0,
            completion_tokens=0,
            cached_tokens=0,
            media_count=0,
            duration_sec=0.0,
            status_code=200,
            error_message=desc,
            project_source="killswitch",
            metadata={"action": action_name, "reason": req.reason or ""},
            timestamp=datetime.now(timezone.utc).isoformat(),
        )
    except Exception:
        pass

    # 3. Синхронизация с хуком в app/services
    try:
        from app.services.api_tracker_hook import toggle_killswitch
        toggle_killswitch(req.action, user_name=user_name, reason=req.reason or "")
    except Exception:
        pass

    is_blocked = (req.action.lower() == "block")
    return {
        "is_blocked": is_blocked,
        "updated_by": user_name,
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "reason": desc,
        "synced": cloud_ok,
    }



def _fetch_export_items(
    *,
    scope: str = "cloud",
    user_name: str | None = None,
    provider: str | None = None,
    model: str | None = None,
    status: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
) -> list[dict[str, Any]]:
    if scope == "cloud":
        try:
            from app.cloud_sync import get_cloud_logs
            items, _ = get_cloud_logs(
                limit=10000,
                offset=0,
                user_name=user_name,
                provider=provider,
                model=model,
                status=status,
                date_from=date_from,
                date_to=date_to,
            )
            if items:
                return items
        except Exception:
            pass
    items, _ = get_logs(
        limit=10000,
        offset=0,
        user_name=user_name,
        provider=provider,
        model=model,
        status=status,
        date_from=date_from,
        date_to=date_to,
    )
    return items


def build_excel_workbook(items: list[dict[str, Any]]) -> bytes:
    """Генерация Excel (.xlsx) с форматированием, границами и итоговой строкой."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Активность API"

    headers = [
        "№", "Дата и время (МСК)", "Пользователь", "Устройство", "Провайдер", "Модель", "Тип",
        "Входные токены", "Выходные токены", "Всего токенов",
        "Длительность (сек)", "Стоимость ($)", "Статус", "Источник", "Ошибка"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="164E63", end_color="164E63", fill_type="solid")
    header_font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    total_cost = 0.0
    total_tokens = 0

    for idx, row in enumerate(items, start=1):
        c_val = float(row.get("cost_usd") or 0.0)
        tok_val = int(row.get("total_tokens") or 0)
        total_cost += c_val
        total_tokens += tok_val

        ws.append([
            idx,
            to_msk_str(row.get("timestamp")),
            row.get("user_name") or "—",
            row.get("device_name") or "—",
            row.get("provider") or "",
            row.get("model") or "",
            (row.get("call_type") or "text").upper(),
            row.get("prompt_tokens") or 0,
            row.get("completion_tokens") or 0,
            tok_val,
            round(float(row.get("duration_sec") or 0.0), 2),
            round(c_val, 6),
            row.get("status_code") or 200,
            row.get("project_source") or "",
            row.get("error_message") or "",
        ])

        curr_row = idx + 1
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=curr_row, column=c)
            cell.border = thin_border
            if c in (1, 7, 13):
                cell.alignment = center_align
            elif c in (8, 9, 10, 11, 12):
                cell.alignment = right_align
            else:
                cell.alignment = left_align

    if items:
        tot_row = len(items) + 2
        ws.cell(row=tot_row, column=2, value="ИТОГО:").font = Font(name="Calibri", bold=True, size=11)
        ws.cell(row=tot_row, column=10, value=total_tokens).font = Font(name="Calibri", bold=True, size=11)
        cost_cell = ws.cell(row=tot_row, column=12, value=round(total_cost, 4))
        cost_cell.font = Font(name="Calibri", bold=True, size=11, color="047857")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


@app.get("/api/export/csv")
async def export_csv(
    scope: str = Query("cloud"),
    user_name: str | None = Query(None),
    provider: str | None = Query(None),
    model: str | None = Query(None),
    status: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
):
    """Экспорт логов в CSV."""
    items = _fetch_export_items(
        scope=scope,
        user_name=user_name,
        provider=provider,
        model=model,
        status=status,
        date_from=date_from,
        date_to=date_to,
    )

    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow([
        "№", "Время (МСК)", "Пользователь", "Устройство", "Провайдер", "Модель", "Тип",
        "Входные токены", "Выходные токены", "Всего токенов",
        "Длительность (сек)", "Стоимость ($)", "Статус", "Источник", "Ошибка"
    ])

    for idx, row in enumerate(items, start=1):
        writer.writerow([
            idx,
            to_msk_str(row.get("timestamp")),
            row.get("user_name") or "—",
            row.get("device_name") or "—",
            row.get("provider") or "",
            row.get("model") or "",
            (row.get("call_type") or "text").upper(),
            row.get("prompt_tokens") or 0,
            row.get("completion_tokens") or 0,
            row.get("total_tokens") or 0,
            row.get("duration_sec") or 0.0,
            f"{float(row.get('cost_usd') or 0.0):.6f}",
            row.get("status_code") or 200,
            row.get("project_source") or "",
            row.get("error_message") or "",
        ])

    data = buf.getvalue().encode("utf-8-sig")
    return Response(
        content=data,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=api_activity.csv"},
    )


@app.get("/api/export/excel")
async def export_excel(
    scope: str = Query("cloud"),
    user_name: str | None = Query(None),
    provider: str | None = Query(None),
    model: str | None = Query(None),
    status: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
):
    """Экспорт логов в Excel (.xlsx) с форматированием."""
    items = _fetch_export_items(
        scope=scope,
        user_name=user_name,
        provider=provider,
        model=model,
        status=status,
        date_from=date_from,
        date_to=date_to,
    )
    content = build_excel_workbook(items)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=api_activity.xlsx"},
    )


@app.post("/api/export/save-to-downloads")
async def save_excel_to_downloads(
    scope: str = Query("cloud"),
    user_name: str | None = Query(None),
    provider: str | None = Query(None),
    model: str | None = Query(None),
    status: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
):
    """Сформировать и сохранить отчёт Excel напрямую в папку Загрузки на ПК."""
    items = _fetch_export_items(
        scope=scope,
        user_name=user_name,
        provider=provider,
        model=model,
        status=status,
        date_from=date_from,
        date_to=date_to,
    )
    content = build_excel_workbook(items)

    downloads_dir = Path.home() / "Downloads"
    if not downloads_dir.exists():
        downloads_dir = Path.home()

    date_str = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filename = f"API_Tracker_{date_str}.xlsx"
    target_path = downloads_dir / filename
    target_path.write_bytes(content)

    return {
        "status": "ok",
        "filename": filename,
        "path": str(target_path),
        "total": len(items),
    }


class OpenFileRequest(BaseModel):
    path: str


@app.post("/api/export/open-file")
async def open_exported_file(req: OpenFileRequest):
    """Открыть сохранённый файл через системную ассоциацию Windows."""
    import os
    target = Path(req.path)
    if target.is_file():
        try:
            os.startfile(str(target))
            return {"status": "ok"}
        except Exception as e:
            return {"status": "error", "message": str(e)}
    return {"status": "error", "message": "Файл не найден"}


# Главная страница дашборда
@app.get("/")
async def index():
    html_path = STATIC_DIR / "index.html"
    if html_path.is_file():
        return FileResponse(html_path)
    return {"message": "API Tracker Backend Running"}


if STATIC_DIR.is_dir():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
