"""Шаг 4b «Предметы» — генерация реф-картинок предметов.

Параллельная ветвь шага 4 «Объекты»: если шаг 4a «Персонажи» делает
hero_reference, то 4b делает item_reference. Логика проще, чем у Hero:
без HITL, без вариаций (1 картинка на предмет).

Источник списка предметов:
1) `project.item_descriptions: list[str]`.
2) Если пуст — лист «Предметы» в `project.xlsx` или сущности `Entity(type="item")`.
По одному непустому описанию = один сгенерированный предмет.
Файлы кладутся в `data/videos/<slug>/items/predmet<N>_<uuid>.png`,
где N — 1-based индекс предмета.

Если шаг падает на каком-то предмете — статус откатывается на
hero_ready (предметы опциональны), юзер правит описание и жмёт
«Предметы» снова.
"""

from __future__ import annotations

import uuid
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, AsyncIterator

import openpyxl
from aiogram import Bot
from loguru import logger
from sqlalchemy import desc, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm.attributes import flag_modified

from app.bots.browser import browser_session
from app.bots.outsee import (
    OutseeBot,
    OutseeContentRejectedError,
    OutseeImageError,
)
from app.bots.outsee_http import outsee_api_configured, outsee_api_enabled_for_image
from app.db import SessionLocal
from app.generation_options import (
    IMAGE_GENERATORS_BY_ID,
    IMAGE_RESOLUTIONS_BY_ID,
)
from app.models import Artifact, ArtifactKind, Entity, Project, ProjectStatus
from app.services.gpt_client import get_gpt_client
from app.services.img_streams import acquire_image_slot
from app.services.media_route import image_provider_for
from app.services.outsee_retry import generate_image_with_retries
from app.services.prompt_library import get_project_prompt
from app.services.step_cancel import raise_if_cancelled
from app.settings import settings

# Aspect ratio и Relax для предметов — как у hero (16:9 + Relax), потому
# что предметы тоже идут как реф-листы.
ITEM_ASPECT_RATIO = "16:9"
ITEM_RELAX = True


def _need_browser_session(model_slug: str | None) -> bool:
    """True, если генератору картинок нужен Chrome CDP (Outsee browser).
    Для KIE HTTP (Flux 2 Pro, Kling, Qwen) и Outsee HTTP API браузер не нужен."""
    backend = image_provider_for(model_slug)
    if backend == "kie":
        return False
    if backend == "outsee" and (outsee_api_configured() or outsee_api_enabled_for_image()):
        return False
    return True


@asynccontextmanager
async def _optional_browser_session(
    need_cdp: bool,
) -> AsyncIterator[Any]:
    if not need_cdp:
        yield None
        return
    async with browser_session() as bs:
        yield bs


def _project_data_dir(project: Project) -> Path | None:
    """Безопасное получение data_dir (возвращает None если slug пуст или не инициализирован)."""
    try:
        if not getattr(project, "slug", None):
            return None
        return project.data_dir
    except Exception:
        return None


async def _existing_item_indices(
    session: AsyncSession, project: Project
) -> set[int]:
    """Какие индексы предметов уже имеют артефакт kind=item_reference."""
    rows = (
        await session.execute(
            select(Artifact)
            .where(
                Artifact.project_id == project.id,
                Artifact.kind == ArtifactKind.item_reference,
            )
            .order_by(desc(Artifact.id))
        )
    ).scalars().all()
    out: set[int] = set()
    for a in rows:
        m = a.meta or {}
        idx = m.get("item_index")
        if isinstance(idx, int):
            out.add(idx)

    # Disk fallback: проверяем файлы в data_dir/items/predmet<N>_*.png
    pdir = _project_data_dir(project)
    if pdir:
        items_dir = Path(pdir) / "items"
        if items_dir.is_dir():
            for p in items_dir.glob("predmet*.png"):
                if p.is_file() and p.stat().st_size > 1000:
                    stem = p.stem.lower()
                    prefix = stem.split("_")[0]
                    if prefix.startswith("predmet"):
                        num_part = prefix[len("predmet"):]
                        if num_part.isdigit():
                            out.add(int(num_part))
    return out


async def _resolve_item_descriptions(
    session: AsyncSession, project: Project
) -> list[str]:
    """Возвращает список непустых описаний предметов.
    Приоритет:
    1) project.item_descriptions
    2) Entity(type='item' | 'prop') в БД
    3) Лист «Предметы» в project.xlsx
    """
    raw = list(project.item_descriptions or [])
    descriptions = [d.strip() for d in raw if isinstance(d, str) and d.strip()]
    if descriptions:
        return descriptions

    # 2) Fallback: Entity в БД
    try:
        ents = (
            await session.execute(
                select(Entity)
                .where(
                    Entity.project_id == project.id,
                    Entity.type.in_(["prop", "item"]),
                )
                .order_by(Entity.id)
            )
        ).scalars().all()
        for e in ents:
            attrs = e.attrs or {}
            desc_val = (attrs.get("description") or attrs.get("описание") or "").strip()
            name_val = (e.name or "").strip()
            text = f"{name_val}: {desc_val}" if name_val and desc_val else (desc_val or name_val)
            if text and text not in descriptions:
                descriptions.append(text)
    except Exception as exc:  # noqa: BLE001
        logger.debug("[#{}] items: entity fallback check failed: {}", project.id, exc)

    if descriptions:
        project.item_descriptions = descriptions
        flag_modified(project, "item_descriptions")
        await session.flush()
        return descriptions

    # 3) Fallback: лист «Предметы» в project.xlsx
    pdir = _project_data_dir(project)
    if pdir:
        xlsx_path = Path(pdir) / "project.xlsx"
        if xlsx_path.is_file():
            try:
                wb = openpyxl.load_workbook(xlsx_path, data_only=True)
                target_ws = None
                for sname in wb.sheetnames:
                    lower_name = sname.lower()
                    if "предмет" in lower_name or "item" in lower_name or "prop" in lower_name:
                        target_ws = wb[sname]
                        break
                if target_ws is not None and target_ws.max_row >= 2:
                    # Проверяем колоночный формат (как в Персонажах: заголовки в Col 1, данные в Col 2..N)
                    for col in range(2, target_ws.max_column + 1):
                        name = target_ws.cell(3, col).value
                        desc = target_ws.cell(4, col).value
                        rules = target_ws.cell(8, col).value
                        parts = [str(x).strip() for x in (name, desc, rules) if x and str(x).strip()]
                        if parts:
                            item_text = " — ".join(parts)
                            descriptions.append(item_text)

                    # Если по колонкам ничего не нашли, проверяем построчный формат (заголовки в Row 1, данные в Row 2..N)
                    if not descriptions and target_ws.max_row > 1:
                        for row in range(2, target_ws.max_row + 1):
                            vals = [
                                str(target_ws.cell(row, c).value).strip()
                                for c in range(1, target_ws.max_column + 1)
                                if target_ws.cell(row, c).value
                            ]
                            if vals:
                                descriptions.append(" — ".join(vals))
            except Exception as exc:  # noqa: BLE001
                logger.warning("[#{}] items: xlsx parse fallback failed: {}", project.id, exc)

    if descriptions:
        project.item_descriptions = descriptions
        flag_modified(project, "item_descriptions")
        await session.flush()

    return descriptions


def _items_style_prompt(project: Project) -> str:
    """Мастер-промт для генерации предметов (prompts/04b_items/<name>.md).
    Если ничего не выбрано/нет файла — пустая строка (юзер должен
    положить хотя бы default.md)."""
    try:
        return get_project_prompt(project, "items").strip()
    except FileNotFoundError:
        logger.warning(
            "items: prompts/04b_items/default.md не найден — генерирую "
            "только из описаний без стиля"
        )
        return ""


async def run(session: AsyncSession, project: Project, bot: Bot) -> None:
    if project.status is not ProjectStatus.generating_items:
        return

    descriptions = await _resolve_item_descriptions(session, project)
    if not descriptions:
        logger.info(
            "[#{}] items: item_descriptions пуст — items_ready без работы",
            project.id,
        )
        project.status = ProjectStatus.items_ready
        await session.flush()
        return

    style = _items_style_prompt(project)
    from app.services.vibecode_catalog import resolve_node_media_settings

    media = resolve_node_media_settings(project, node_type="items")
    img_gid = media["image_generator_id"]
    img_gen = IMAGE_GENERATORS_BY_ID.get(img_gid)
    ir = IMAGE_RESOLUTIONS_BY_ID.get(media["resolution_id"])
    quality_slug = media["quality_slug"]
    item_aspect = media["aspect_slug"] or ITEM_ASPECT_RATIO

    already_done = await _existing_item_indices(session, project)
    pdir = _project_data_dir(project)
    if not pdir:
        raise RuntimeError(f"Project #{project.id} data_dir is invalid (slug is empty)")
    out_dir = pdir / "items"
    out_dir.mkdir(parents=True, exist_ok=True)

    model_slug = img_gen.outsee_slug if img_gen else None
    need_cdp = _need_browser_session(model_slug)

    # Идём по предметам последовательно, пропускаем уже сгенерированные.
    for idx, desc_text in enumerate(descriptions, start=1):
        raise_if_cancelled(project.id)

        if idx in already_done:
            logger.info(
                "[#{}] items: predmet{} уже есть, пропускаю",
                project.id,
                idx,
            )
            continue
        logger.info(
            "[#{}] items: предмет {}/{} — '{}' (need_cdp={})",
            project.id,
            idx,
            len(descriptions),
            desc_text[:60],
            need_cdp,
        )

        full_prompt = (
            (style + "\n\n---\n\n" if style else "")
            + f"Описание предмета (predmet{idx}): {desc_text}"
        )

        short_uuid = uuid.uuid4().hex[:8]
        file_name = f"predmet{idx}_{short_uuid}.png"
        out_path = out_dir / file_name
        prompt_id_prefix = f"[ID: P{project.id}-ITEM{idx}-{short_uuid}]"

        try:
            async with _optional_browser_session(need_cdp=need_cdp) as bs:
                outsee = OutseeBot(bs) if bs is not None else None
                gpt = get_gpt_client()
                async with acquire_image_slot():
                    result = await generate_image_with_retries(
                        outsee, gpt,
                        prompt=full_prompt,
                        out_path=out_path,
                        max_attempts_per_prompt=3,
                        gpt_rewrite=True,
                        aspect_ratio=item_aspect,
                        model_slug=model_slug,
                        resolution=ir.outsee_slug if ir else None,
                        quality=quality_slug,
                        relax=ITEM_RELAX,
                        prompt_id_prefix=prompt_id_prefix,
                        reference_image=None,
                        timeout=600,
                        project_id=project.id,
                    )
        except OutseeImageError as e:
            is_moderation = isinstance(e, OutseeContentRejectedError)
            logger.error(
                "[#{}] items: predmet{} 6 попыток провалились "
                "(moderation={}): {}",
                project.id, idx, is_moderation,
                getattr(e, "reason", None) or str(e),
            )
            # Откат на hero_ready: предметы опциональны, юзер может
            # пропустить и идти дальше.
            project.status = ProjectStatus.hero_ready
            await session.flush()
            raise RuntimeError(
                f"items: predmet{idx} не удалось сгенерить "
                f"(см. логи). Статус откатил на hero_ready — поправь "
                f"описание предмета и жми «Предметы» снова."
            ) from e

        # Сохраняем артефакт в project.db
        art_uuid = uuid.uuid4().hex
        art_meta = {
            "item_index": idx,
            "item_id": f"predmet{idx}",
            "description": desc_text,
            "prompt": full_prompt,
        }
        a = Artifact(
            project_id=project.id,
            frame_id=None,
            kind=ArtifactKind.item_reference,
            uuid=art_uuid,
            path=str(result.file_path),
            meta=art_meta,
        )
        session.add(a)
        await session.flush()

        # Синхронизируем Artifact в master state.db для веб-UI студии
        try:
            async with SessionLocal() as master_sess:
                m_art = Artifact(
                    project_id=project.id,
                    kind=ArtifactKind.item_reference,
                    uuid=art_uuid,
                    path=str(result.file_path),
                    meta=art_meta,
                )
                master_sess.add(m_art)
                await master_sess.commit()
        except Exception as exc:  # noqa: BLE001
            logger.debug("[#{}] master artifact sync for item {}: {}", project.id, idx, exc)

        logger.info(
            "[#{}] items: predmet{} → {}",
            project.id,
            idx,
            result.file_path,
        )

    # Все предметы готовы.
    project.status = ProjectStatus.items_ready
    await session.flush()
    logger.info(
        "[#{}] items: все {} предметов готовы → items_ready",
        project.id,
        len(descriptions),
    )

