"""Общие хелперы xlsx-flow для шагов ChatGPT.

Правило отправки:
  - Мастер-промт (выбранный variant с диска) → всегда во временный файл.
  - Текст в чат → только `gpt_text_builder.get_effective_text()` (override
    пользователя или дефолтное сопр. сообщение без мастер-промта).
"""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from loguru import logger
from sqlalchemy.ext.asyncio import AsyncSession

from app.bots.chatgpt import ChatGPTBot
from app.models import Project
from app.services import gpt_text_builder as gtb
from app.services.prompt_library import get_project_prompt
from app.services.xlsx_sync import reload_from_xlsx
from app.services.xlsx_v8_import import has_v8_plan_sheet, import_v8_xlsx
from app.services.xlsx_versioning import backup_to_old, replace_with, validate_xlsx

from app.bots.browser import _looks_like_cdp_connect_failure


def _sync_had_changes(info: dict | None) -> bool:
    if not info:
        return False
    return bool(
        info.get("project_fields_changed")
        or info.get("frames_created")
        or info.get("frames_updated")
        or info.get("frames_changed")
    )


def _log_sync_result(project_id: int, label: str, info: dict | None) -> None:
    if _sync_had_changes(info):
        logger.info("[#{}] sync_project_xlsx {}: {}", project_id, label, info)
    else:
        logger.debug("[#{}] sync_project_xlsx {}: no changes", project_id, label)


def project_xlsx_stat(path: Path) -> tuple[float, int]:
    """(mtime, size) для проверки, обновился ли xlsx после GPT."""
    if path.exists():
        st = path.stat()
        return st.st_mtime, st.st_size
    return 0.0, 0


def should_accept_xlsx_after_gpt_error(
    path: Path,
    stat_before: tuple[float, int],
    exc: BaseException,
) -> bool:
    """True только если GPT реально записал новый xlsx (не старый файл на диске)."""
    if _looks_like_cdp_connect_failure(exc):
        return False
    if not path.exists() or path.stat().st_size < 1024:
        return False
    if validate_xlsx(path) is not None:
        return False
    st = path.stat()
    mtime_b, size_b = stat_before
    return st.st_mtime > mtime_b + 0.5 or st.st_size != size_b



def tmp_gpt_dir(project: Project) -> Path:
    d = project.data_dir / "tmp_gpt"
    d.mkdir(parents=True, exist_ok=True)
    return d


# Кэш GPT round-trip по шагам — чистим перед явным «Запустить шаг».
_STEP_TMP_GLOBS: dict[str, tuple[str, ...]] = {
    "plan": ("plan_*.xlsx", "prompt_plan_*"),
    "script": ("script_*.txt", "prompt_script_*"),
    "split": ("split_*.xlsx", "prompt_split_*"),
    "img_pr": ("prompt_img_pr_*"),
    "anim_pr": ("prompt_anim_pr_*"),
    "enrich_1": ("prompt_enrich_1_*", "enrich_1_*.xlsx"),
    "enrich_2": ("prompt_enrich_2_*", "enrich_2_*.xlsx"),
    "enrich_3": ("prompt_enrich_3_*", "enrich_3_*.xlsx"),
    "enrich_4": ("prompt_enrich_4_*", "enrich_4_*.xlsx"),
    "enrich_5": ("prompt_enrich_5_*", "enrich_5_*.xlsx"),
}


def purge_tmp_gpt_for_step(project: Project, step_code: str) -> int:
    """Удалить кэш GPT-файлов шага в tmp_gpt (перед повторным запуском)."""
    tmp_dir = project.data_dir / "tmp_gpt"
    if not tmp_dir.is_dir():
        return 0
    removed = 0
    protect_enrich = not step_code.startswith("enrich_")
    for pattern in _STEP_TMP_GLOBS.get(step_code, ()):
        for path in tmp_dir.glob(pattern):
            if protect_enrich and path.name.startswith("prompt_enrich_"):
                continue
            try:
                path.unlink(missing_ok=True)
                removed += 1
            except OSError as e:
                logger.warning(
                    "[#{}] purge_tmp_gpt {}: cannot delete {}: {}",
                    project.id,
                    step_code,
                    path.name,
                    e,
                )
    if removed:
        logger.info(
            "[#{}] purge_tmp_gpt {}: removed {} cached file(s)",
            project.id,
            step_code,
            removed,
        )
    return removed


def _timestamp() -> str:
    return datetime.utcnow().strftime("%Y%m%d_%H%M%S")


def _get_master_or_fallback(project: Project, step_code: str, fallback: str) -> str:
    try:
        return get_project_prompt(project, step_code).strip()
    except FileNotFoundError:
        return fallback


def write_plan_prompt_file(
    project: Project,
    tmp_dir: Path,
    *,
    topic: str | None = None,
    ts: str | None = None,
) -> Path:
    actual_topic = topic if topic is not None else (project.topic or "")
    master = _get_master_or_fallback(
        project,
        "plan",
        "# plan\n\nМастер-промт для шага «План» ещё не настроен.",
    )
    hero_hint = {
        "hero": (
            "Игнорируй автоматическое определение hero_needed, "
            "выставь hero_needed=true."
        ),
        "no_hero": (
            "Игнорируй автоматическое определение hero_needed, "
            "выставь hero_needed=false."
        ),
        "auto": "",
    }.get(project.hero_mode, "")
    extra = f"\n\nДополнительное указание: {hero_hint}" if hero_hint else ""
    prompt_file = tmp_dir / f"prompt_plan_{ts or _timestamp()}.txt"
    from app.services.gpt_text_builder import inject_topic_placeholders

    master = inject_topic_placeholders(master, actual_topic)
    prompt_file.write_text(
        f"Тема ролика: ({actual_topic})\n\n{master}{extra}",
        encoding="utf-8",
    )
    return prompt_file


def write_script_prompt_file(
    project: Project, tmp_dir: Path, *, ts: str | None = None
) -> Path:
    topic = (project.topic or "").strip()
    prompt_text = _get_master_or_fallback(
        project,
        "script",
        "Мастер-промт для шага «Закадровый текст» ещё не настроен.",
    )
    hero_hint = {
        "hero": (
            "РЕЖИМ: A (герой). В плане hero_needed=true. "
            "Найди главного персонажа по теме и плану, "
            "пиши персонажный сценарий — см. раздел «РЕЖИМ A» в инструкции."
        ),
        "no_hero": (
            "РЕЖИМ: B (тема). hero_needed=false. "
            "Без биографического героя — подробно раскрой тему, "
            "см. раздел «РЕЖИМ B» в инструкции."
        ),
        "auto": (
            "РЕЖИМ: определи сам по листу «Общий план» в xlsx "
            "(hero_needed и содержание плана) — A или B."
        ),
    }.get(project.hero_mode or "auto", "")
    from app.services.gpt_text_builder import inject_topic_placeholders

    prompt_text = inject_topic_placeholders(prompt_text, topic)
    extra = f"\n\n---\n\nУКАЗАНИЕ ПРОЕКТА:\n{hero_hint}\n" if hero_hint else ""
    prompt_file = tmp_dir / f"prompt_script_{ts or _timestamp()}.txt"
    prompt_file.write_text(
        f"# Инструкция для GPT (шаг 2 «Закадровый текст»)\n"
        f"# Тема ролика: «{topic}»\n\n"
        f"{prompt_text}{extra}\n",
        encoding="utf-8",
    )
    return prompt_file


def write_split_prompt_file(
    project: Project, tmp_dir: Path, *, ts: str | None = None
) -> Path:
    topic = (project.topic or "").strip()
    prompt_text = _get_master_or_fallback(
        project,
        "split",
        "Мастер-промт для шага «Разбивка на блоки» ещё не настроен.",
    )
    prompt_file = tmp_dir / f"prompt_split_{ts or _timestamp()}.txt"
    prompt_file.write_text(
        f"# Инструкция для GPT (шаг 3 «Разбивка на блоки»)\n"
        f"# Тема ролика: «{topic}»\n\n"
        f"{prompt_text}\n",
        encoding="utf-8",
    )
    return prompt_file


def write_img_pr_prompt_file(
    project: Project, tmp_dir: Path, *, ts: str | None = None
) -> Path:
    master = _get_master_or_fallback(
        project,
        "img_pr",
        "# img_pr\n\nМастер-промт для шага «Промты картинок» ещё не настроен.",
    )
    prompt_file = tmp_dir / f"prompt_img_pr_{ts or _timestamp()}.txt"
    prompt_file.write_text(master, encoding="utf-8")
    return prompt_file


def write_anim_pr_prompt_file(
    project: Project, tmp_dir: Path, *, ts: str | None = None
) -> Path:
    """Мастер-промт шага 8 «Промты анимации» — отдельный файл в ChatGPT."""
    master = _get_master_or_fallback(
        project,
        "anim_pr",
        "# anim_pr\n\nМастер-промт для шага «Промты анимации» ещё не настроен.",
    )
    prompt_file = tmp_dir / f"prompt_anim_pr_{ts or _timestamp()}.md"
    prompt_file.write_text(master, encoding="utf-8")
    return prompt_file


def chat_message(project: Project, step_code: str, **ctx) -> str:
    """Текст сообщения в ChatGPT (без мастер-промта)."""
    return gtb.get_effective_text(project, step_code, **ctx).strip()


async def ask_with_prompt_files(
    gpt: ChatGPTBot,
    chat_msg: str,
    attachments: list[Path],
    *,
    timeout: int = 1800,
    project_id: int | None = None,
) -> str:
    """Deprecated: используйте xlsx_gpt_flow.telegram_style_* (bot — источник правды)."""
    await gpt.new_conversation()
    return await gpt.ask_with_files(
        (chat_msg or "").strip(),
        attachments,
        timeout=timeout,
        project_id=project_id,
        expect_file_download=True,
    )


async def download_and_replace_xlsx(
    gpt: ChatGPTBot,
    project_xlsx: Path,
    download_path: Path,
    *,
    timeout: int = 600,
) -> None:
    await gpt.download_attachment_from_last_reply(download_path, timeout=timeout)
    validation_err = validate_xlsx(download_path)
    if validation_err is not None:
        raise RuntimeError(f"скачанный xlsx невалиден: {validation_err}")
    backup_to_old(project_xlsx)
    replace_with(project_xlsx, download_path)


async def download_text_attachment(
    gpt: ChatGPTBot,
    target: Path,
    *,
    timeout: int = 1800,
) -> str:
    await gpt.download_attachment_from_last_reply(target, timeout=timeout)
    if not target.exists() or target.stat().st_size < 10:
        raise RuntimeError(f"скачанный текстовый файл пустой: {target}")
    return target.read_text(encoding="utf-8").strip()


async def sync_project_xlsx(
    session: AsyncSession,
    project: Project,
    xlsx_path: Path,
    *,
    keep_fields: bool = False,
    update_frames_voiceover: bool = False,
) -> dict | None:
    """Импортирует project.xlsx в БД (v8 + fallback v7)."""
    validation_err = validate_xlsx(xlsx_path)
    if validation_err is not None:
        raise RuntimeError(validation_err)

    sync_info: dict | None = None
    v8_error: Exception | None = None
    v7_error: Exception | None = None
    try:
        from openpyxl import load_workbook

        wb = load_workbook(
            filename=str(xlsx_path), data_only=True, read_only=True
        )
        from app.services.xlsx_v8_import import has_v8_plan_sheet

        is_v8 = has_v8_plan_sheet(wb)
        wb.close()
    except Exception as e:  # noqa: BLE001
        is_v8 = False
        logger.warning(
            "[#{}] sync_project_xlsx: cannot peek sheet names: {}",
            project.id,
            e,
        )

    if is_v8:
        try:
            sync_info = await import_v8_xlsx(
                session,
                project,
                xlsx_path,
                keep_fields=keep_fields,
                update_frames_voiceover=update_frames_voiceover,
            )
            _log_sync_result(project.id, "v8", sync_info)
        except Exception as e:  # noqa: BLE001
            v8_error = e
            logger.warning(
                "[#{}] sync_project_xlsx v8 failed: {}", project.id, e
            )
            if update_frames_voiceover:
                raise RuntimeError(f"xlsx-sync v8: {e}") from e
    try:
        info_v7 = await reload_from_xlsx(session, project, xlsx_path)
        _log_sync_result(project.id, "v7", info_v7)
        if sync_info is None:
            sync_info = info_v7
    except Exception as e:  # noqa: BLE001
        v7_error = e
        logger.warning("[#{}] sync_project_xlsx v7 failed: {}", project.id, e)

    if sync_info is None:
        parts: list[str] = []
        if is_v8 and v8_error is not None:
            parts.append(f"v8: {v8_error}")
        if v7_error is not None:
            parts.append(f"v7: {v7_error}")
        msg = "xlsx-sync: импорт не удался"
        if parts:
            msg += f" ({'; '.join(parts)})"
        from app.services.run_sync import mark_running_node_failed

        await mark_running_node_failed(
            session, project, msg[:2000], initiator="worker"
        )
        raise RuntimeError(msg)

    if sync_info.get("error") and update_frames_voiceover:
        raise RuntimeError(f"xlsx-sync: {sync_info['error']}")
    return sync_info


def _sync_voiceover_from_script_text(project: Project) -> Path | None:
    """Возвращает путь к voiceover.txt, синхронизируя из script_text / бэкапа при необходимости."""
    voiceover_path = project.data_dir / "voiceover.txt"
    if voiceover_path.is_file() and voiceover_path.stat().st_size > 0:
        return voiceover_path
    text = (project.script_text or "").strip()
    if text:
        voiceover_path.parent.mkdir(parents=True, exist_ok=True)
        voiceover_path.write_text(text, encoding="utf-8")
        logger.info(
            "[#{}] ensure_current_voiceover: voiceover.txt из script_text ({} симв)",
            project.id,
            len(text),
        )
        return voiceover_path
    old_dir = voiceover_path.parent / "old"
    if old_dir.is_dir():
        backups = sorted(old_dir.glob("*_voiceover.txt"), reverse=True)
        for backup in backups:
            if backup.is_file() and backup.stat().st_size > 0:
                voiceover_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(backup, voiceover_path)
                logger.info(
                    "[#{}] ensure_current_voiceover: voiceover.txt из бэкапа {}",
                    project.id,
                    backup.name,
                )
                return voiceover_path
    return None


def ensure_current_voiceover(project: Project) -> Path | None:
    """Актуальный voiceover для split/music и шагов после script."""
    return _sync_voiceover_from_script_text(project)


def ensure_script_input_voiceover(project: Project) -> Path | None:
    """Исходный voiceover для шага script — самый ранний бэкап, иначе текущий файл."""
    from app.services.voiceover_recovery import oldest_voiceover_backup

    oldest = oldest_voiceover_backup(project)
    if oldest is not None:
        return oldest
    return ensure_current_voiceover(project)


def ensure_source_voiceover(project: Project) -> Path | None:
    """Обратная совместимость: актуальный voiceover (не исходный черновик)."""
    return ensure_current_voiceover(project)


def save_voiceover_text(project: Project, voiceover_path: Path, text: str) -> None:
    """Сохраняет voiceover.txt с бэкапом предыдущей версии."""
    voiceover_path.parent.mkdir(parents=True, exist_ok=True)
    if voiceover_path.exists():
        old_dir = voiceover_path.parent / "old"
        old_dir.mkdir(parents=True, exist_ok=True)
        backup = old_dir / f"{_timestamp()}_voiceover.txt"
        shutil.copy2(voiceover_path, backup)
    voiceover_path.write_text(text, encoding="utf-8")
