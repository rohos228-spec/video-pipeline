"""R0c-B: проверки не удаляют сцены и кадры — пометка + WARN, данные на месте."""

from __future__ import annotations

from pathlib import Path

import pytest
from loguru import logger
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.models import (
    Artifact,
    ArtifactKind,
    Base,
    Frame,
    FrameStatus,
    FrameText,
    Project,
    ProjectStatus,
    PromptVersion,
)


@pytest.fixture
def warn_log():
    lines: list[str] = []
    sink = logger.add(lambda message: lines.append(str(message)), level="WARNING")
    try:
        yield lines
    finally:
        logger.remove(sink)


def _ids(rows: list[dict]) -> list[str]:
    return [str(row.get("id")) for row in rows]


def test_align_marks_foreign_room_and_keeps_shot_ids(warn_log) -> None:
    from app.services.scene_shot_grammar import align_shots_to_action_chain

    action = (
        "1. кабинет следствия — следователь открывает папку\n"
        "(Следователь открыл толстую следственную папку на столе.)"
    )
    shots = [
        {"id": "a", "сцена": 1, "место": "кабинет следствия", "действие": "открывает"},
        {"id": "lab", "сцена": 1, "место": "лаборатория криминалистики", "действие": "снимает"},
        {"id": "c", "сцена": 1, "место": "кабинет следствия", "действие": "кладёт"},
    ]
    ops = [{"frame_uuid": "u1", "fields": {"кадры": shots}}]
    frames = [{"uuid": "u1", "attrs": {"главное_действие": action}}]
    assert align_shots_to_action_chain(ops, frames) == 1
    kept = ops[0]["fields"]["кадры"]
    assert _ids(kept) == ["a", "lab", "c"]
    marked = kept[1]
    assert marked["карантин"] is True
    assert "чужое место" in marked["пометка"]
    assert any("WARN align_shots_to_action_chain" in line for line in warn_log)


def test_vo_cap_marks_extra_shots_without_dropping_them(warn_log) -> None:
    from app.services.scene_shot_grammar import merge_gpt_shots_with_action

    action = "1. кабинет — следователь открывает папку"
    shots = [
        {"id": "s1", "действие": "следователь открывает папку", "закадр": "Открыл папку."},
        {"id": "s2", "действие": "кладёт папку на стол", "закадр": "Положил папку."},
        {"id": "s3", "действие": "смотрит в окно", "закадр": "Смотрит в окно кабинета."},
    ]
    out, _err = merge_gpt_shots_with_action(shots, action, cell_number=1, vo="")
    assert out is not None
    assert _ids(out) == ["s1", "s2", "s3"]
    marked = [s for s in out if s.get("карантин")]
    assert {s["id"] for s in marked} >= {"s2", "s3"}
    assert any("WARN merge_gpt_shots_with_action" in line and "сверх лимита" in line for line in warn_log)


def test_obshiy_vid_shot_is_marked_not_dropped(warn_log) -> None:
    from app.services.scene_shot_grammar import merge_gpt_shots_with_action

    action = "1. кабинет — следователь открывает папку"
    shots = [
        {"id": "est", "действие": "общий вид кабинета без людей"},
        {"id": "act", "действие": "следователь открывает папку"},
    ]
    out, _err = merge_gpt_shots_with_action(shots, action, vo="")
    assert out is not None
    assert set(_ids(out)) == {"est", "act"}
    est = next(s for s in out if s["id"] == "est")
    assert est["карантин"] is True
    assert "общий вид" in est["пометка"]
    assert any("общий вид" in line for line in warn_log)


def test_repair_shot_vo_marks_empty_shell_and_keeps_it(warn_log) -> None:
    from app.services.scene_shot_grammar import repair_shot_vo_lengths

    vo = "Следователь открыл толстую папку на столе."
    shots = [
        {"id": "real", "действие": "открывает папку", "закадр": vo},
        {"id": "shell"},
    ]
    repair_shot_vo_lengths(shots, vo)
    assert _ids(shots) == ["real", "shell"]
    assert shots[1]["карантин"] is True
    assert "пустая оболочка" in shots[1]["пометка"]
    assert any("WARN repair_shot_vo_lengths" in line for line in warn_log)


def test_spread_vo_sentences_keeps_shot_ids() -> None:
    """L1793 в ревью — spread_vo_sentences: кадры не выкидывает."""
    from app.services.scene_shot_grammar import spread_vo_sentences

    shots = [
        {"id": "k1", "закадр": "Первое предложение. Второе предложение."},
        {"id": "k2", "закадр": ""},
    ]
    spread_vo_sentences(shots)
    assert _ids(shots) == ["k1", "k2"]
    assert len(shots) == 2


def test_trim_to_cap_marks_bridges_and_keeps_every_id(warn_log) -> None:
    from app.services.montage_scene_improve import _trim_to_cap

    shots = [
        {"id": "a", "роль": "действие"},
        {"id": "b", "роль": "мост"},
        {"id": "c", "роль": "перебивка"},
    ]
    out = _trim_to_cap(shots, 1)
    assert _ids(out) == ["a", "b", "c"]
    assert out[0].get("карантин") is not True
    assert out[1]["карантин"] is True and out[2]["карантин"] is True
    assert any("WARN _trim_to_cap" in line for line in warn_log)


def test_repair_shots_keeps_duplicate_and_unmatched_anchor(warn_log) -> None:
    from app.services.montage_scene_improve import repair_shots

    units = [{"n": 1, "бюджет": 2, "закадр": "Он идёт домой и открывает дверь."}]
    shots = [
        {
            "id": "a",
            "якорь_n": 1,
            "действие": "идёт домой",
            "закадр": "Он идёт домой",
            "роль": "действие",
            "план": "СРЕДНИЙ",
            "ракурс": "3/4",
            "объект": "тело",
            "порядок": 1,
        },
        {
            "id": "dup",
            "якорь_n": 1,
            "действие": "идёт домой",
            "закадр": "и открывает дверь.",
            "роль": "действие",
            "план": "КРУПНЫЙ",
            "ракурс": "фронт",
            "объект": "тело",
            "порядок": 2,
        },
        {
            "id": "lost",
            "якорь_n": 9,
            "действие": "смотрит в окно",
            "закадр": "смотрит",
            "роль": "следствие",
            "план": "КРУПНЫЙ",
            "ракурс": "сверху",
            "объект": "взгляд",
            "порядок": 3,
        },
    ]
    repair_shots(shots, units)
    assert _ids(shots) == ["a", "dup", "lost"]
    assert shots[1]["карантин"] is True
    assert "повтор" in shots[1]["пометка"]
    assert shots[1].get("fw_qc")
    assert shots[2]["карантин"] is True
    assert "не совпал" in shots[2]["пометка"]
    assert any("WARN repair_shots" in line for line in warn_log)


def test_assign_scene_vo_keeps_shots_when_voiceover_empty(warn_log) -> None:
    from app.services.shot_templates import _assign_scene_vo, fill_kadry_from_catalog

    shots = [
        {"id": "a", "закадр": "уже было", "действие": "стоит"},
        {"id": "b", "закадр": "", "действие": "садится"},
    ]
    out = _assign_scene_vo(shots, "")
    assert _ids(out) == ["a", "b"]
    assert all(s.get("карантин") for s in out)
    assert any("WARN _assign_scene_vo" in line for line in warn_log)
    filled = fill_kadry_from_catalog(
        [{"id": "keep-me", "сцена": 1, "действие": "герой стоит у стола", "закадр": ""}],
        "1. комната — герой стоит у стола",
        cell_number=1,
    )
    assert filled
    assert len(filled) >= 1


def test_drop_extra_stops_at_required_and_keeps_ids(warn_log) -> None:
    from app.services.shot_templates import _drop_extra_shots

    rows = [
        {"shot_id": "T1-K1", "required": 0},
        {"shot_id": "T1-K2", "required": 1},
        {"shot_id": "T1-K3", "required": 1},
        {"shot_id": "T1-K4", "required": 0},
    ]
    shots = [{"id": f"s{i}", "шаблон": "T1", "закадр": ""} for i in range(4)]
    out = _drop_extra_shots(shots, 1, rows)
    assert _ids(out) == ["s0", "s1", "s2", "s3"]
    assert out[1].get("карантин") is not True
    assert out[2].get("карантин") is not True
    assert out[3]["карантин"] is True
    assert any("required=1" in line for line in warn_log)


def test_qc_coalesce_keeps_ids_when_incoming_is_shorter(warn_log) -> None:
    from app.services.apply_ops_batches import coalesce_cell_shot_ops

    uid = "abc"
    frames = [
        {
            "uuid": uid,
            "кадры": [
                {"id": "1-S1-K1", "закадр": "раз", "действие": "старое"},
                {"id": "1-S1-K2", "закадр": "", "действие": "пауза"},
                {"id": "1-S1-K3", "закадр": "три", "действие": "хвост"},
            ],
        }
    ]
    ops = [
        {
            "frame_uuid": uid,
            "fields": {"кадры": [{"id": "1-S1-K1", "закадр": "раз", "действие": "новое"}]},
        }
    ]
    merged = coalesce_cell_shot_ops(ops, frames, patch=True)
    rows = merged[0]["fields"]["кадры"]
    assert _ids(rows) == ["1-S1-K1", "1-S1-K2", "1-S1-K3"]
    assert rows[0]["действие"] == "новое"
    assert rows[1]["карантин"] is True
    assert rows[2]["карантин"] is True
    assert any("WARN coalesce_cell_shot_ops" in line for line in warn_log)


def test_apply_generated_scenario_keeps_operator_blocks(warn_log) -> None:
    from app.services.nii67_card import apply_generated_scenario, default_card, insert_block

    card = insert_block(default_card(), "door1", 6, "custom")
    door = card["doors"][0]
    world = next(b for b in door["blocks"] if b["kind"] == "world_frame")
    world["text"] = "текст оператора"
    world["image_prompt"] = "промт оператора"
    world["image_name"] = "slot.png"
    before_ids = [b["id"] for b in door["blocks"]]
    scenario = {
        "theme": "Мокрый цех",
        "law": "Капли падают сами",
        "illustration": "герой выходит",
        "door_design": "стекло",
        "door_anim": "пар",
        "beats": [
            {
                "id": "entry",
                "text": "сгенерированный вход",
                "frames": [{"start_frame": "новый промт кадра"}],
            },
            {"id": "hook", "text": "сгенерированный выход"},
        ],
    }
    out = apply_generated_scenario(card, {"door1": scenario}, {"standard": "финал"})
    blocks = out["doors"][0]["blocks"]
    after_ids = [b["id"] for b in blocks]
    assert before_ids == [i for i in after_ids if i in before_ids]
    assert len(after_ids) == len(set(after_ids))
    saved = next(b for b in blocks if b["id"] == world["id"])
    assert saved["text"] == "текст оператора"
    assert saved["image_prompt"] == "промт оператора"
    assert saved["image_name"] == "slot.png"
    assert saved.get("пометка")
    custom = next(b for b in blocks if b["kind"] == "custom")
    assert custom["id"] in before_ids
    assert custom.get("карантин") is True
    assert any("WARN apply_generated_scenario" in line for line in warn_log)


@pytest.fixture
async def session(tmp_path: Path) -> AsyncSession:
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'keep.db'}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as s:
        yield s
    await engine.dispose()


@pytest.fixture
def project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Project:
    data_root = tmp_path / "data"
    data_root.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr("app.settings.settings.data_dir", str(data_root))
    return Project(
        slug="keep-frames",
        topic="t",
        hero_mode="auto",
        status=ProjectStatus.frames_ready,
        meta={},
    )


@pytest.mark.asyncio
async def test_coverage_delete_in_auto_apply_marks_frame(
    session: AsyncSession, project: Project, warn_log
) -> None:
    from app.services.frame_tombstones import frame_is_active
    from app.services.montage_coverage_ops import apply_coverage_op

    parent_uid = "aa" * 12
    parent = Frame(
        project_id=1,
        number=1,
        uuid=parent_uid,
        voiceover_text="ячейка",
        status=FrameStatus.planned,
        attrs={"camera_subdivide": {"role": "vo_parent", "parent_uuid": parent_uid}},
    )
    child = Frame(
        project_id=1,
        number=2,
        uuid="bb" * 12,
        voiceover_text="кусок",
        status=FrameStatus.image_generated,
        attrs={
            "camera_subdivide": {
                "role": "shot",
                "parent_uuid": parent_uid,
                "coverage_kind": "child",
            }
        },
    )
    session.add(project)
    await session.flush()
    parent.project_id = project.id
    child.project_id = project.id
    session.add_all([parent, child])
    await session.flush()
    session.add_all(
        [
            PromptVersion(
                project_id=project.id,
                frame_id=child.id,
                kind="img",
                version=1,
                text="промт ребёнка",
                is_active=True,
            ),
            FrameText(
                project_id=project.id,
                frame_id=child.id,
                kind="voiceover",
                text="кусок",
            ),
            Artifact(
                project_id=project.id,
                frame_id=child.id,
                kind=ArtifactKind.scene_image,
                uuid="art-child",
                path="scenes/frame_002.png",
            ),
        ]
    )
    await session.flush()
    child_id = int(child.id)
    await apply_coverage_op(
        session, project, {"type": "coverage_delete", "frame_number": 2}
    )
    await session.flush()
    rows = (
        await session.execute(select(Frame).where(Frame.project_id == project.id))
    ).scalars().all()
    assert sorted(fr.number for fr in rows) == [1, 2]
    live = await session.get(Frame, child_id)
    assert live is not None
    assert (live.attrs or {}).get("удалён") is True
    assert frame_is_active(live) is False
    assert live.status == FrameStatus.image_generated
    prompt = (
        await session.execute(select(PromptVersion).where(PromptVersion.frame_id == child_id))
    ).scalar_one()
    text = (
        await session.execute(select(FrameText).where(FrameText.frame_id == child_id))
    ).scalar_one()
    art = (
        await session.execute(select(Artifact).where(Artifact.uuid == "art-child"))
    ).scalar_one()
    assert prompt.text == "промт ребёнка"
    assert text.text == "кусок"
    assert art.frame_id is None
    assert 2 in (project.meta or {}).get("deleted_frame_numbers", [])
    assert any("запрещён в авто-apply" in line for line in warn_log)


@pytest.mark.asyncio
async def test_delete_montage_frame_marks_parent_and_keeps_children(
    session: AsyncSession, project: Project, warn_log
) -> None:
    from app.services.frame_tombstones import frame_is_active
    from app.services.montage_board_frames import delete_montage_frame

    session.add(project)
    await session.flush()
    parent = Frame(
        project_id=project.id,
        number=1,
        uuid="cc" * 12,
        voiceover_text="голова",
        status=FrameStatus.planned,
        attrs={"camera_subdivide": {"role": "vo_parent", "parent_uuid": "cc" * 12}},
    )
    child = Frame(
        project_id=project.id,
        number=2,
        uuid="dd" * 12,
        voiceover_text="шот",
        status=FrameStatus.planned,
        attrs={
            "camera_subdivide": {
                "role": "shot",
                "parent_uuid": "cc" * 12,
            }
        },
    )
    session.add_all([parent, child])
    await session.flush()
    parent_id, child_id = int(parent.id), int(child.id)
    result = await delete_montage_frame(session, project, parent_id)
    assert result["deleted"] == 0 and result["marked"] == 1
    left = (
        await session.execute(
            select(Frame).where(Frame.project_id == project.id).order_by(Frame.number)
        )
    ).scalars().all()
    assert [fr.id for fr in left] == [parent_id, child_id]
    assert frame_is_active(left[0]) is False
    assert (left[0].attrs or {}).get("пометка")
    assert frame_is_active(left[1]) is True
    assert any("WARN delete_montage_frame" in line for line in warn_log)


@pytest.mark.asyncio
async def test_marked_frame_is_not_revived_or_duplicated_by_xlsx_import(
    session: AsyncSession, project: Project, tmp_path: Path, warn_log
) -> None:
    from openpyxl import Workbook

    from app.services.frame_tombstones import (
        deleted_numbers,
        forget_deleted,
        frame_is_active,
        frame_marked_deleted,
        mark_frame_deleted,
    )
    from app.services.xlsx_sync import reload_from_xlsx
    from app.services.xlsx_v8_import import (
        ROW_IMAGE_PROMPT_V8,
        ROW_VOICEOVER_V8,
        bootstrap_frames_for_image_step,
        import_v8_xlsx,
    )
    from app.storage.project_sheet import (
        ROW_HEADER,
        ROW_IMAGE_PROMPT,
        ROW_VOICEOVER,
        SHEET_FRAMES,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "план"
    for i, text in enumerate(["один", "два", "три новых"], start=3):
        ws.cell(row=ROW_VOICEOVER_V8, column=i, value=text)
        ws.cell(row=ROW_IMAGE_PROMPT_V8, column=i, value=f"промт-{i}")
    kadry = wb.create_sheet(SHEET_FRAMES)
    for col, number in ((2, 1), (3, 2), (4, 3)):
        kadry.cell(row=ROW_HEADER, column=col, value=number)
    kadry.cell(row=ROW_VOICEOVER, column=4, value="три из листа Кадры")
    kadry.cell(row=ROW_IMAGE_PROMPT, column=4, value="промт который не должен ожить")
    xlsx = tmp_path / "project.xlsx"
    wb.save(xlsx)

    session.add(project)
    await session.flush()
    frames = []
    for number, text in ((1, "один"), (2, "два"), (3, "три старых")):
        frames.append(
            Frame(
                project_id=project.id,
                number=number,
                uuid=f"{number:02d}" * 12,
                voiceover_text=text,
                image_prompt="старый промт" if number == 3 else None,
                status=FrameStatus.planned,
            )
        )
    session.add_all(frames)
    await session.flush()
    marked = frames[2]
    marked_id = int(marked.id)
    mark_frame_deleted(marked, project, "тест: кадр помечен удалён")
    # Список в meta пуст: импорт обязан остановиться по пометке на кадре,
    # а не завести активную копию с тем же номером.
    forget_deleted(project, 3)
    await session.flush()
    assert deleted_numbers(project) == set()
    assert frame_marked_deleted(marked) is True

    summary = await import_v8_xlsx(
        session, project, xlsx, keep_fields=True, update_frames_voiceover=True
    )
    await bootstrap_frames_for_image_step(session, project, xlsx)
    sync = await reload_from_xlsx(session, project, xlsx)
    await session.flush()
    assert 3 in summary.get("frames_kept_marked", [])
    assert 3 not in sync.get("frames_changed", [])
    assert 3 not in sync.get("frames_created", [])

    rows = (
        await session.execute(
            select(Frame).where(Frame.project_id == project.id).order_by(Frame.number)
        )
    ).scalars().all()
    assert [fr.number for fr in rows] == [1, 2, 3]
    assert [fr.id for fr in rows] == [frames[0].id, frames[1].id, marked_id]
    live = rows[2]
    assert live.voiceover_text == "три старых"
    assert live.image_prompt == "старый промт"
    assert live.status == FrameStatus.planned
    assert (live.attrs or {}).get("удалён") is True
    assert frame_is_active(live) is False
    assert deleted_numbers(project) == {3}
    assert sum(1 for fr in rows if fr.number == 3) == 1
    assert any("помечен «удалён»" in line for line in warn_log)
