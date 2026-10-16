from __future__ import annotations
from app.services import revisions

import asyncio
from base64 import b64decode
from contextlib import contextmanager, ExitStack
from copy import deepcopy
from importlib import import_module
import json
from pathlib import Path
import unittest
from unittest.mock import AsyncMock, patch


EXECUTION = "00000000-0000-4000-8000-000000000007"
IMAGE_EXECUTION = "00000000-0000-4000-8000-000000000008"
ATTEMPT = "00000000-0000-4000-8000-000000000010"
REQUEST_A = "00000000-0000-4000-8000-000000000011"
REQUEST_B = "00000000-0000-4000-8000-000000000022"
REQUEST_RACE = "00000000-0000-4000-8000-000000000099"
ANIMATION = "The camera slowly tracks the person opening the blue book."
OUTPUT_FIELDS = ("id", "project_id", "uuid", "number", "sort_key", "scene_id",
                 "voiceover_text", "image_prompt", "attrs", "animation_prompt")


class NativeAnimationObservation(unittest.IsolatedAsyncioTestCase):
    fixture_node_type = "animation_prompts"

    async def asyncSetUp(self):
        from test_actions_v1_contract import DurableActions
        from sqlalchemy import select
        from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
        from app.models import Base, Frame, NodeRun, Project, ProjectStatus, WorkflowRun
        from app.services import observation
        from app.services.node_status_machine import queue_node_for_start, start_node_running
        from app.settings import settings

        await DurableActions.asyncSetUp(self)
        self.settings_patch = patch.object(settings, "data_dir", Path(self.tmp.name))
        self.settings_patch.start()
        self.observer = observation.get_observer()
        async with self.sessions() as master:
            project = await master.get(Project, 7)
            project.status = ProjectStatus.generating_animation_prompts
            project.meta = {"canvas_graph": {"nodes": self.nodes, "edges": []}}
            node = await master.scalar(select(NodeRun).where(NodeRun.node_key == "n1"))
            queue_node_for_start(node, project_id=7, initiator="worker")
            start_node_running(node, project_id=7, initiator="worker")
            scope = await master.get(observation.ExecutionControl, "exec-durable")
            scope.execution_id = EXECUTION
            scope.node_keys = ["n1"]
            run = await master.get(WorkflowRun, scope.workflow_run_id)
            self.native_node_type = next(item["type"] for item in run.nodes_snapshot
                                         if item["id"] == node.node_key)
            self.assertEqual(self.native_node_type, node.node_type)
            await master.commit()
            self.scope = scope
        self.observer.scopes[self.scope.workflow_run_id] = self.scope

        directory = Path(self.tmp.name) / "videos" / "actions-fixture"
        directory.mkdir(parents=True)
        self.database = directory / "project.db"
        self.business_engine = create_async_engine(f"sqlite+aiosqlite:///{self.database}")
        self.business_sessions = async_sessionmaker(self.business_engine, expire_on_commit=False)
        async with self.business_engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
            await observation.migrate_observation_schema(connection)
        async with self.business_sessions() as seed:
            seed.add_all([
                Project(id=7, slug="actions-fixture", topic="synthetic",
                    status=ProjectStatus.generating_animation_prompts,
                    meta={"canvas_graph": {"nodes": self.nodes, "edges": []}}),
                Project(id=8, slug="foreign-animation-fixture", topic="outside scope", meta={}),
                Frame(project_id=7, number=1, sort_key=2.0, uuid="af" * 16,
                    voiceover_text="Альфа открывает книгу.\n", image_prompt="A person opens a blue book.",
                    attrs={"existing_note": {"keep": ["first"]}}),
                Frame(project_id=7, number=2, sort_key=1.5, uuid="bf" * 16,
                    voiceover_text="Бета закрывает книгу.  ", image_prompt="A person closes a blue book.",
                    attrs={"existing_note": {"keep": ["second"]}}),
                Frame(project_id=8, number=1, uuid="ef" * 16, voiceover_text="Foreign input.",
                    image_prompt="Foreign usable image prompt.", attrs={}),
            ])
            await seed.commit()

    async def asyncTearDown(self):
        from test_actions_v1_contract import DurableActions

        try:
            await self.business_engine.dispose()
        finally:
            self.settings_patch.stop()
            await DurableActions.asyncTearDown(self)

    def api(self, name):
        module = import_module("app.services.native_animation_contracts")
        function = getattr(module, name, None)
        self.assertTrue(callable(function), "missing bounded animation observation API: " + name)
        return function

    async def inputs(self, session):
        from sqlalchemy import select
        from app.models import Frame, Project

        project = await session.get(Project, 7)
        frames = list((await session.scalars(select(Frame).where(Frame.project_id == 7)
                       .order_by(Frame.number))).all())
        return project, frames

    def parsed(self, project, frames, *, chosen=None, reply=None, shot=1):
        native = import_module("app.services.animation_prompt_gpt")
        if reply is None:
            selected = frames if chosen is None else chosen
            reply = json.dumps({"ops": [{"frame_uuid": frame.uuid,
                "fields": {"промт_видео": f"Frame {frame.number}: {ANIMATION}"}}
                for frame in selected]}, ensure_ascii=False)
        batch = native.collect_batch_items(project, frames)
        pairs = native.parse_animation_reply(reply, frames, batch_items=batch)
        ops = native.build_apply_ops_from_pairs(pairs, frames, shot=shot)
        return reply, pairs, ops

    @contextmanager
    def context(self, session, *, completed=None):
        from app.services import observation

        ctx = observation.AttemptContext(self.observer, self.scope, session,
                                        "project", "n1", ATTEMPT)
        ctx.native_node_type = self.native_node_type
        ctx.last_request_id = REQUEST_RACE
        attempt_token = observation._attempt.set(ctx)
        request_token = observation._request.set(None)
        completed_token = observation._completed_model_request.set((ctx, completed) if completed else None)
        try:
            yield ctx
        finally:
            observation._completed_model_request.reset(completed_token)
            observation._request.reset(request_token)
            observation._attempt.reset(attempt_token)

    @contextmanager
    def business_sql(self):
        from sqlalchemy import event
        from sqlalchemy.engine import Engine

        statements = []

        def capture(connection, cursor, statement, parameters, context, executemany):
            if Path(str(connection.engine.url.database)).resolve() == self.database.resolve():
                statements.append(statement)

        event.listen(Engine, "before_cursor_execute", capture)
        try:
            yield statements
        finally:
            event.remove(Engine, "before_cursor_execute", capture)

    def no_business_writes(self, statements):
        writes = [statement for statement in statements if statement.lstrip().split(None, 1)[0].upper()
                  in {"INSERT", "UPDATE", "DELETE", "REPLACE", "CREATE", "ALTER", "DROP"}]
        self.assertEqual(writes, [], "observation may read, but must not write/autoflush business data")

    async def state(self):
        from sqlalchemy import select
        from app.models import Frame, Project

        async with self.business_sessions() as reader:
            frames = [dict(row) for row in (await reader.execute(select(Frame.__table__)
                .order_by(Frame.project_id, Frame.number))).mappings().all()]
            project = dict((await reader.execute(select(Project.__table__).where(Project.id == 7))).mappings().one())
        return {"project": project, "frames": frames}

    async def facts(self, execution_id=EXECUTION):
        from sqlalchemy import select
        from app.services import observation

        await self.observer.consume(execution_id, "master", self.sessions)
        async with self.sessions() as reader:
            return list((await reader.scalars(select(observation.ObservationEvent).where(
                observation.ObservationEvent.execution_id == execution_id)
                .order_by(observation.ObservationEvent.id))).all())

    def one(self, facts, kind):
        found = [fact for fact in facts if fact.kind == kind]
        self.assertEqual(len(found), 1, (kind, [(fact.kind, fact.payload) for fact in facts]))
        self.assertEqual((found[0].execution_id, found[0].project_id, found[0].node_key),
                         (EXECUTION, 7, "n1"))
        return found[0]

    def body(self, fact):
        reference = fact.payload["body"]
        self.assertNotIn("gap", reference)
        path = (self.observer.body_root / reference["path"]).resolve()
        self.assertTrue(path.is_relative_to(self.observer.body_root.resolve()))
        return json.loads(path.read_text(encoding="utf-8"))

    def unclaimed(self, proof):
        for key in ("managed_eligible", "business_commit_claim", "semantic_quality_verified", "video_quality_verified"):
            self.assertIs(proof[key], False, proof)

    async def unchanged_authority(self):
        from sqlalchemy import select
        from app.models import NodeRun, NodeRunStatus, Project, ProjectStatus
        from app.services.actions_v1 import ActionExecutionControl, ActionRecord, get_permission
        from app.services.node_capabilities import capability_for
        from app.services.observation import ExecutionControl

        self.assertNotIn("retry_profile", capability_for("animation_prompts").actions)
        self.assertIsNone(get_permission(7))
        async with self.sessions() as reader:
            self.assertEqual((await reader.get(ExecutionControl, EXECUTION)).mode, "observe")
            self.assertIsNone(await reader.get(ActionExecutionControl, EXECUTION))
            self.assertEqual(list((await reader.scalars(select(ActionRecord))).all()), [])
            self.assertEqual((await reader.get(Project, 7)).status, ProjectStatus.generating_animation_prompts)
            states = {node.node_key: node.status for node in (await reader.scalars(select(NodeRun))).all()}
            self.assertEqual(states, {"n1": NodeRunStatus.running, "n2": NodeRunStatus.pending})

    async def apply_native(self, session, project, ops):
        from app.services.db_apply import apply_ops

        result = await apply_ops(session, project, ops, export_xlsx=False, node_kind="anim_pr")
        self.assertTrue(result["ok"], result)
        await session.commit()

    def declare_native_second_shot(self, project, frames):
        from openpyxl import Workbook
        from app.services.xlsx_v8_import import ROW_IMAGE_PROMPT_V8, ROW_VOICEOVER_V8

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "план"
        for frame in frames:
            column = frame.number + 2
            sheet.cell(ROW_VOICEOVER_V8, column, frame.voiceover_text)
            sheet.cell(ROW_IMAGE_PROMPT_V8, column, frame.image_prompt)
        sheet.cell(46, frames[0].number + 2, "The same person seen from a second camera angle.")
        workbook.save(project.data_dir / "project.xlsx")
        workbook.close()
        directory = project.data_dir / "scenes"
        directory.mkdir(exist_ok=True)
        (directory / "frame_001_s2_native.png").write_bytes(b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
            "/x8AAwMCAO+aL1sAAAAASUVORK5CYII="))

    async def test_public_observation_seams_are_present(self):
        for name in ("record_animation_input", "record_animation_parse", "observe_animation_result"):
            with self.subTest(api=name):
                self.api(name)

    async def test_default_off_none_context_touches_no_inputs_sql_bodies_or_events(self):
        from app.services import observation

        record_input = self.api("record_animation_input")
        record_parse = self.api("record_animation_parse")
        observe_result = self.api("observe_animation_result")
        self.assertIsNone(observation._attempt.get())
        before = await self.state()
        bodies = sorted(self.observer.body_root.glob("**/*.json"))
        with self.business_sql() as statements:
            await record_input(None, None)
            await record_parse(None, None, None, shot=1, target_uuids=[])
            await observe_result(None, None)
        self.assertEqual(statements, [])
        self.assertEqual(await self.facts(), [])
        self.assertEqual(sorted(self.observer.body_root.glob("**/*.json")), bodies)
        self.assertEqual(await self.state(), before)
        await self.unchanged_authority()

    async def test_input_body_is_detached_original_snapshot_without_business_write(self):
        from app.services.native_animation_contracts import animation_prompt_input_snapshot
        from app.services.observation import digest

        record_input = self.api("record_animation_input")
        before = await self.state()
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            self.declare_native_second_shot(project, frames)
            expected = animation_prompt_input_snapshot(frames, project_id=7, project=project)
            self.assertIs(expected["frames"][0]["required_shot2"], True)
            with self.context(business), self.business_sql() as statements:
                await record_input(project, frames)
            self.no_business_writes(statements)
            frames[0].attrs["existing_note"]["keep"].append("uncommitted mutation after capture")
        fact = self.one(await self.facts(), "node_input")
        self.assertEqual(fact.payload["contract"], "native_animation_prompts")
        self.assertEqual(fact.payload["frame_count"], 2)
        self.assertEqual(fact.payload["input_sha256"], digest(expected))
        self.assertEqual(self.body(fact), expected)
        self.assertEqual(await self.state(), before)
        await self.unchanged_authority()

    async def test_sidecar_input_preserves_finalize_false_in_native_capture(self):
        from app.services import animation_prompt_gpt as native

        record_input = self.api("record_animation_input")
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            self.declare_native_second_shot(project, frames)
            self.assertTrue(native.frame_needs_shot2_video_prompt(project, frames[0]))
            with self.context(business):
                await record_input(project, frames, finalize_status=False)
        snapshot = self.body(self.one(await self.facts(), "node_input"))
        self.assertIs(snapshot["finalize_status"], False)
        self.assertTrue(all(row["required_prompt"] for row in snapshot["frames"]))
        self.assertFalse(any(row["required_shot2"] for row in snapshot["frames"]))
        await self.unchanged_authority()

    async def test_valid_native_parse_binds_completed_request_ops_and_does_not_apply(self):
        from app.services.observation import digest

        record_parse = self.api("record_animation_parse")
        before = await self.state()
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            reply, pairs, ops = self.parsed(project, frames, chosen=list(reversed(frames)))
            self.assertEqual([op["frame_uuid"] for op in ops], [frames[1].uuid, frames[0].uuid])
            original_pairs, original_ops = deepcopy(pairs), deepcopy(ops)
            with self.context(business, completed=REQUEST_A), self.business_sql() as statements:
                await record_parse(reply, pairs, ops, shot=1, target_uuids=[frame.uuid for frame in frames])
            self.assertEqual(pairs, original_pairs)
            self.assertEqual(ops, original_ops)
            self.no_business_writes(statements)
        facts = await self.facts()
        parsed = self.one(facts, "parse_result")
        self.assertEqual(parsed.request_id, REQUEST_A)
        self.assertEqual(parsed.payload["contract"], "native_animation_prompts")
        for key in ("syntax_valid", "parsed", "schema_valid", "batch_coverage_complete"):
            self.assertIs(parsed.payload[key], True)
        self.assertEqual(parsed.payload["shot"], 1)
        self.assertEqual(parsed.payload["sha256"], digest(ops))
        self.assertEqual(self.body(parsed)["extracted_ops"], ops)
        self.assertEqual(self.body(parsed)["text"], reply)
        self.assertFalse(any(fact.kind in {"apply_input", "data_applied"} for fact in facts))
        self.assertEqual(await self.state(), before)
        await self.unchanged_authority()

    async def test_partial_native_parse_keeps_valid_extraction_and_incomplete_coverage(self):
        record_parse = self.api("record_animation_parse")
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            reply, pairs, ops = self.parsed(project, frames, chosen=frames[1:])
            self.assertEqual([pair.frame_number for pair in pairs], [2])
            with self.context(business, completed=REQUEST_A):
                await record_parse(reply, pairs, ops, shot=1, target_uuids=[frame.uuid for frame in frames])
        fact = self.one(await self.facts(), "parse_result")
        self.assertIs(fact.payload["parsed"], True)
        self.assertIs(fact.payload["schema_valid"], True)
        self.assertIs(fact.payload["batch_coverage_complete"], False)
        self.assertEqual(self.body(fact)["extracted_ops"], ops)
        await self.unchanged_authority()

    async def test_parallel_parse_uses_task_local_completion_not_shared_last_request(self):
        from app.services import observation
        from app.services.observation import digest

        record_parse = self.api("record_animation_parse")
        expected = {}
        ready = asyncio.Event()
        arrivals = 0
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            with self.context(business) as ctx:
                async def parse_one(frame, request_id):
                    nonlocal arrivals
                    token = observation._completed_model_request.set((ctx, request_id))
                    try:
                        reply, pairs, ops = self.parsed(project, frames, chosen=[frame])
                        expected[request_id] = digest(ops)
                        ctx.last_request_id = request_id
                        arrivals += 1
                        if arrivals == 2:
                            ctx.last_request_id = REQUEST_RACE
                            ready.set()
                        await ready.wait()
                        self.assertEqual(observation.completed_model_request_id(), request_id)
                        await record_parse(reply, pairs, ops, shot=1, target_uuids=[frame.uuid])
                    finally:
                        observation._completed_model_request.reset(token)
                await asyncio.gather(parse_one(frames[0], REQUEST_A), parse_one(frames[1], REQUEST_B))
        facts = [fact for fact in await self.facts() if fact.kind == "parse_result"]
        self.assertEqual(len(facts), 2)
        self.assertEqual({fact.request_id: fact.payload["sha256"] for fact in facts}, expected)
        self.assertNotIn(REQUEST_RACE, {fact.request_id for fact in facts})
        await self.unchanged_authority()

    async def test_invalid_actual_native_replies_emit_false_parse_without_apply(self):
        record_parse = self.api("record_animation_parse")
        before = await self.state()
        replies = ('{"ops": }', json.dumps({"ops": [{"frame_uuid": "foreign",
                   "fields": {"промт_видео": ANIMATION}}]}))
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            with self.context(business, completed=REQUEST_A), self.business_sql() as statements:
                for reply in replies:
                    native_reply, pairs, ops = self.parsed(project, frames, reply=reply)
                    self.assertEqual(pairs, [])
                    self.assertEqual(ops, [])
                    await record_parse(native_reply, pairs, ops, shot=1, target_uuids=[frame.uuid for frame in frames])
            self.no_business_writes(statements)
        facts = await self.facts()
        parses = [fact for fact in facts if fact.kind == "parse_result"]
        self.assertEqual(len(parses), 2)
        self.assertIs(parses[0].payload["syntax_valid"], False)
        for fact in parses:
            self.assertEqual(fact.request_id, REQUEST_A)
            for key in ("parsed", "schema_valid", "batch_coverage_complete"):
                self.assertIs(fact.payload[key], False)
            self.assertEqual(self.body(fact)["extracted_ops"], [])
            self.assertIs(fact.payload["payload_required"], True)
            self.assertIs(fact.payload["payload_present"], False)
        from types import SimpleNamespace
        from app.services.diagnosis import derive_incidents
        scope = SimpleNamespace(project_id=7, execution_id=EXECUTION, node_keys=["n1"],
                                observer_state="active", gaps=[])
        rows = [{key: getattr(fact, key) for key in ("id", "event_id", "execution_id", "project_id", "node_key",
                 "attempt_id", "request_id", "kind", "db_scope", "payload", "created_at")} for fact in facts]
        self.assertIn("invalid_json", {item["type"] for item in derive_incidents(scope, rows, {"nodes": {}})})
        self.assertFalse(any(fact.kind in {"apply_input", "data_applied"} for fact in facts))
        checks = [fact for fact in facts if fact.kind == "runtime_check"]
        self.assertTrue(checks)
        self.assertTrue(all(fact.payload["status"] == "failed" for fact in checks))
        self.assertEqual(await self.state(), before)
        await self.unchanged_authority()

    async def test_filled_committed_output_without_captured_input_stays_unknown(self):
        observe_result = self.api("observe_animation_result")
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            _, _, ops = self.parsed(project, frames)
            await self.apply_native(business, project, ops)
            before = await self.state()
            with self.context(business), self.business_sql() as statements:
                await observe_result(business, project)
            self.no_business_writes(statements)
        facts = await self.facts()
        proof = self.one(facts, "node_result_verification").payload
        self.assertEqual(proof["status"], "unknown")
        self.unclaimed(proof)
        self.assertTrue(all(frame["animation_prompt"] for frame in self.body(self.one(facts, "node_output"))["frames"]))
        self.assertEqual(await self.state(), before)
        await self.unchanged_authority()

    async def test_actual_committed_native_results_verify_once_without_managed_grant(self):
        record_input = self.api("record_animation_input")
        observe_result = self.api("observe_animation_result")
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            with self.context(business):
                await record_input(project, frames)
                _, _, ops = self.parsed(project, frames)
                await self.apply_native(business, project, ops)
                before = await self.state()
                with self.business_sql() as statements:
                    await observe_result(business, project, node_type="animation_prompts")
                self.no_business_writes(statements)
        facts = await self.facts()
        proof = self.one(facts, "node_result_verification").payload
        self.assertEqual(proof["status"], "passed")
        self.assertEqual(proof["target_frame_count"], 2)
        self.unclaimed(proof)
        output = self.body(self.one(facts, "node_output"))
        self.assertEqual(output["project_id"], 7)
        self.assertEqual([{key: frame[key] for key in OUTPUT_FIELDS} for frame in output["frames"]],
            [{key: frame[key] for key in OUTPUT_FIELDS} for frame in before["frames"] if frame["project_id"] == 7])
        self.assertEqual(await self.state(), before)
        await self.unchanged_authority()

    async def test_real_observe_attempt_success_calls_result_exactly_once(self):
        from app.services import observation

        record_input = self.api("record_animation_input")
        result = self.api("observe_animation_result")
        module = import_module("app.services.native_animation_contracts")
        spy = AsyncMock(wraps=result)
        with patch.object(module, "observe_animation_result", new=spy):
            async with self.business_sessions() as business:
                project, frames = await self.inputs(business)
                async with observation.observe_attempt(business, project):
                    await record_input(project, frames)
                    _, _, ops = self.parsed(project, frames)
                    await self.apply_native(business, project, ops)
                self.assertEqual(spy.await_count, 1)
                call = spy.await_args
                self.assertIs(call.args[0] if call.args else call.kwargs["session"], business)
                self.assertIs(call.args[1] if len(call.args) > 1 else call.kwargs["project"], project)
        facts = await self.facts()
        self.one(facts, "node_start")
        self.one(facts, "node_end")
        self.assertEqual(self.one(facts, "node_result_verification").payload["status"], "passed")
        self.assertFalse(any(fact.kind == "node_error" for fact in facts))
        await self.unchanged_authority()

    async def test_real_sql_failure_reads_committed_partial_output_without_repair_or_error_replacement(self):
        from sqlalchemy.exc import IntegrityError, PendingRollbackError
        from sqlalchemy import select
        from app.models import Frame, ProjectStatus
        from app.services import observation

        record_input = self.api("record_animation_input")
        result = self.api("observe_animation_result")
        module = import_module("app.services.native_animation_contracts")
        spy = AsyncMock(wraps=result)
        original_error = None
        with patch.object(module, "observe_animation_result", new=spy):
            async with self.business_sessions() as business:
                project, frames = await self.inputs(business)
                with ExitStack() as diagnostics:
                    with self.assertRaises(IntegrityError) as caught:
                        async with observation.observe_attempt(business, project):
                            await record_input(project, frames)
                            _, _, ops = self.parsed(project, frames, chosen=frames[:1])
                            await self.apply_native(business, project, ops)
                            committed = await self.state()
                            frames[1].voiceover_text = "Uncommitted rewritten voiceover."
                            frames[1].animation_prompt = "Uncommitted second animation."
                            business.add(Frame(project_id=7, number=1, uuid="cf" * 16,
                                               voiceover_text="Duplicate native number.", attrs={}))
                            try:
                                await business.flush()
                            except IntegrityError as exc:
                                original_error = exc
                                self.assertFalse(business.is_active)
                                statements = diagnostics.enter_context(self.business_sql())
                                rollback = diagnostics.enter_context(patch.object(business, "rollback", new=AsyncMock(wraps=business.rollback)))
                                commit = diagnostics.enter_context(patch.object(business, "commit", new=AsyncMock(wraps=business.commit)))
                                flush = diagnostics.enter_context(patch.object(business, "flush", new=AsyncMock(wraps=business.flush)))
                                raise
                            self.fail("real SQLite must reject duplicate (project_id, number)")
                    self.assertIs(caught.exception, original_error)
                    self.assertEqual(spy.await_count, 1)
                    rollback.assert_not_awaited()
                    commit.assert_not_awaited()
                    flush.assert_not_awaited()
                    self.assertFalse(business.is_active)
                    with self.assertRaises(PendingRollbackError):
                        await business.execute(select(Frame.id))
                    self.assertTrue(any(statement.lstrip().upper().startswith("SELECT") and "frames" in statement.lower()
                                        for statement in statements), "a separate reader must inspect committed native frames")
                    self.no_business_writes(statements)
        self.assertEqual(await self.state(), committed)
        self.assertEqual(committed["project"]["status"], ProjectStatus.generating_animation_prompts)
        facts = await self.facts()
        proof = self.one(facts, "node_result_verification").payload
        self.assertEqual(proof["status"], "failed", "the second required committed prompt is still missing")
        self.unclaimed(proof)
        output = self.body(self.one(facts, "node_output"))
        self.assertEqual([{key: frame[key] for key in OUTPUT_FIELDS} for frame in output["frames"]],
            [{key: frame[key] for key in OUTPUT_FIELDS} for frame in committed["frames"] if frame["project_id"] == 7])
        self.one(facts, "node_error")
        self.assertFalse(any(fact.kind == "node_end" for fact in facts))
        await self.unchanged_authority()

    async def test_observe_attempt_diagnostic_failure_never_replaces_original_native_error(self):
        from app.services import observation

        self.api("observe_animation_result")
        module = import_module("app.services.native_animation_contracts")
        original_error = RuntimeError("original native animation failure")
        diagnostic_error = LookupError("bounded result evidence unavailable")
        spy = AsyncMock(side_effect=diagnostic_error)
        with patch.object(module, "observe_animation_result", new=spy):
            async with self.business_sessions() as business:
                project, _ = await self.inputs(business)
                with self.assertRaises(RuntimeError) as caught:
                    async with observation.observe_attempt(business, project):
                        raise original_error
                self.assertIs(caught.exception, original_error)
                self.assertEqual(spy.await_count, 1)
                self.assertTrue(business.is_active)
        self.assertIn(EXECUTION, self.observer.failures)
        self.one(await self.facts(), "node_error")
        await self.unchanged_authority()

    async def test_valid_native_save_marks_parse_diagnostic_gap_and_preserves_actual_apply(self):
        from sqlalchemy import select
        from app.models import FrameStatus, PromptVersion
        from app.orchestrator.steps import make_animation_prompts as step
        from app.services import animation_prompt_gpt as native

        module = import_module("app.services.native_animation_contracts")
        self.api("record_animation_parse")
        diagnostic_error = LookupError("native batch parse journal unavailable")
        parse_spy = AsyncMock(side_effect=diagnostic_error)
        apply_spy = AsyncMock(wraps=step.db_apply.apply_ops)
        before = await self.state()
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            reply, pairs, ops = self.parsed(project, frames)
            expected = {op["frame_uuid"]: op["fields"]["промт_видео"] for op in ops}
            expected_versions = {frame.id: expected[frame.uuid] for frame in frames}
            batch = native.collect_batch_items(project, frames)
            self.assertEqual(len(batch), 2)
            with self.context(business, completed=REQUEST_A), \
                    patch.object(module, "record_animation_parse", new=parse_spy), \
                    patch.object(step.db_apply, "apply_ops", new=apply_spy):
                try:
                    delivered = await step._save_anim_pr_batch(
                        business, project, frames, batch, reply, None, shot=1)
                except Exception as exc:
                    self.fail("diagnostic fault interrupted a valid original batch: "
                              + type(exc).__name__ + ": " + str(exc))
            self.assertEqual(delivered, 2)
            self.assertTrue(business.is_active)
            self.assertEqual(native.collect_batch_items(project, frames), [])
        parse_spy.assert_awaited_once_with(
            reply, pairs, ops, shot=1, target_uuids=[frame.uuid for frame in frames])
        apply_spy.assert_awaited_once()
        self.assertEqual(apply_spy.await_args.args[2], ops)
        self.assertEqual(apply_spy.await_args.kwargs,
                         {"export_xlsx": False, "node_kind": "anim_pr"})
        self.assertEqual(self.observer.failures.get(EXECUTION), "LookupError")
        after = await self.state()
        self.assertEqual(len(after["frames"]), len(before["frames"]))
        for previous, current in zip(before["frames"], after["frames"]):
            if previous["project_id"] != 7:
                self.assertEqual(current, previous)
                continue
            self.assertEqual(current["animation_prompt"], expected[current["uuid"]])
            self.assertEqual(current["status"], FrameStatus.animation_prompt_ready)
            for field in OUTPUT_FIELDS:
                if field != "animation_prompt":
                    self.assertEqual(current[field], previous[field], field)
        self.assertEqual(after["project"]["status"], before["project"]["status"])
        async with self.business_sessions() as reader:
            versions = list((await reader.scalars(select(PromptVersion).where(
                PromptVersion.project_id == 7, PromptVersion.kind == "video",
                PromptVersion.is_active.is_(True)))).all())
            self.assertEqual({version.frame_id: version.text for version in versions},
                             expected_versions, "the original writer must create prompt versions")
        await self.observer.consume(EXECUTION, "project", self.business_sessions)
        facts = await self.facts()
        self.one(facts, "apply_input")
        self.one(facts, "data_applied")
        await self.unchanged_authority()

    async def test_malformed_native_save_preserves_original_error_when_parse_callback_faults(self):
        from app.orchestrator.steps import make_animation_prompts as step
        from app.services import animation_prompt_gpt as native

        module = import_module("app.services.native_animation_contracts")
        self.api("record_animation_parse")
        reply = '{"ops": }'
        diagnostic_error = LookupError("invalid-reply parse journal unavailable")
        parse_spy = AsyncMock(side_effect=diagnostic_error)
        apply_spy = AsyncMock(wraps=step.db_apply.apply_ops)
        before = await self.state()
        async with self.business_sessions() as business:
            project, frames = await self.inputs(business)
            batch = native.collect_batch_items(project, frames)
            self.assertTrue(native.is_apply_ops_blob(reply))
            self.assertEqual(native.parse_apply_ops_animation_reply(reply, frames), [])
            with self.assertRaises(RuntimeError) as original:
                await step._save_anim_pr_batch(
                    business, project, frames, batch, reply, None, shot=1)
            with self.context(business, completed=REQUEST_A), self.business_sql() as statements, \
                    patch.object(module, "record_animation_parse", new=parse_spy), \
                    patch.object(step.db_apply, "apply_ops", new=apply_spy):
                with self.assertRaises(RuntimeError) as caught:
                    await step._save_anim_pr_batch(
                        business, project, frames, batch, reply, None, shot=1)
            self.assertIs(type(caught.exception), type(original.exception))
            self.assertEqual(str(caught.exception), str(original.exception))
            self.assertIsNot(caught.exception, diagnostic_error)
            self.assertTrue(business.is_active)
            self.no_business_writes(statements)
        parse_spy.assert_awaited_once_with(
            reply, [], [], shot=1, target_uuids=[frame.uuid for frame in frames])
        apply_spy.assert_not_awaited()
        self.assertEqual(self.observer.failures.get(EXECUTION), "LookupError")
        self.assertEqual(await self.state(), before)
        self.assertFalse(any(fact.kind in {"apply_input", "data_applied"}
                             for fact in await self.facts()))
        await self.unchanged_authority()

    async def test_sidecar_inherits_real_images_attempt_without_animation_node_facts(self):
        from datetime import datetime
        from sqlalchemy import select
        from app.models import Frame, NodeRun, NodeRunStatus, Project, ProjectStatus, Workflow, WorkflowRun
        from app.services import observation
        from app.services.actions_v1 import ActionExecutionControl, ActionRecord, get_permission
        from app.services.node_capabilities import capability_for
        from app.services.node_status_machine import queue_node_for_start, start_node_running

        record_input = self.api("record_animation_input")
        record_parse = self.api("record_animation_parse")
        observe_result = self.api("observe_animation_result")
        image_nodes = [{"id": "image-node", "type": "images", "data": {"groupId": "images"}}]
        async with self.sessions() as master:
            image_node = NodeRun(workflow_run_id=21, node_key="image-node", node_type="images",
                                 status=NodeRunStatus.pending)
            image_scope = observation.ExecutionControl(
                execution_id=IMAGE_EXECUTION, project_id=8, workflow_run_id=21,
                node_keys=["image-node"], graph_version=revisions.graph_version(image_nodes, []), mode="observe",
                observer_state="active", created_at=datetime.utcnow(), cursors={}, gaps=[])
            master.add_all([
                Project(id=8, slug="foreign-animation-fixture", topic="synthetic images",
                        status=ProjectStatus.generating_images,
                        meta={"canvas_graph": {"nodes": image_nodes, "edges": []}}),
                Workflow(id=19, name="native images sidecar fixture", nodes=image_nodes, edges=[]),
                WorkflowRun(id=21, workflow_id=19, project_id=8,
                            nodes_snapshot=image_nodes, edges_snapshot=[]),
                image_node, image_scope,
            ])
            self.assertTrue(queue_node_for_start(image_node, project_id=8, initiator="worker"))
            self.assertTrue(start_node_running(image_node, project_id=8, initiator="worker"))
            await master.commit()
            run = await master.get(WorkflowRun, 21)
            self.assertEqual(run.nodes_snapshot[0]["type"], image_node.node_type)
            self.assertEqual(image_node.status, NodeRunStatus.running)
        self.observer.scopes[21] = image_scope
        async with self.business_sessions() as business:
            project = await business.get(Project, 8)
            project.status = ProjectStatus.generating_images
            project.meta = {"canvas_graph": {"nodes": image_nodes, "edges": []}}
            await business.commit()
            frames = list((await business.scalars(select(Frame).where(
                Frame.project_id == 8).order_by(Frame.number))).all())
            reply, pairs, ops = self.parsed(project, frames)
            self.assertTrue(pairs and ops)
            before = await self.state()
            bodies = sorted(self.observer.body_root.glob("**/*.json"))
            async with observation.observe_attempt(business, project):
                image_ctx = observation._attempt.get()
                self.assertIsNotNone(image_ctx)
                self.assertEqual((image_ctx.scope.workflow_run_id, image_ctx.scope.project_id,
                                  image_ctx.node_key), (21, 8, "image-node"))

                async def inherited_sidecar():
                    self.assertIs(observation._attempt.get(), image_ctx)
                    with self.business_sql() as statements:
                        await record_input(project, frames, finalize_status=False)
                        await record_parse(reply, pairs, ops, shot=1,
                                           target_uuids=[frame.uuid for frame in frames])
                        await observe_result(business, project)
                    self.no_business_writes(statements)

                await asyncio.create_task(inherited_sidecar())
            self.assertIsNone(observation._attempt.get())
            self.assertEqual(project.status, ProjectStatus.generating_images)
        facts = await self.facts(IMAGE_EXECUTION)
        self.assertTrue(any(fact.kind == "node_start" for fact in facts))
        self.assertTrue(any(fact.kind == "node_end" for fact in facts))
        animation_facts = [fact for fact in facts if (
            fact.payload.get("contract") == "native_animation_prompts"
            or fact.payload.get("node_type") == "animation_prompts"
            or fact.payload.get("check") == "native_animation_prompt_batch"
            or fact.kind in {"parse_result", "node_output", "node_result_verification"})
            and fact.payload.get("contract") != "native_media"]
        self.assertEqual(animation_facts, [],
                         "sidecar capture/result must not be attributed to an images NodeRun")
        self.assertEqual(getattr(image_ctx, "native_node_type", None), "images")
        self.assertFalse(hasattr(image_ctx, "animation_prompt_input"))
        self.assertFalse(getattr(image_ctx, "animation_verification_recorded", False))
        import json as _json
        new_bodies = set(self.observer.body_root.glob("**/*.json")) - set(bodies)
        for body in new_bodies:
            self.assertEqual(_json.loads(body.read_text(encoding="utf-8")).get("node_type"), "images", body)
        self.assertEqual(await self.state(), before)
        self.assertNotIn("retry_profile", capability_for("animation_prompts").actions)
        self.assertIsNone(get_permission(7))
        self.assertIsNone(get_permission(8))
        async with self.sessions() as reader:
            for execution_id in (EXECUTION, IMAGE_EXECUTION):
                self.assertEqual((await reader.get(observation.ExecutionControl, execution_id)).mode,
                                 "observe")
                self.assertIsNone(await reader.get(ActionExecutionControl, execution_id))
            self.assertEqual(list((await reader.scalars(select(ActionRecord))).all()), [])
            nodes = list((await reader.scalars(select(NodeRun))).all())
            self.assertEqual({(node.workflow_run_id, node.node_key): node.status for node in nodes},
                             {(11, "n1"): NodeRunStatus.running, (11, "n2"): NodeRunStatus.pending,
                              (21, "image-node"): NodeRunStatus.running})
            self.assertEqual((await reader.get(WorkflowRun, 11)).nodes_snapshot, self.nodes)
            self.assertEqual((await reader.get(WorkflowRun, 21)).nodes_snapshot, image_nodes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
