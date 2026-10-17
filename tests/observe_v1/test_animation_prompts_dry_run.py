import hashlib
import json
import unittest
from contextlib import ExitStack, asynccontextmanager
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import test_actions_v1_contract as base_tests
from app.models import (
    Base, Frame, NodeRun, NodeRunStatus, Project, ProjectStatus, Workflow, WorkflowRun,
)
from app.orchestrator.steps import make_animation_prompts
from app.services import decision_config, observation
from app.services.actions_recovery import RecoveryAction, finish_job_tracking, reconcile
from app.services.diagnosis import Incident, get_diagnoser
from app.services.gpt_api import GptApiError, GptChatResult
from app.services.node_status_machine import (
    complete_node, fail_node, queue_node_for_start, start_node_running,
)
from app.services.observation import observe_model
from app.settings import settings

EXECUTION = "00000000-0000-4000-8000-0000000000b1"
SCRIPT = {"replies": [], "bodies": [], "attachments": []}
UUID1, UUID2 = "af" * 16, "bf" * 16
ANIM = "The camera slowly tracks the person opening the blue book, a single continuous take."
ANIM2 = "A second camera angle shows the same person turning a page, a single continuous take."
OUTAGE = lambda: GptApiError("GPT: server exception, try again later", context={"retryable": True})


def ops_reply(*uuids, field="промт_видео", text=ANIM):
    return json.dumps({"ops": [{"frame_uuid": uid, "fields": {field: text}} for uid in uuids]},
                      ensure_ascii=False)


def write_png(path, color):
    from PIL import Image
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (64, 96), color).save(path)


@observe_model
async def fake_http_post(*, url, headers, body, timeout, use_model):
    SCRIPT["bodies"].append(body)
    if not SCRIPT["replies"]:
        raise AssertionError("unscripted model request: the dry run allows exactly the scripted ones")
    reply = SCRIPT["replies"].pop(0)
    if isinstance(reply, BaseException):
        raise reply
    return GptChatResult(text=reply, model=use_model, finish_reason="stop",
                         usage={"prompt_tokens": 10, "completion_tokens": 5},
                         raw={"status": "completed"}, response_id=f"resp-{len(SCRIPT['bodies'])}")


async def fake_ask_with_files(self, text, files, *, timeout=600, project_id=None, expect_file_download=False,
                              history=None, treat_txt_as_prompt=True, system=None, max_retries=None,
                              auto_pack=True, on_delta=None):
    master, data = "", []
    for item in files:
        path = Path(item)
        raw = path.read_bytes()
        if treat_txt_as_prompt and path.suffix.lower() in {".txt", ".md"} and not master:
            master = raw.decode("utf-8", errors="replace")
        else:
            data.append({"name": path.name, "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw),
                         "magic": raw[:4].hex()})
    SCRIPT["attachments"].append(data)
    content = "\n\n".join(part for part in (master, text) if part)
    body = {"model": "gpt-5-6-sol", "system": system or "",
            "input": [{"role": "user", "content": content}], "attachments": data}
    result = await fake_http_post(url="https://synthetic.invalid/codex/v1/responses", headers={}, body=body,
                                  timeout=timeout, use_model="gpt-5-6-sol")
    return result.text


class DryRun(unittest.IsolatedAsyncioTestCase):
    fixture_node_type = "animation_prompts"

    async def asyncSetUp(self):
        SCRIPT["replies"], SCRIPT["bodies"], SCRIPT["attachments"] = [], [], []
        await base_tests.DurableActions.asyncSetUp(self)
        self.stack = ExitStack()
        self.data = Path(self.tmp.name) / "videos" / "actions-fixture"
        self.data.mkdir(parents=True)
        from app.services.gpt_client import ApiGptClient
        for patcher in (patch.object(settings, "data_dir", Path(self.tmp.name)),
                        patch.object(settings, "gpt_api_key", "synthetic-observe-test-key"),
                        patch.object(settings, "gpt_base_url", "http://127.0.0.1:18766"),
                        patch.object(ApiGptClient, "ask_with_files", new=fake_ask_with_files),
                        patch.object(make_animation_prompts, "_BATCH_BACKOFF_S", (0, 0, 0, 0, 0))):
            self.stack.enter_context(patcher)
        cfg = decision_config.validate_config({'version': 1, 'mode': 'synthetic', 'online_enabled': True,
            'endpoint': 'http://127.0.0.1:18766/codex/v1/responses',
            'decision': {'catalog_id': 'gpt-kie', 'model_id': 'gpt-5-6-sol', 'reasoning_effort': 'low',
                         'max_output_tokens': 1024, 'timeout_s': 30},
            'profiles': {}, 'allowed_actions': sorted(decision_config.ALL_ACTIONS)})
        self.stack.enter_context(patch.object(decision_config, "runtime_config", return_value=cfg))
        self.observer = observation._observer

        write_png(self.data / "scenes" / "frame_001_a1b2c3d4.png", (200, 30, 30))
        write_png(self.data / "scenes" / "frame_002_e5f6a7b8.png", (30, 30, 200))

        self.engine = create_async_engine(f"sqlite+aiosqlite:///{self.data / 'project.db'}")
        self.business = async_sessionmaker(self.engine, expire_on_commit=False)
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
            await observation.migrate_observation_schema(connection)
        canvas = {"canvas_graph": {"nodes": self.nodes, "edges": []}}
        async with self.business() as seed:
            seed.add_all([
                Project(id=7, slug="actions-fixture", topic="synthetic animation",
                        status=ProjectStatus.generating_animation_prompts, auto_mode=False, meta=dict(canvas)),
                Workflow(id=9, name="actions fixture", nodes=self.nodes, edges=[]),
                WorkflowRun(id=11, workflow_id=9, project_id=7, nodes_snapshot=self.nodes, edges_snapshot=[]),
                NodeRun(workflow_run_id=11, node_key="n1", node_type="animation_prompts",
                        status=NodeRunStatus.pending),
                NodeRun(workflow_run_id=11, node_key="n2", node_type="script", status=NodeRunStatus.pending),
                Frame(project_id=7, number=1, uuid=UUID1, sort_key=1, voiceover_text="Альфа открывает книгу.",
                      image_prompt="A person opens a blue book at a desk.", duration_seconds=3.0, attrs={}),
                Frame(project_id=7, number=2, uuid=UUID2, sort_key=2, voiceover_text="Бета закрывает книгу.",
                      image_prompt="A person closes a blue book at a desk.", duration_seconds=3.0, attrs={}),
            ])
            await seed.commit()
        async with self.sessions() as master:
            project = await master.get(Project, 7)
            project.status, project.meta = ProjectStatus.generating_animation_prompts, dict(canvas)
            scope = await master.get(observation.ExecutionControl, "exec-durable")
            scope.execution_id, scope.node_keys = EXECUTION, ["n1", "n2"]
            await master.commit()

        @asynccontextmanager
        async def project_scope(_project_id):
            async with self.business() as session:
                yield session
                await session.commit()

        async def push(session, project):
            await self.copy_runtime(session, project)

        async def no_pull(*_args):
            return None

        self.stack.enter_context(patch("app.project_db.project_db_session_scope", project_scope))
        self.stack.enter_context(patch("app.project_db.push_runtime_to_master", push))
        self.stack.enter_context(patch("app.project_db.pull_master_runtime_into_project", no_pull))
        await self.supervisor.enable(project_id=7, execution_id=EXECUTION, group_id="g1", node_keys=["n1"],
            graph_version=self.graph, initiator="dry-run", rationale="animation_prompts dry run", request_limit=30)
        await self.observer.reload()

    async def asyncTearDown(self):
        try:
            await self.engine.dispose()
            self.stack.close()
        finally:
            await base_tests.DurableActions.asyncTearDown(self)

    async def copy_runtime(self, session, project):
        from app.project_db import _apply_runtime_fields, _copy_noderuns_between_sessions
        async with self.sessions() as master:
            _apply_runtime_fields(project, await master.get(Project, 7))
            await _copy_noderuns_between_sessions(session, master, 7)
            await master.commit()

    async def sync(self):
        async with self.business() as session:
            await self.copy_runtime(session, await session.get(Project, 7))

    async def deliver(self):
        await self.observer.consume(EXECUTION, "master", self.sessions)
        await self.observer.consume(EXECUTION, "project", self.business)

    async def diagnose(self):
        await self.sync()
        await self.deliver()
        await get_diagnoser(self.observer).update(EXECUTION)
        self.assertEqual(get_diagnoser(self.observer).failures, {})

    async def facts(self, kind=None):
        await self.deliver()
        async with self.sessions() as session:
            rows = (await session.scalars(select(observation.ObservationEvent).where(
                observation.ObservationEvent.execution_id == EXECUTION)
                .order_by(observation.ObservationEvent.id))).all()
        return [r for r in rows if kind is None or r.kind == kind]

    async def incidents(self):
        async with self.sessions() as session:
            return (await session.scalars(select(Incident).where(Incident.execution_id == EXECUTION))).all()

    async def frames(self):
        async with self.business() as session:
            return list((await session.scalars(select(Frame).where(Frame.project_id == 7)
                                               .order_by(Frame.number))).all())

    async def worker(self, replies, *, start=True):
        SCRIPT["replies"] = list(replies)
        async with self.business() as session:
            project = await session.get(Project, 7)
            node = await session.scalar(select(NodeRun).where(NodeRun.node_key == "n1"))
            if start:
                self.assertTrue(queue_node_for_start(node, project_id=7, initiator="worker"))
                self.assertTrue(start_node_running(node, project_id=7, initiator="worker"))
                await session.commit()
            try:
                async with observation.observe_attempt(session, project):
                    await make_animation_prompts.run(session, project, None)
            except Exception:
                await session.rollback()
                raise
            complete_node(node, project_id=7, initiator="worker")
            await session.commit()

    async def exhaust_native_retries(self, error="native step retries exhausted"):
        async with self.business() as session:
            project = await session.get(Project, 7)
            node = await session.scalar(select(NodeRun).where(NodeRun.node_key == "n1"))
            fail_node(node, error, project_id=7, initiator="worker")
            node.finished_at, node.attempts = datetime.utcnow(), 9
            project.status = ProjectStatus.paused
            project.meta = {**project.meta, "step_failure": {"total_fails": {"generating_animation_prompts": 9},
                "last_running": "generating_animation_prompts", "abandoned_at": "synthetic"}}
            await session.commit()
        await self.diagnose()

    async def failed_attempt(self, replies):
        with self.assertRaises(RuntimeError) as caught:
            await self.worker(replies)
        await self.exhaust_native_retries()
        return caught.exception

    async def typed_incident(self):
        found = [i for i in await self.incidents() if i.type in {"invalid_json", "missing_payload"}]
        self.assertEqual(len(found), 1, [(i.type, i.status) for i in await self.incidents()])
        return found[0]

    def decision_model(self, actions):
        queue = list(actions)

        async def reply(**kwargs):
            context = json.loads(kwargs["prompt"])
            action = queue.pop(0)
            self.assertIn(action, context["allowed_actions"])
            proposal = dict(action=action, arguments={}, evidence_ids=context["evidence_ids"][:1],
                            context_version=context["context_version"],
                            explanation="synthetic facts, no hidden reasoning")
            return GptChatResult(text=json.dumps(proposal), model="synthetic", raw={"status": "completed"})
        return patch("app.services.gpt_api.chat", side_effect=reply)

    def declare_native_second_shot(self):
        from openpyxl import Workbook
        from app.services.xlsx_v8_import import ROW_IMAGE_PROMPT_V8, ROW_VOICEOVER_V8
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "план"
        for number, voice, prompt in ((1, "Альфа открывает книгу.", "A person opens a blue book at a desk."),
                                      (2, "Бета закрывает книгу.", "A person closes a blue book at a desk.")):
            sheet.cell(ROW_VOICEOVER_V8, number + 2, voice)
            sheet.cell(ROW_IMAGE_PROMPT_V8, number + 2, prompt)
        sheet.cell(46, 1 + 2, "The same person seen from a second camera angle.")
        workbook.save(self.data / "project.xlsx")
        workbook.close()
        write_png(self.data / "scenes" / "frame_001_s2_native.png", (30, 200, 30))

    async def test_batch_request_carries_the_png_strip_and_a_bad_reply_is_typed_invalid_json(self):
        await self.failed_attempt(["phase one reply", "no pairs here at all"])
        self.assertEqual(len(SCRIPT["attachments"]), 2)
        self.assertEqual(SCRIPT["attachments"][0], [], "phase 1: the master prompt .md is inlined, no images")
        strip = SCRIPT["attachments"][1]
        self.assertEqual(len(strip), 1, "the batch carries exactly one strip of both frame images")
        self.assertRegex(strip[0]["name"], r"^anim_pr_strip_001_002\.(png|jpg)$")
        self.assertGreater(strip[0]["bytes"], 0)
        self.assertIn(strip[0]["magic"][:6], {"89504e", "ffd8ff"})
        parse = (await self.facts("parse_result"))[-1].payload
        self.assertEqual(parse["contract"], "native_animation_prompts")
        self.assertIs(parse["payload_required"], True)
        self.assertIs(parse["payload_present"], False)
        self.assertIs(parse["syntax_valid"], False)
        incident = await self.typed_incident()
        self.assertEqual(incident.type, "invalid_json")
        self.assertFalse(incident.details["evidence_incomplete"])
        self.assertEqual(await self.facts("data_applied"), [])
        self.assertFalse(any(i.type == "animation_degraded" for i in await self.incidents()))

    async def test_json_without_pairs_is_typed_missing_payload(self):
        await self.failed_attempt(["phase one reply", '{"note": "nothing usable"}'])
        self.assertIs((await self.facts("parse_result"))[-1].payload["syntax_valid"], True)
        self.assertEqual((await self.typed_incident()).type, "missing_payload")

    async def test_success_applies_both_prompts_and_the_verifier_confirms_them(self):
        await self.worker(["phase one reply", ops_reply(UUID1, UUID2)])
        self.assertEqual([len(f.animation_prompt or "") > 10 for f in await self.frames()], [True, True])
        proof = (await self.facts("node_result_verification"))[-1].payload
        self.assertEqual((proof["status"], proof["target_frame_count"]), ("passed", 2))
        self.assertIs(proof["voiceover_preserved"] and proof["image_prompts_preserved"], True)
        self.assertGreaterEqual(len(await self.facts("data_applied")), 1)
        await self.diagnose()
        self.assertEqual(await self.incidents(), [], "a clean run raises no incident")

    async def test_local_fallback_after_a_provider_outage_is_a_visible_degradation(self):
        await self.worker(["phase one reply", OUTAGE(), OUTAGE(), OUTAGE()])
        self.assertEqual([len(f.animation_prompt or "") > 10 for f in await self.frames()], [True, True])
        self.assertIn("NO VOICE", (await self.frames())[0].animation_prompt)
        degraded = await self.facts("animation_degraded")
        self.assertEqual(len(degraded), 1)
        fact = degraded[0].payload
        self.assertEqual((fact["reason"], fact["shot"], fact["frame_numbers"]), ("local_fallback", 1, [1, 2]))
        self.assertEqual(fact["cause_class"], "GptApiError")
        self.assertIs(fact["semantic_quality_verified"], False)
        await self.diagnose()
        found = [i for i in await self.incidents() if i.type == "animation_degraded"]
        self.assertEqual(len(found), 1)
        item = found[0]
        self.assertEqual((item.status, item.severity, item.details["blocking"]), ("human_review", "warning", False))
        self.assertIn("локальным шаблоном", item.details["fact"])
        self.assertIn("1, 2", item.details["fact"])
        self.assertFalse(any(i.type in {"invalid_json", "missing_payload"} for i in await self.incidents()))
        async with self.business() as session:
            self.assertEqual((await session.get(Project, 7)).status, ProjectStatus.animation_prompts_ready)

    async def test_skipped_second_shot_is_a_visible_degradation_and_the_verifier_still_fails_it(self):
        self.declare_native_second_shot()
        await self.worker(["phase one reply", ops_reply(UUID1, UUID2), OUTAGE(), OUTAGE(), OUTAGE()])
        degraded = await self.facts("animation_degraded")
        self.assertEqual(len(degraded), 1)
        self.assertEqual((degraded[0].payload["reason"], degraded[0].payload["shot"],
                          degraded[0].payload["frame_numbers"]), ("shot2_skipped", 2, [1]))
        proof = (await self.facts("node_result_verification"))[-1].payload
        self.assertEqual(proof["status"], "failed")
        self.assertIn("missing_native_shot2_animation:" + UUID1, proof["errors"])
        await self.diagnose()
        found = [i for i in await self.incidents() if i.type == "animation_degraded"]
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0].details["degradation"], "shot2_skipped")
        self.assertIn("shot_02", found[0].details["fact"])
        self.assertFalse(found[0].details["blocking"])
        self.assertEqual(len(SCRIPT["attachments"][2]), 1, "the shot_02 batch carries its own strip too")
        self.assertRegex(SCRIPT["attachments"][2][0]["name"], r"^anim_pr_strip_001\.(png|jpg)$")

    async def test_a_degradation_incident_never_opens_a_recovery(self):
        await self.worker(["phase one reply", OUTAGE(), OUTAGE(), OUTAGE()])
        await self.diagnose()
        incident = next(i for i in await self.incidents() if i.type == "animation_degraded")
        with self.assertRaisesRegex(ValueError, "current mandatory JSON/payload incident required"):
            await self.supervisor.recover_node(project_id=7, execution_id=EXECUTION,
                incident_id=incident.incident_id, node_key="n1", idempotency_key="degraded-retry-0001",
                initiator="dry-run", rationale="a degradation is not a failure")
        async with self.sessions() as session:
            self.assertIsNone(await session.scalar(select(RecoveryAction)))

    async def test_failed_attempt_to_decision_to_one_verified_retry_with_images(self):
        from app.services import decision_v1
        await self.failed_attempt(["phase one reply", "no pairs here at all"])
        first_bodies = list(SCRIPT["bodies"])
        self.assertEqual(len(first_bodies), 2)
        incident = await self.typed_incident()
        self.assertEqual(incident.type, "invalid_json")
        loop = decision_v1.DecisionSupervisor(self.observer)
        await loop.configure(project_id=7, execution_id=EXECUTION, enabled=True,
                             initiator="dry-run", rationale="animation_prompts dry run decision")
        record = await loop.begin(project_id=7, execution_id=EXECUTION, incident_id=incident.incident_id)
        context = await loop.context(record["decision_id"])
        self.assertEqual(context["node_capability"]["node_type"], "animation_prompts")
        self.assertIn("retry_same", context["allowed_actions"])
        self.assertNotIn("retry_profile", context["allowed_actions"])
        contract = await loop.read_contract(record["decision_id"])
        self.assertEqual(contract["contract"]["node_type"], "animation_prompts")
        self.assertIs(contract["contract"]["payload_required"], True)
        self.assertEqual(contract["contract"]["degradation_facts"]["fact_kind"], "animation_degraded")
        with self.decision_model(["read_contract", "retry_same"]):
            await loop.step(record["decision_id"])
            await loop.step(record["decision_id"])
        async with self.sessions() as session:
            ticket = await session.scalar(select(RecoveryAction))
            self.assertEqual(ticket.status, "queued")
            self.assertEqual(ticket.preconditions["action_type"], "retry_same")
            self.assertEqual(ticket.preconditions["effective_node_type"], "animation_prompts")
            self.assertEqual(ticket.preconditions["first_body_normalization"], "attachment_timestamps")
            self.assertEqual(ticket.budget_state["recovery_limit"], 1)
        before = len(SCRIPT["bodies"])
        await self.worker(["phase one reply", ops_reply(UUID1, UUID2)], start=False)
        finish_job_tracking(7)
        retry_bodies = SCRIPT["bodies"][before:]
        self.assertEqual(len(retry_bodies), 2)
        from app.services.actions_recovery import body_identity
        self.assertEqual(body_identity(retry_bodies[0], normalize_names=True),
                         body_identity(first_bodies[0], normalize_names=True),
                         "recovery's first request equals the source's, modulo the attachment-name stamp")
        self.assertEqual(len(SCRIPT["attachments"][-1]), 1, "the retry's batch carries the strip again")
        await self.sync()
        await self.deliver()
        await reconcile(self.observer)
        async with self.sessions() as session:
            ticket = await session.scalar(select(RecoveryAction))
            self.assertEqual(ticket.status, "verified", ticket.verification_result)
            checks = ticket.verification_result
            self.assertTrue(checks["data_applied"] and checks["accepted_payload"] and checks["result_confirmed"])
            self.assertTrue(checks["node_done"] and checks["transactional_done"] and checks["requests_complete"])
            self.assertEqual(ticket.result["continuation"], "original runtime; mandatory HITL untouched")
            self.assertFalse((await session.get(base_tests.ActionExecutionControl, EXECUTION)).held)
        self.assertEqual([len(f.animation_prompt or "") > 10 for f in await self.frames()], [True, True])
        async with self.business() as session:
            self.assertEqual((await session.get(Project, 7)).status, ProjectStatus.animation_prompts_ready)
        with self.assertRaisesRegex(ValueError, "decision owner"):
            await self.supervisor.recover_node(project_id=7, execution_id=EXECUTION,
                incident_id=incident.incident_id, node_key="n1", idempotency_key="second-retry-0001",
                initiator="dry-run", rationale="second retry must be refused")
        from sqlalchemy.exc import IntegrityError
        async with self.sessions() as session:
            session.add(RecoveryAction(action_id="second", execution_id=EXECUTION, project_id=7,
                incident_id=incident.incident_id, node_key="n1", source_attempt_id="s", new_attempt_id="n",
                idempotency_key="second-ticket-0001", preconditions={}, budget_state={}, status="reserved"))
            with self.assertRaises(IntegrityError):
                await session.commit()
        self.assertEqual(len(SCRIPT["bodies"]), 4, "two failed requests plus two recovery requests, nothing else")

    async def test_retry_is_refused_after_a_batch_was_applied_partial_apply(self):
        await self.failed_attempt(["phase one reply", ops_reply(UUID1), "no pairs here at all"])
        frames = await self.frames()
        self.assertGreater(len((frames[0].animation_prompt or "")), 10, "frame 1 was applied")
        self.assertFalse((frames[1].animation_prompt or "").strip())
        self.assertGreaterEqual(len(await self.facts("data_applied")), 1)
        incident = await self.typed_incident()
        with self.assertRaisesRegex(ValueError, "partial apply"):
            await self.supervisor.recover_node(project_id=7, execution_id=EXECUTION,
                incident_id=incident.incident_id, node_key="n1", idempotency_key="blocked-retry-0001",
                initiator="dry-run", rationale="applied batch must block the retry")
        async with self.sessions() as session:
            self.assertIsNone(await session.scalar(select(RecoveryAction)))

    async def test_profile_retry_is_denied_by_capability(self):
        await self.failed_attempt(["phase one reply", "no pairs here at all"])
        incident = await self.typed_incident()
        with self.assertRaisesRegex(ValueError, "denied by node capability"):
            await self.supervisor.recover_node(project_id=7, execution_id=EXECUTION,
                incident_id=incident.incident_id, node_key="n1", idempotency_key="profile-retry-0001",
                initiator="dry-run", rationale="profile retry is not an animation_prompts capability",
                profile_overlay_id="no-such-overlay")


if __name__ == "__main__":
    unittest.main(verbosity=2)
