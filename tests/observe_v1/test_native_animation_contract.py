
from __future__ import annotations

from base64 import b64decode
from copy import deepcopy
from importlib import import_module
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest


PROJECT_ID = 7
ANIMATION_TEXT = "The camera slowly tracks the person opening the blue book."
SECOND_ANIMATION = "The camera orbits the same person holding the opened book."
PNG = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
    "/x8AAwMCAO+aL1sAAAAASUVORK5CYII="
)


def contracts():
    try:
        return import_module("app.services.native_animation_contracts")
    except ModuleNotFoundError as exc:
        if exc.name == "app.services.native_animation_contracts":
            raise AssertionError("missing bounded native animation contracts") from exc
        raise


def native():
    return import_module("app.services.animation_prompt_gpt")


def frames_fixture():
    return [
        SimpleNamespace(
            id=number, project_id=PROJECT_ID, uuid=f"{number:024x}",
            number=number, sort_key=sort_key, scene_id=31,
            voiceover_text=voiceover, image_prompt=image,
            animation_prompt=None,
            attrs={"existing_note": {"keep": ["original"]}},
        )
        for number, sort_key, voiceover, image in (
            (1, 2.0, "Альфа открывает книгу.\n", "A person opens a blue book."),
            (2, 1.5, "Бета закрывает книгу.  ", "A person closes a blue book."),
        )
    ]


def assert_unclaimed(test, proof):
    for key in ("business_commit_claim", "semantic_quality_verified",
                "video_quality_verified", "managed_eligible"):
        test.assertIs(proof[key], False, proof)


class NativeAnimationContracts(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.project = SimpleNamespace(id=PROJECT_ID, data_dir=self.root)
        self.frames = frames_fixture()

    def snapshot(self, frames=None, *, finalize_status=True):
        return contracts().animation_prompt_input_snapshot(
            self.frames if frames is None else frames,
            project_id=PROJECT_ID, project=self.project,
            finalize_status=finalize_status,
        )

    def filled(self, frames=None):
        result = deepcopy(self.frames if frames is None else frames)
        for frame in result:
            if native().frame_has_usable_image_prompt(frame):
                frame.animation_prompt = ANIMATION_TEXT
        return result

    def check(self, expected, frames, snapshot):
        before = deepcopy((frames, snapshot))
        proof = contracts().verify_animation_prompts(
            frames, project_id=PROJECT_ID, project=self.project, snapshot=snapshot,
        )
        self.assertEqual(proof["status"], expected, proof)
        assert_unclaimed(self, proof)
        self.assertEqual((frames, snapshot), before, "verification is read-only")
        return proof

    def plan(self, *, declared=(), second_animation=None, primary_animation=None):
        from openpyxl import Workbook
        from app.services.xlsx_v8_import import ROW_IMAGE_PROMPT_V8, ROW_VOICEOVER_V8

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "план"
        for frame in self.frames:
            column = frame.number + 2
            sheet.cell(ROW_VOICEOVER_V8, column, frame.voiceover_text)
            sheet.cell(ROW_IMAGE_PROMPT_V8, column, frame.image_prompt)
            if frame.number in declared:
                sheet.cell(46, column, "The same person seen from a second camera angle.")
            if second_animation is not None:
                sheet.cell(64, column, second_animation)
            if primary_animation is not None:
                sheet.cell(48, column, primary_animation)
        workbook.save(self.root / "project.xlsx")
        workbook.close()

    def second_image(self, number=1):
        directory = self.root / "scenes"
        directory.mkdir(exist_ok=True)
        path = directory / f"frame_{number:03d}_s2_native.png"
        path.write_bytes(PNG)
        return path

    def test_contract_names_original_executor_parser_writer_and_harness(self):
        from app.orchestrator.steps import make_animation_prompts
        from app.services import agent_harness, node_write_contract
        from app.services.node_capabilities import READ_ACTIONS

        contract = contracts().animation_prompts_contract()
        self.assertEqual(contract["node_type"], "animation_prompts")
        self.assertEqual(contract["executor"],
                         "app.orchestrator.steps.make_animation_prompts.run")
        self.assertEqual(contract["parser"],
                         "app.services.animation_prompt_gpt.parse_animation_reply")
        self.assertEqual(contract["writer"], "app.services.db_apply.apply_ops")
        self.assertEqual(contract["harness"], "app.services.agent_harness.harness_gate_or_raise")
        self.assertEqual(contract["write_contract"],
                         "node_write_contract.filter_ops_for_node(node_kind=anim_pr)")
        self.assertEqual(set(contract["allowed_fields"]), set(node_write_contract.ANIM_PROMPT_FIELDS))
        self.assertEqual(contract["min_characters"], native().MIN_ANIM_PROMPT_LEN)
        self.assertEqual(contract["shot2_min_characters"], native().MIN_SHOT2_VIDEO_PROMPT_LEN)
        self.assertFalse(contract["primary_png_required"])
        self.assertFalse(contract["child_exemption"])
        self.assertTrue(contract["managed"])
        self.assertEqual(tuple(contract["actions"]), READ_ACTIONS + ("retry_same", "finish"))
        self.assertTrue(contract["payload_required"])
        assert_unclaimed(self, contract)
        self.assertTrue(callable(make_animation_prompts.run))
        self.assertTrue(callable(make_animation_prompts._save_anim_pr_batch))
        self.assertTrue(callable(agent_harness.harness_gate_or_raise))

    def test_snapshot_detaches_raw_inputs_in_native_number_order(self):
        before = deepcopy(self.frames)
        snapshot = self.snapshot(list(reversed(self.frames)))
        self.assertEqual(snapshot["project_id"], PROJECT_ID)
        self.assertEqual([row["number"] for row in snapshot["frames"]], [1, 2])
        self.assertEqual([row["sort_key"] for row in snapshot["frames"]], [2.0, 1.5])
        for row, frame in zip(snapshot["frames"], self.frames):
            for field in ("id", "project_id", "uuid", "number", "sort_key", "scene_id",
                          "voiceover_text", "image_prompt", "attrs"):
                self.assertEqual(row[field], getattr(frame, field))
            self.assertIs(row["required_prompt"], True)
            self.assertIs(row["required_shot2"], False)
        snapshot["frames"][0]["attrs"]["existing_note"]["keep"].append("detached")
        self.assertEqual(self.frames, before)

    def test_positive_primary_prompts_need_no_image_files_and_make_no_quality_claim(self):
        snapshot = self.snapshot()
        self.assertFalse((self.root / "scenes").exists())
        proof = self.check("passed", self.filled(), snapshot)
        self.assertTrue(proof["voiceover_preserved"])
        self.assertTrue(proof["image_prompts_preserved"])
        self.assertEqual(proof["target_frame_count"], 2)

    def test_missing_or_incomplete_input_provenance_is_unknown(self):
        good = self.snapshot()
        incomplete = deepcopy(good)
        del incomplete["frames"][0]["image_prompt"]
        no_order = deepcopy(good)
        del no_order["frames"][0]["sort_key"]
        no_target = deepcopy(good)
        del no_target["frames"][0]["required_prompt"]
        for snapshot in (None, {}, {"frames": []}, incomplete, no_order, no_target):
            with self.subTest(snapshot=snapshot):
                self.check("unknown", self.filled(), snapshot)

    def test_identity_order_voiceover_and_image_prompt_drift_fail(self):
        snapshot = self.snapshot()
        for field, value in (
            ("id", 99), ("project_id", 8), ("uuid", "f" * 24),
            ("number", 9), ("sort_key", 9.5), ("scene_id", 99),
            ("voiceover_text", "Rewritten voiceover."),
            ("voiceover_text", self.frames[0].voiceover_text.strip()),
            ("image_prompt", "Rewritten image prompt."),
        ):
            changed = self.filled()
            setattr(changed[0], field, value)
            with self.subTest(field=field, value=value):
                self.check("failed", changed, snapshot)

    def test_added_removed_duplicate_and_blank_frame_identities_fail(self):
        snapshot = self.snapshot()
        extra = deepcopy(self.frames[0])
        extra.id, extra.number, extra.uuid = 3, 3, f"{3:024x}"
        extra.animation_prompt = ANIMATION_TEXT
        duplicates = self.filled()
        duplicates[1].uuid = duplicates[0].uuid
        duplicate_number = self.filled()
        duplicate_number[1].number = duplicate_number[0].number
        blank = self.filled()
        blank[0].uuid = None
        for changed in (self.filled()[:1], self.filled() + [extra],
                        duplicates, duplicate_number, blank):
            with self.subTest(identities=[f.uuid for f in changed]):
                self.check("failed", changed, snapshot)

    def test_database_iteration_order_does_not_reassign_prompts(self):
        self.check("passed", list(reversed(self.filled())), self.snapshot())

    def test_actual_rows_must_have_every_required_primary_animation(self):
        snapshot = self.snapshot()
        for text in (None, "", "   ", "short", "x" * (native().MIN_ANIM_PROMPT_LEN - 1),
                     "нет исходных данных для заполнения",
                     '{"ops":[{"frame_uuid":"foreign","fields":{"промт_видео":"text"}}]}'):
            changed = self.filled()
            changed[0].animation_prompt = text
            with self.subTest(text=text):
                self.check("failed", changed, snapshot)

    def test_original_minimum_is_preserved_without_inventing_a_stronger_threshold(self):
        changed = self.filled()
        changed[0].animation_prompt = "x" * native().MIN_ANIM_PROMPT_LEN
        self.assertTrue(native().has_animation_prompt_for_frame(self.project, changed[0]))
        self.check("passed", changed, self.snapshot())

    def test_raw_parser_reply_is_insufficient_without_applied_result_rows(self):
        reply = json.dumps({"ops": [
            {"frame_uuid": f.uuid, "fields": {"промт_видео": ANIMATION_TEXT}}
            for f in self.frames
        ]}, ensure_ascii=False)
        batch = native().collect_batch_items(self.project, self.frames)
        self.assertEqual(len(native().parse_animation_reply(reply, self.frames, batch_items=batch)), 2)
        self.check("failed", deepcopy(self.frames), self.snapshot())

    def test_second_prompt_does_not_replace_missing_primary_animation(self):
        changed = self.filled()
        changed[0].animation_prompt = None
        changed[0].attrs["animation_prompt_shot2"] = SECOND_ANIMATION
        self.check("failed", changed, self.snapshot())

    def test_xlsx_primary_mirror_does_not_replace_missing_db_animation(self):
        self.plan(primary_animation=ANIMATION_TEXT)
        changed = self.filled()
        changed[0].animation_prompt = None
        self.assertFalse(native().has_animation_prompt_for_frame(self.project, changed[0]))
        self.check("failed", changed, self.snapshot())

    def test_nonusable_primary_inputs_keep_native_no_work_semantics(self):
        for text in (None, "", "нет исходных данных для заполнения", "[ID: P7-F1-abcd]"):
            rows = deepcopy(self.frames)
            rows[0].image_prompt = text
            snapshot = self.snapshot(rows)
            self.assertFalse(snapshot["frames"][0]["required_prompt"])
            changed = self.filled(rows)
            self.assertIsNone(changed[0].animation_prompt)
            self.check("passed", changed, snapshot)

    def test_target_boolean_cannot_hide_a_required_input(self):
        snapshot = self.snapshot()
        snapshot["frames"][0]["required_prompt"] = False
        changed = self.filled()
        changed[0].animation_prompt = None
        self.check("failed", changed, snapshot)

    def test_output_cannot_erase_primary_input_to_remove_coverage_requirement(self):
        changed = self.filled()
        changed[0].image_prompt = ""
        changed[0].animation_prompt = None
        self.check("failed", changed, self.snapshot())

    def test_native_child_with_usable_image_prompt_is_not_exempt(self):
        child = deepcopy(self.frames[0])
        child.id, child.number, child.sort_key, child.uuid = 3, 3, 2.5, f"{3:024x}"
        child.voiceover_text = ""
        child.attrs["camera_subdivide"] = {"role": "shot", "parent_uuid": self.frames[0].uuid}
        rows = self.frames + [child]
        self.assertEqual([it.frame.number for it in native().collect_batch_items(self.project, rows)], [1, 2, 3])
        snapshot = self.snapshot(rows)
        self.assertIs(snapshot["frames"][2]["required_prompt"], True)
        changed = self.filled(rows)
        changed[2].animation_prompt = None
        self.check("failed", changed, snapshot)
        self.check("passed", self.filled(rows), snapshot)

    def test_child_without_usable_primary_input_needs_no_invented_prompt(self):
        child = deepcopy(self.frames[0])
        child.id, child.number, child.sort_key, child.uuid = 3, 3, 2.5, f"{3:024x}"
        child.voiceover_text, child.image_prompt = "", None
        child.attrs["camera_subdivide"] = {"role": "shot", "parent_uuid": self.frames[0].uuid}
        rows = self.frames + [child]
        self.assertFalse(self.snapshot(rows)["frames"][2]["required_prompt"])
        self.check("passed", self.filled(rows), self.snapshot(rows))

    def test_native_harness_child_coverage_is_exposed_without_replacing_the_gate(self):
        from app.services.agent_harness import _append_prompt_nn_checks

        frame = self.frames[0]
        checks, repair = [], []
        _append_prompt_nn_checks(checks, repair, [
            (frame.number, "", frame.image_prompt, "", json.dumps({
                "camera_subdivide": {"role": "shot", "parent_uuid": "parent"}
            }), frame.uuid)
        ], status="generating_animation_prompts", step="anim_pr")
        coverage = next(check for check in checks if check.name == "anim_pr_coverage")
        self.assertFalse(coverage.ok)
        self.assertIn("anim_pr", repair)

    def test_unrelated_native_attrs_and_second_image_prompt_are_preserved(self):
        self.frames[0].attrs["image_prompt_shot2"] = "Native optional second image input."
        snapshot = self.snapshot()
        for field, value in (("image_prompt_shot2", "rewritten"),
                             ("existing_note", {"keep": []}),
                             ("camera_subdivide", {"role": "shot"})):
            changed = self.filled()
            changed[0].attrs[field] = value
            with self.subTest(field=field):
                self.check("failed", changed, snapshot)

    def test_optional_shot02_is_not_required_from_attrs_alone(self):
        self.frames[0].attrs["image_prompt_shot2"] = "An optional second image prompt."
        self.assertFalse(native().frame_needs_shot2_video_prompt(self.project, self.frames[0]))
        self.assertFalse(self.snapshot()["frames"][0]["required_shot2"])
        self.check("passed", self.filled(), self.snapshot())

    def test_optional_shot02_needs_both_native_declaration_and_image(self):
        self.plan(declared=(1,))
        self.assertFalse(native().frame_needs_shot2_video_prompt(self.project, self.frames[0]))
        self.check("passed", self.filled(), self.snapshot())
        self.second_image(2)
        self.assertFalse(native().frame_needs_shot2_video_prompt(self.project, self.frames[1]))
        self.assertFalse(self.snapshot()["frames"][1]["required_shot2"])

    def test_declared_available_shot02_requires_original_completion_predicate(self):
        self.plan(declared=(1,))
        self.second_image()
        snapshot = self.snapshot()
        self.assertTrue(snapshot["frames"][0]["required_shot2"])
        self.check("failed", self.filled(), snapshot)
        changed = self.filled()
        changed[0].attrs["animation_prompt_shot2"] = SECOND_ANIMATION
        self.assertFalse(native().has_animation_prompt_shot2_for_frame(self.project, changed[0]))
        self.check("failed", changed, snapshot)
        self.plan(declared=(1,), second_animation=SECOND_ANIMATION)
        self.assertTrue(native().has_animation_prompt_shot2_for_frame(self.project, changed[0]))
        self.check("passed", changed, snapshot)

    def test_existing_native_shot02_completion_does_not_queue_new_work(self):
        self.plan(declared=(1,), second_animation=SECOND_ANIMATION)
        self.second_image()
        self.assertEqual(native().collect_shot2_batch_items(self.project, self.frames), [])
        self.assertFalse(self.snapshot()["frames"][0]["required_shot2"])
        self.check("passed", self.filled(), self.snapshot())

    def test_sidecar_does_not_invent_shot02_obligations(self):
        self.plan(declared=(1,))
        self.second_image()
        self.assertTrue(native().frame_needs_shot2_video_prompt(self.project, self.frames[0]))
        snapshot = self.snapshot(finalize_status=False)
        self.assertFalse(snapshot["frames"][0]["required_shot2"])
        self.check("passed", self.filled(), snapshot)

    def test_original_shot02_attrs_threshold_when_no_workbook_exists(self):
        frame = deepcopy(self.frames[0])
        for text, expected in (("x" * 9, False), ("x" * 10, True)):
            frame.attrs["animation_prompt_shot2"] = text
            with self.subTest(text=text):
                self.assertIs(native().has_animation_prompt_shot2_for_frame(self.project, frame), expected)

    def test_original_parser_maps_reordered_uuids_and_keeps_partial_coverage_visible(self):
        batch = native().collect_batch_items(self.project, self.frames)
        reply = json.dumps({"ops": [
            {"frame_uuid": f.uuid, "fields": {"промт_видео": f"Frame {f.number}: {ANIMATION_TEXT}"}}
            for f in reversed(self.frames)
        ]}, ensure_ascii=False)
        pairs = native().parse_animation_reply(reply, self.frames, batch_items=batch)
        ops = native().build_apply_ops_from_pairs(pairs, self.frames, shot=1)
        self.assertEqual([op["frame_uuid"] for op in ops], [self.frames[1].uuid, self.frames[0].uuid])
        self.assertTrue(ops[0]["fields"]["промт_видео"].startswith("Frame 2:"))
        partial = json.dumps({"ops": [{"frame_uuid": self.frames[1].uuid,
                                       "fields": {"animation_prompt": ANIMATION_TEXT}}]})
        pairs = native().parse_animation_reply(partial, self.frames, batch_items=batch)
        self.assertEqual([p.frame_number for p in pairs], [2])
        changed = self.filled()
        changed[0].animation_prompt = None
        self.check("failed", changed, self.snapshot())

    def test_original_parser_rejects_foreign_invalid_blob_and_unidentified_replies(self):
        batch = native().collect_batch_items(self.project, self.frames)
        for reply in (
            '{"ops": }',
            json.dumps({"ops": [{"frame_uuid": "foreign", "fields": {"промт_видео": ANIMATION_TEXT}}]}),
            json.dumps({"ops": [{"frame_uuid": self.frames[0].uuid, "fields": {"промт_видео": "short"}}]}),
            json.dumps({"ops": [{"frame_uuid": self.frames[0].uuid, "fields": {"промт_видео": '{"ops":[]}'}}]}),
            f"текст анимации: {ANIMATION_TEXT}",
        ):
            with self.subTest(reply=reply):
                self.assertEqual(native().parse_animation_reply(reply, self.frames, batch_items=batch), [])

    def test_original_parser_limits_json_to_current_batch_and_keeps_legacy_ids(self):
        batch = native().collect_batch_items(self.project, self.frames)
        reply = json.dumps({"ops": [
            {"frame_uuid": f.uuid, "fields": {"промт_видео": ANIMATION_TEXT}}
            for f in self.frames
        ]})
        pairs = native().parse_animation_reply(reply, self.frames, batch_items=batch[:1])
        self.assertEqual([p.frame_number for p in pairs], [1])
        legacy = f"ID изображения: {batch[1].image_id}\nтекст анимации: {SECOND_ANIMATION}"
        pairs = native().parse_animation_reply(legacy, self.frames, batch_items=batch)
        self.assertEqual([p.frame_number for p in pairs], [2])
        self.assertEqual(pairs[0].animation_text, SECOND_ANIMATION)

    def test_original_writer_builds_shot_fields_without_uuid_or_vo_rewrites(self):
        from app.services.node_write_contract import filter_ops_for_node

        pairs = [native().ParsedAnimationPair(
            image_id=self.frames[0].uuid, animation_text=ANIMATION_TEXT, frame_number=1,
        )]
        self.assertEqual(native().build_apply_ops_from_pairs(pairs, self.frames, shot=1), [
            {"frame_uuid": self.frames[0].uuid, "fields": {"промт_видео": ANIMATION_TEXT}}
        ])
        self.assertEqual(native().build_apply_ops_from_pairs(pairs, self.frames, shot=2), [
            {"frame_uuid": self.frames[0].uuid, "fields": {"промт_видео_2": ANIMATION_TEXT}}
        ])
        ops = [{"frame_uuid": self.frames[0].uuid, "fields": {
            "промт_видео": ANIMATION_TEXT, "промт_видео_2": SECOND_ANIMATION,
            "закадр": "forbidden", "image_prompt": "forbidden",
            "image_prompt_shot2": "forbidden", "child_prompts": {"forbidden": True},
            "characters": ["forbidden"], "sort_key": 100,
        }}]
        before = deepcopy(ops)
        self.assertEqual(filter_ops_for_node(ops, node_kind="anim_pr"), [
            {"frame_uuid": self.frames[0].uuid, "fields": {
                "промт_видео": ANIMATION_TEXT, "промт_видео_2": SECOND_ANIMATION,
            }}
        ])
        self.assertEqual(ops, before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
