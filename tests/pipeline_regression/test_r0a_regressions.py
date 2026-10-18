"""Регрессии R0a: дедуп xlsx, full-VO в b_user, shot grammar, only_node."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.models import Base, Frame, Project, ProjectStatus
from app.services.apply_ops_batches import ACTION_STAGGER_SEC, SHOTS_PARALLEL
from app.services.scene_shot_grammar import merge_trivial_shots
from app.services.scene_skeleton_batches import b_user
from app.services.xlsx_sync import reload_from_xlsx
from app.storage.project_sheet import ROW_HEADER, ROW_VOICEOVER, SHEET_FRAMES


def _frames_xlsx(path: Path, columns: list[tuple[int, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_FRAMES
    ws.cell(row=ROW_HEADER, column=1, value="Инфа")
    for col, (number, voice) in enumerate(columns, start=2):
        ws.cell(row=ROW_HEADER, column=col, value=number)
        ws.cell(row=ROW_VOICEOVER, column=col, value=voice)
    wb.save(path)


@pytest.mark.asyncio
async def test_repeat_xlsx_import_does_not_duplicate_frames(tmp_path: Path) -> None:
    """Баг 747 кадров: повторный импорт при уже существующих кадрах не плодит строки и VO."""
    xlsx = tmp_path / "project.xlsx"
    existing = "Уже лежащий в базе закадр кадра один."
    _frames_xlsx(
        xlsx,
        [
            (1, existing),
            (2, "Новый хвост, которого в базе нет и быть не должно."),
            (99, existing),
        ],
    )
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as session:
            project = Project(
                slug="no-dup",
                topic="t",
                status=ProjectStatus.frames_ready,
            )
            session.add(project)
            await session.flush()
            session.add(
                Frame(
                    project_id=project.id,
                    number=1,
                    voiceover_text=existing,
                )
            )
            await session.commit()

            first = await reload_from_xlsx(session, project, xlsx)
            await session.commit()
            second = await reload_from_xlsx(session, project, xlsx)
            await session.commit()

            rows = (
                await session.execute(
                    select(Frame).where(Frame.project_id == project.id)
                )
            ).scalars().all()
            voices = [r.voiceover_text for r in rows]
            assert [r.number for r in rows] == [1]
            assert voices == [existing]
            assert 2 in first.get("frames_create_skipped", [])
            assert second.get("frames_created", []) == []
            assert first.get("frames_created", []) == []
    finally:
        await engine.dispose()


def test_full_vo_block_is_in_b_user() -> None:
    frags = ["первый фрагмент", "второй фрагмент", "третий фрагмент"]
    text = b_user(frags, 0, 1, [], [], [])
    assert "===ВЕСЬ_ЗАКАДР" in text
    assert text.index("===ВЕСЬ_ЗАКАДР") < text.index("===УЖЕ_ЕСТЬ===")
    block = text.split("===ВЕСЬ_ЗАКАДР", 1)[1].split("===УЖЕ_ЕСТЬ===", 1)[0]
    for frag in frags:
        assert frag in block


def test_shot_vo_fill_calls_merge_spread_wide() -> None:
    src = Path("app/services/apply_ops_batches.py").read_text(encoding="utf-8")
    start = src.index('if kind == "shots_coverage"')
    block = src[start : start + 1800]
    assert 'shot_vo_fill' in block
    assert "merge_trivial_shots" in block
    assert "spread_vo_sentences" in block
    assert "ensure_wide_before_dialogue" in block
    assert SHOTS_PARALLEL == 16
    assert ACTION_STAGGER_SEC == 0.0


def _shot(sid: str, action: str, vo: str) -> dict:
    return {
        "id": sid,
        "место": "кухня",
        "кто": "c01",
        "действие": action,
        "план": "средний",
        "ракурс": "фронт",
        "объект": "стол",
        "закадр": vo,
    }


def test_merge_trivial_shots_keeps_vo_and_action() -> None:
    merged = [
        _shot("s1", "c01 сидит за столом", "Он сидит молча."),
        _shot("s2", "c01 продолжая держать чашку", ""),
    ]
    notes = merge_trivial_shots(merged)
    assert notes
    from app.orchestrator.validators.activity_flags import shot_is_active
    live = [s for s in merged if shot_is_active(s)]
    assert len(live) == 1
    assert "Он сидит молча." in live[0]["закадр"]
    assert "сидит за столом" in live[0]["действие"]
    assert "чашку" in live[0]["действие"]
    assert [s["id"] for s in merged if not shot_is_active(s)] == ["s2"]

    kept = [
        _shot("a", "c01 сидит за столом", "Первое предложение целиком."),
        _shot("b", "c01 продолжая держать чашку", "Второе предложение целиком."),
    ]
    assert merge_trivial_shots(kept) == []
    assert len(kept) == 2
    assert kept[0]["закадр"] == "Первое предложение целиком."
    assert kept[1]["закадр"] == "Второе предложение целиком."


@pytest.mark.asyncio
async def test_only_node_stop_key_written_before_step_start(monkeypatch) -> None:
    from app.web.routers import projects as pr

    events: list = []
    project = SimpleNamespace(
        id=7,
        meta={},
        status=SimpleNamespace(value="script_ready"),
    )

    class Sess:
        async def get(self, model, pid):
            return project

        async def refresh(self, obj):
            events.append("refresh")

    async def _commit(session):
        events.append(
            ("commit", (project.meta or {}).get("excel_gpt_stop_after_key"))
        )

    async def _start(session, p, step_code, **kwargs):
        events.append(
            ("start", (p.meta or {}).get("excel_gpt_stop_after_key"), step_code)
        )

    async def _noop(*_a, **_k):
        return None

    monkeypatch.setattr(pr, "commit_with_retry", _commit)
    monkeypatch.setattr(pr, "start_step", _start)
    monkeypatch.setattr(pr, "sync_run_for_project", _noop)
    monkeypatch.setattr(pr, "publish_project_event", _noop)
    monkeypatch.setattr(pr, "project_to_detail", lambda p: {"meta": dict(p.meta or {})})

    import app.project_db as pdb
    import sqlalchemy.orm.attributes as attrs

    monkeypatch.setattr(pdb, "push_runtime_to_master", _noop)
    monkeypatch.setattr(pdb, "sync_project_row_to_project_db", _noop)
    monkeypatch.setattr(attrs, "flag_modified", lambda *_a, **_k: None)

    await pr.run_project_step(
        7,
        "excel_gpt",
        dry_run=False,
        node_key="n_vo",
        only_node=True,
        session=Sess(),
    )
    start_i = next(
        i for i, ev in enumerate(events) if isinstance(ev, tuple) and ev[0] == "start"
    )
    commit_i = next(
        i for i, ev in enumerate(events) if isinstance(ev, tuple) and ev[0] == "commit"
    )
    assert commit_i < start_i
    assert events[commit_i][1] == "n_vo"
    assert events[start_i][1] == "n_vo"
    assert project.meta["excel_gpt_stop_after_key"] == "n_vo"
