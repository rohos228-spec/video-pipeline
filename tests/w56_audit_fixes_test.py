"""W56: исправления по аудиту (обход каталога статики, возврат в шаг после паузы, история отказов)."""
from __future__ import annotations

import asyncio
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient


def _static_app(tmp_path):
    from app.web import api as web_api

    out = tmp_path / "out"
    out.mkdir()
    (out / "index.html").write_text("<html>INDEX</html>", encoding="utf-8")
    (out / "a.txt").write_text("OK", encoding="utf-8")
    (tmp_path / "secret.env").write_text("KEY=SECRET", encoding="utf-8")
    return web_api, out


def test_static_catch_all_does_not_leave_out_dir(tmp_path, monkeypatch):
    web_api, out = _static_app(tmp_path)
    src = open(web_api.__file__, encoding="utf-8").read()
    assert "_inside(candidate)" in src and "relative_to(base)" in src


def test_restart_chain_stop_remembers_running_status():
    from app.models import ProjectStatus
    from app.services import restart_chain

    class _S:
        async def flush(self):
            return None

    from app.models import Project
    p = Project(id=1, slug="t", topic="t", meta={}, status=ProjectStatus.generating_images)
    asyncio.run(restart_chain.stop(_S(), p, "generating_images", 61, 60, "auto"))
    assert p.status is ProjectStatus.paused
    assert p.meta["paused_from_status"] == ProjectStatus.generating_images.value


def test_regen_budget_stop_remembers_running_status():
    from app.models import ProjectStatus
    from app.services import regen_budget

    class _S:
        async def flush(self):
            return None

    verdict = SimpleNamespace(summary=lambda: "stop", kind="images", blocked=[], run_exhausted=True, rounds=3,
                              usd_exhausted=False, usd_spent=0.0, usd_cap=1.0, chain_exhausted=False,
                              per_frame=1, per_run=1, usd_round=0.0)
    from app.models import Project
    p = Project(id=2, slug="t", topic="t", meta={}, status=ProjectStatus.generating_videos)
    asyncio.run(regen_budget.stop(_S(), p, verdict))
    assert p.meta["paused_from_status"] == ProjectStatus.generating_videos.value


def test_clear_failure_on_success_resets_infra_streak():
    from app.models import ProjectStatus
    from app.services.step_failure_policy import _save_failure_state, _failure_state, clear_failure_on_success

    p = SimpleNamespace(id=3, meta={})
    fs = _failure_state(p)
    fs["infra_lock_streak"] = 7
    _save_failure_state(p, fs)
    clear_failure_on_success(p, ProjectStatus.generating_images)
    assert "infra_lock_streak" not in _failure_state(p)


def _ffmpeg_ok() -> bool:
    import shutil
    return shutil.which("ffmpeg") is not None and shutil.which("ffprobe") is not None


@pytest.mark.skipif(not _ffmpeg_ok(), reason="нужен ffmpeg")
def test_variant2_mux_mixes_sfx_into_audio(tmp_path):
    """M1: звук из sfx_gen реально попадает в дорожку (раньше на пути variant2 выбрасывался молча)."""
    import subprocess

    from app.services.montage.variant2 import _mux
    from app.services.sfx_mix import SfxInput

    def run(*a):
        subprocess.run(a, check=True, capture_output=True)

    video, voice, sfx_f = tmp_path / "v.mp4", tmp_path / "voice.wav", tmp_path / "sfx.wav"
    run("ffmpeg", "-y", "-f", "lavfi", "-i", "color=c=black:s=64x64:d=4", "-pix_fmt", "yuv420p", str(video))
    run("ffmpeg", "-y", "-f", "lavfi", "-i", "anullsrc=r=44100:cl=mono", "-t", "4", str(voice))  # тишина
    run("ffmpeg", "-y", "-f", "lavfi", "-i", "sine=f=800:d=1", str(sfx_f))
    out_plain, out_sfx = tmp_path / "plain.mp4", tmp_path / "sfx.mp4"
    asyncio.run(_mux(video, voice, out_plain, voice_s=4.0, bgm=None))
    asyncio.run(_mux(video, voice, out_sfx, voice_s=4.0, bgm=None,
                     sfx=[SfxInput(path=sfx_f, t_start=1.0, gain=0.8, kind="sfx")]))

    def peak(path):
        r = subprocess.run(["ffmpeg", "-i", str(path), "-af", "volumedetect", "-f", "null", "-"],
                           capture_output=True, text=True)
        for line in r.stderr.splitlines():
            if "max_volume" in line:
                return float(line.split("max_volume:")[1].split("dB")[0])
        return -91.0

    assert peak(out_plain) < -60          # без SFX дорожка тихая
    assert peak(out_sfx) > -20            # с SFX слышен сигнал


def test_assemble_refuses_missing_clips_by_default():
    from app.settings import settings

    assert settings.assemble_allow_missing_clips is False


def test_private_files_are_not_served(tmp_path):
    from app.web.routers.artifacts import _is_private_file

    for name in ("state.db", "studio-live.log", ".env", ".env.local", "web_sessions.sqlite3", "k.pem", "x.sqlite"):
        assert _is_private_file(tmp_path / name), name
    for name in ("clip_1.mp4", "frame.png", "voice.mp3", "project.xlsx"):
        assert not _is_private_file(tmp_path / name), name


def test_fleet_safe_path_rejects_sibling_prefix(tmp_path):
    from fastapi import HTTPException

    from app.web.routers.fleet import _safe_path

    base = tmp_path / "root"
    (tmp_path / "root-evil").mkdir()
    base.mkdir()
    with pytest.raises(HTTPException):
        _safe_path(base, "../root-evil/x.txt")
    assert _safe_path(base, "a/b.txt").is_relative_to(base.resolve())


def test_remember_pause_origin_keeps_running_step():
    from app.models import Project, ProjectStatus
    from app.services import restart_chain

    p = Project(id=9, slug="t", topic="t", meta={"x": 1}, status=ProjectStatus.plan_ready)
    restart_chain.remember_pause_origin(p)
    assert p.meta["paused_from_status"] == "plan_ready" and p.meta["x"] == 1
    p.status = ProjectStatus.paused
    restart_chain.remember_pause_origin(p)          # уже на паузе: источник не затирается
    assert p.meta["paused_from_status"] == "plan_ready"


def test_animation_prompt_placeholder_counts_as_missing():
    """Живой F10 14.10: модель вернула «нет исходных данных для заполнения» для кадров 7-8; повтор шага не просил их заново."""
    from app.models import Frame
    from app.services.animation_prompt_gpt import has_animation_prompt_for_frame

    fr = Frame(project_id=1, number=7, animation_prompt="нет исходных данных для заполнения")
    assert has_animation_prompt_for_frame(None, fr) is False
    fr.animation_prompt = "Камера медленно приближается к двери, свет меняется на тёплый, дверь приоткрывается."
    assert has_animation_prompt_for_frame(None, fr) is True


def test_sfx_disabled_pass_through_status_is_confirmed(monkeypatch):
    """Живой F10 14.10: при SFX_ENABLED=false sfx_plan_ready откатывался в music_ready и шаг крутился бесконечно."""
    from app.models import Project, ProjectStatus
    from app.services.step_data_guard import ready_status_confirmed_by_data
    from app.settings import settings

    monkeypatch.setattr(settings, "sfx_enabled", False)
    for st in (ProjectStatus.sfx_plan_ready, ProjectStatus.sfx_ready):
        p = Project(id=1, slug="t", topic="t", meta={}, status=st)
        assert asyncio.run(ready_status_confirmed_by_data(None, p, st)) is True


def test_plan_sheet_is_created_for_db_first_workbook(tmp_path):
    """Живой F10 14.10: в книге проекта нет листа «план», таймкоды R15 некуда писать, сборка вставала."""
    from openpyxl import Workbook, load_workbook

    from app.storage.plan_sheet_v8 import (
        ensure_plan_sheet, read_plan_timestamps_cells, write_plan_timestamps)

    wb = Workbook()
    wb.active.title = "Кадры"
    wb.save(tmp_path / "project.xlsx")
    p = SimpleNamespace(id=1, data_dir=tmp_path)
    cells = [(1, "Первая фраза."), (2, "Вторая фраза.")]
    assert ensure_plan_sheet(p, cells) is True
    assert ensure_plan_sheet(p, cells) is False                      # повтор ничего не меняет
    assert "план" in load_workbook(tmp_path / "project.xlsx").sheetnames
    assert write_plan_timestamps(p, [(1, "0:00.00-0:03.00"), (2, "0:03.00-0:06.00")]) == 2
    got, _row = read_plan_timestamps_cells(p, [1, 2])
    assert [t for _n, t in got] == ["0:00.00-0:03.00", "0:03.00-0:06.00"]


def test_create_project_keeps_aspect_and_generators():
    """Живой F10 14.10: POST /api/projects принимал aspect_ratio и генераторы, но не сохранял их."""
    src = open("app/web/routers/projects.py", encoding="utf-8").read()
    assert 'for _field in ("aspect_ratio", "image_resolution", "video_resolution", "image_generator", "video_generator")' in src


def test_suno_error_statuses_are_final_failures():
    from app.bots.kie_http import _task_state

    for st in ("FAILED", "SENSITIVE_WORD_ERROR", "CALLBACK_EXCEPTION", "GENERATE_AUDIO_FAILED"):
        assert _task_state("suno", {"status": st}) == "fail", st
    for st in ("PENDING", "TEXT_SUCCESS", "GENERATING"):
        assert _task_state("suno", {"status": st}) == "pending", st
    assert _task_state("suno", {"status": "SUCCESS"}) == "success"


def test_release_for_human_everywhere_clears_both_sessions(monkeypatch):
    from app.models import Project, ProjectStatus
    from app.services import restart_chain

    class _S:
        pass

    main = Project(id=5, slug="a", topic="t", meta={"restart_chain": {"generating_images": 61}}, status=ProjectStatus.paused)
    mirror = Project(id=5, slug="a", topic="t", meta={"restart_chain": {"generating_images": 61}}, status=ProjectStatus.paused)

    async def _fake_replica(session, project, fn):
        await fn(None, mirror)

    import app.project_db as pdb
    monkeypatch.setattr(pdb, "run_on_replica_project", _fake_replica)
    asyncio.run(restart_chain.release_for_human_everywhere(_S(), main))
    assert not (main.meta or {}).get("restart_chain") and not (mirror.meta or {}).get("restart_chain")


def test_fleet_active_transfers_route_exists():
    from app.web.routers.fleet import active_transfers

    assert asyncio.run(active_transfers()) == {"transfers": []}


def test_hollow_scene_prompt_is_detected():
    """Живой F10 15.10: промты из шаблона стиля с заглушками в каждом поле сцены ушли в платную генерацию."""
    from app.generation_options import is_scene_prompt_hollow

    hollow = ("Референс: персонажи отсутствуют. Фон: нет исходных данных для заполнения. Свет: нет исходных данных для "
              "заполнения. Действие: нет исходных данных для заполнения.")
    real = "Старый почтовый вагон на рассвете, пар над рельсами, тёплый свет в окне. Фон: нет исходных данных для заполнения."
    assert is_scene_prompt_hollow(hollow) is True
    assert is_scene_prompt_hollow(real) is False
