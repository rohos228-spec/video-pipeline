from __future__ import annotations

from types import SimpleNamespace

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.services.xlsx_versioning import validate_xlsx_sheets
from app.storage.project_sheet import SHEET_FRAMES, SHEET_GENERAL
from app.services.xlsx_v8_import import SHEET_PLAN_V8


def _book(tmp_path, names):
    wb = openpyxl.Workbook()
    wb.active.title = names[0]
    for n in names[1:]:
        wb.create_sheet(n)
    p = tmp_path / "b.xlsx"
    wb.save(p)
    return p


def test_v7_with_plan_sheet_layout_is_valid(tmp_path):
    assert validate_xlsx_sheets(_book(tmp_path, [SHEET_FRAMES, SHEET_GENERAL, SHEET_PLAN_V8])) is None


def test_unknown_layout_still_rejected(tmp_path):
    assert validate_xlsx_sheets(_book(tmp_path, [SHEET_FRAMES, "Левый лист"])) is not None


def test_reload_xlsx_turns_validation_error_into_400():
    import inspect

    from app.web.routers import project_ops

    src = inspect.getsource(project_ops.reload_xlsx)
    assert "except RuntimeError" in src and "status_code=400" in src


def test_publications_router_reads_project_db():
    import inspect

    from app.web.routers import publications

    src = inspect.getsource(publications)
    assert "get_project_session" in src and "Depends(get_session)" not in src


def test_skipped_publication_reason_has_russian_text():
    from app.services.timeline_ru import publication_code_text

    assert "Автопубликация выключена" in publication_code_text("SOCIAL_PUBLISH_ENABLED=false")


def test_deleted_frame_is_not_resurrected_by_xlsx_import():
    from app.services.frame_tombstones import deleted_numbers, forget_deleted, remember_deleted

    p = SimpleNamespace(meta={})
    remember_deleted(p, 8)
    remember_deleted(p, 3)
    assert deleted_numbers(p) == {3, 8}
    forget_deleted(p, 3)
    assert deleted_numbers(p) == {8}
    forget_deleted(p)
    assert deleted_numbers(p) == set()


async def test_v8_import_skips_deleted_numbers(tmp_path):
    from openpyxl import Workbook
    from sqlalchemy import select
    from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

    from app.models import Base, Frame, Project, ProjectStatus
    from app.services.frame_tombstones import remember_deleted
    from app.services.xlsx_v8_import import ROW_VOICEOVER_V8, import_v8_xlsx

    wb = Workbook()
    ws = wb.active
    ws.title = "план"
    for i, text in enumerate(["один", "два", "три"], start=3):
        ws.cell(row=ROW_VOICEOVER_V8, column=i, value=text)
    xlsx = tmp_path / "project.xlsx"
    wb.save(xlsx)

    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    async with factory() as session:
        project = Project(slug="tomb", topic="t", status=ProjectStatus.frames_ready, meta={})
        session.add(project)
        await session.commit()
        remember_deleted(project, 3)
        await import_v8_xlsx(session, project, xlsx, keep_fields=True)
        nums = sorted(f.number for f in (await session.execute(select(Frame))).scalars().all())
        assert nums == [1, 2]
    await engine.dispose()


def test_websocket_requires_login_when_web_auth_enabled(monkeypatch):
    from starlette.websockets import WebSocketDisconnect

    from app.settings import settings
    from app.web import auth_sessions
    from app.web.api import create_app

    monkeypatch.setattr(settings, "web_auth_user", "boss")
    monkeypatch.setattr(settings, "web_auth_password", "s3cret-pass")
    auth_sessions.forget_memory_sessions()
    token = auth_sessions.issue_session_token("boss")
    with TestClient(create_app()) as client:
        with pytest.raises(WebSocketDisconnect):
            with client.websocket_connect("/ws/global") as ws:
                ws.receive_text()
        client.cookies.set("vp_session", token)
        with client.websocket_connect("/ws/global") as ws:
            assert '"subscribed"' in ws.receive_text()


async def test_inspector_title_edit_reaches_master_list(tmp_path, monkeypatch):
    from contextlib import asynccontextmanager

    from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

    import app.db as app_db
    from app import project_db
    from app.models import Base, Project, ProjectStatus

    master_engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'master.db'}")
    proj_engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'project.db'}")
    for eng in (master_engine, proj_engine):
        async with eng.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
    master_maker = async_sessionmaker(master_engine, expire_on_commit=False, class_=AsyncSession)
    proj_maker = async_sessionmaker(proj_engine, expire_on_commit=False, class_=AsyncSession)

    @asynccontextmanager
    async def fake_scope():
        async with master_maker() as s:
            yield s

    monkeypatch.setattr(app_db, "session_scope", fake_scope)
    async with master_maker() as m:
        m.add(Project(id=7, slug="p7", title="старое", topic="т", status=ProjectStatus.new, meta={}))
        await m.commit()
    async with proj_maker() as ps:
        ps.add(Project(id=7, slug="p7", title="новое", topic="т", status=ProjectStatus.new, meta={}))
        await ps.commit()
        project = await ps.get(Project, 7)
        await project_db.push_display_fields_to_master(ps, project, ["meta"])
        async with master_maker() as m:
            assert (await m.get(Project, 7)).title == "старое"
        await project_db.push_display_fields_to_master(ps, project, ["title"])
    async with master_maker() as m:
        assert (await m.get(Project, 7)).title == "новое"
    await master_engine.dispose()
    await proj_engine.dispose()


def test_slug_of_leftover_project_folder_is_taken(tmp_path, monkeypatch):
    from app.services.project_slug import slug_taken_on_disk
    from app.settings import settings

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    assert slug_taken_on_disk("doch-1") is False
    folder = tmp_path / "videos" / "doch-1"
    folder.mkdir(parents=True)
    assert slug_taken_on_disk("doch-1") is False
    (folder / "project.db").write_bytes(b"x")
    assert slug_taken_on_disk("doch-1") is True


def test_partial_hollow_image_prompts_also_block_paid_generation():
    import inspect

    from app.orchestrator.steps import generate_images

    src = inspect.getsource(generate_images.run)
    assert "if hollow:" in src and "len(hollow) == len(non_child)" not in src


async def test_public_image_hosts_are_off_by_default(monkeypatch):
    import pytest

    from app.bots import outsee_http
    from app.settings import settings

    monkeypatch.setattr("app.bots.yandex_storage.yandex_storage_configured", lambda: False)
    monkeypatch.setattr(settings, "outsee_allow_public_image_hosts", False)
    called = []

    async def boom(*a, **k):
        called.append(1)
        raise AssertionError("public host must not be called")

    for name in ("_host_via_litterbox", "_host_via_catbox", "_host_via_uguu", "_host_via_0x0"):
        monkeypatch.setattr(outsee_http, name, boom)
    png = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    with pytest.raises(outsee_http.OutseeApiError):
        await outsee_http.ensure_public_image_url(png)
    assert called == []
