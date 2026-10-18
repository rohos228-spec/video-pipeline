"""Ревью #352: неактивный кадр не попадает в колонку листа «план»."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from app.services.db_apply import export_project_xlsx
from app.services.xlsx_v8_import import (
    ROW_DURATION_V8,
    ROW_IMAGE_PROMPT_V8,
    ROW_VIDEO_PROMPT_V8,
    ROW_VOICEOVER_V8,
    hidden_frame_marker,
)


def test_export_skips_marked_frame_column(tmp_path: Path) -> None:
    path = tmp_path / "project.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "план"
    ws.cell(row=ROW_VOICEOVER_V8, column=4, value="старый-закадр")
    wb.save(path)
    wb.close()
    marked = SimpleNamespace(
        number=2,
        image_prompt="картинка",
        animation_prompt="видео",
        voiceover_text="секретный закадр",
        duration_seconds=3.0,
        start_ts=0.0,
        end_ts=3.0,
        attrs={"удалён": True, "активен": False, "карантин": True},
    )
    live = SimpleNamespace(
        number=1,
        image_prompt="",
        animation_prompt="",
        voiceover_text="живой",
        duration_seconds=1.0,
        start_ts=None,
        end_ts=None,
        attrs={},
    )
    project = SimpleNamespace(data_dir=tmp_path, meta={})
    export_project_xlsx(project, [live, marked])  # type: ignore[arg-type]
    wb2 = load_workbook(path)
    sheet = wb2["план"]
    assert sheet.cell(ROW_VOICEOVER_V8, 3).value == "живой"
    # Колонка остаётся на месте: маркер, не озвучка кадра и не пустота.
    assert sheet.cell(ROW_VOICEOVER_V8, 4).value == hidden_frame_marker(2)
    assert sheet.cell(ROW_IMAGE_PROMPT_V8, 4).value is None
    assert sheet.cell(ROW_VIDEO_PROMPT_V8, 4).value is None
    assert sheet.cell(ROW_DURATION_V8, 4).value is None
    wb2.close()
    assert marked.voiceover_text == "секретный закадр"
    assert marked.attrs["удалён"] is True
