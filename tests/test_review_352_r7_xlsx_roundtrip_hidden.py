"""Ревью #352: неактивный кадр не сдвигает соседей при export→import."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.models import Base, Frame, FrameStatus, Project, ProjectStatus
from app.services.db_apply import export_project_xlsx
from app.services.xlsx_v8_import import (
    INACTIVE_COLUMN_MARK,
    ROW_DURATION_V8,
    ROW_FRAME_ACTIVITY_V8,
    ROW_IMAGE_PROMPT_V8,
    ROW_VIDEO_PROMPT_V8,
    ROW_VOICEOVER_V8,
    SHEET_GENERAL_V8,
    SHEET_PLAN_V8,
)


def _snapshot(frames: list[Frame]) -> dict[int, tuple]:
    return {
        fr.number: (
            fr.voiceover_text,
            fr.image_prompt,
            fr.animation_prompt,
            fr.duration_seconds,
        )
        for fr in frames
    }


def _frames(project_id: int) -> list[Frame]:
    rows = [
        (1, "голос-1", "картинка-1", "видео-1", 1.5, {}),
        (
            2,
            "голос-2",
            "картинка-2",
            "видео-2",
            2.5,
            {"удалён": True, "активен": False, "карантин": True},
        ),
        (3, "голос-3", "картинка-3", "видео-3", 3.5, {}),
        (4, "голос-4", "картинка-4", "видео-4", 4.5, {}),
    ]
    return [
        Frame(
            project_id=project_id,
            number=number,
            voiceover_text=voice,
            image_prompt=image,
            animation_prompt=anim,
            duration_seconds=dur,
            status=FrameStatus.planned,
            attrs=attrs,
        )
        for number, voice, image, anim, dur, attrs in rows
    ]


async def _project(session: AsyncSession, slug: str, root: Path) -> tuple[Project, list[Frame]]:
    project = Project(
        slug=slug,
        topic="t",
        status=ProjectStatus.frames_ready,
        script_text="голос-1 голос-3 голос-4",
        meta={},
    )
    session.add(project)
    await session.flush()
    frames = _frames(project.id)
    session.add_all(frames)
    await session.flush()
    folder = root / "videos" / slug
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "project.xlsx"
    wb = Workbook()
    wb.active.title = SHEET_PLAN_V8
    wb.create_sheet(SHEET_GENERAL_V8)
    wb.save(path)
    wb.close()
    return project, frames


@pytest.mark.asyncio
async def test_hidden_frame_roundtrip_keeps_neighbors(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("app.settings.settings.data_dir", str(tmp_path))
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)

    async with factory() as session:
        bot_project, bot_frames = await _project(session, "bot-round", tmp_path)
        gpt_project, gpt_frames = await _project(session, "gpt-round", tmp_path)
        before_bot = _snapshot(bot_frames)
        before_gpt = _snapshot(gpt_frames)
        export_project_xlsx(bot_project, bot_frames)
        exported = bot_project.data_dir / "project.xlsx"
        target = gpt_project.data_dir / "project.xlsx"
        target.write_bytes(exported.read_bytes())

        wb = load_workbook(exported)
        sheet = wb[SHEET_PLAN_V8]
        assert sheet.cell(ROW_VOICEOVER_V8, 3).value == "голос-1"
        assert sheet.cell(ROW_VOICEOVER_V8, 4).value == "голос-2"
        assert sheet.cell(ROW_IMAGE_PROMPT_V8, 4).value == "картинка-2"
        assert sheet.cell(ROW_VIDEO_PROMPT_V8, 4).value == "видео-2"
        assert sheet.cell(ROW_DURATION_V8, 4).value == 2.5
        assert sheet.cell(ROW_VOICEOVER_V8, 5).value == "голос-3"
        assert sheet.cell(ROW_VOICEOVER_V8, 6).value == "голос-4"
        assert sheet.cell(ROW_FRAME_ACTIVITY_V8, 1).value == "активен"
        assert sheet.cell(ROW_FRAME_ACTIVITY_V8, 3).value is None
        assert sheet.cell(ROW_FRAME_ACTIVITY_V8, 4).value == INACTIVE_COLUMN_MARK
        assert sheet.cell(ROW_FRAME_ACTIVITY_V8, 5).value is None
        assert sheet.cell(ROW_FRAME_ACTIVITY_V8, 6).value is None
        wb.close()

        from app.services.xlsx_v8_import import import_v8_xlsx

        await import_v8_xlsx(
            session,
            bot_project,
            exported,
            keep_fields=False,
            update_frames_voiceover=True,
        )
        for fr in bot_frames:
            await session.refresh(fr)
        assert _snapshot(bot_frames) == before_bot
        assert bot_frames[1].attrs["активен"] is False
        assert bot_frames[1].attrs["удалён"] is True
        assert bot_frames[1].attrs["карантин"] is True

        from app.services.chatgpt_xlsx import sync_project_xlsx

        await sync_project_xlsx(session, gpt_project, target)
        for fr in gpt_frames:
            await session.refresh(fr)
        assert _snapshot(gpt_frames) == before_gpt
        assert gpt_frames[1].attrs["активен"] is False
        assert gpt_frames[1].attrs["удалён"] is True
        assert gpt_frames[1].attrs["карантин"] is True

    await engine.dispose()


def _write_gapped_plan(path: Path) -> None:
    """Кадр 2 помечен, колонка кадра 3 пустая, соседи 1 и 4 со своими номерами."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PLAN_V8
    wb.create_sheet(SHEET_GENERAL_V8)
    ws.cell(ROW_VOICEOVER_V8, 3, "голос-1")
    ws.cell(ROW_IMAGE_PROMPT_V8, 3, "новая-1")
    ws.cell(ROW_VOICEOVER_V8, 4, "яд-голос")
    ws.cell(ROW_IMAGE_PROMPT_V8, 4, "яд-картинка")
    ws.cell(ROW_VIDEO_PROMPT_V8, 4, "яд-видео")
    ws.cell(ROW_DURATION_V8, 4, 9.0)
    ws.cell(ROW_FRAME_ACTIVITY_V8, 1, "активен")
    ws.cell(ROW_FRAME_ACTIVITY_V8, 4, INACTIVE_COLUMN_MARK)
    ws.cell(ROW_VOICEOVER_V8, 6, "голос-4")
    ws.cell(ROW_IMAGE_PROMPT_V8, 6, "новая-4")
    wb.save(path)
    wb.close()


def _assert_gap_import(frames: list[Frame]) -> None:
    assert frames[0].voiceover_text == "голос-1"
    assert frames[0].image_prompt == "новая-1"
    assert frames[0].animation_prompt == "видео-1"
    assert frames[1].voiceover_text == "голос-2"
    assert frames[1].image_prompt == "картинка-2"
    assert frames[1].animation_prompt == "видео-2"
    assert frames[1].duration_seconds == 2.5
    assert frames[1].attrs["активен"] is False
    assert frames[1].attrs["удалён"] is True
    assert frames[2].voiceover_text == "голос-3"
    assert frames[2].image_prompt == "картинка-3"
    assert frames[2].animation_prompt == "видео-3"
    assert frames[2].duration_seconds == 3.5
    assert frames[3].voiceover_text == "голос-4"
    assert frames[3].image_prompt == "новая-4"
    assert frames[3].animation_prompt == "видео-4"


@pytest.mark.asyncio
async def test_empty_column_and_mark_do_not_shift_neighbors(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("app.settings.settings.data_dir", str(tmp_path))
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)

    async with factory() as session:
        bot_project, bot_frames = await _project(session, "bot-gap", tmp_path)
        gpt_project, gpt_frames = await _project(session, "gpt-gap", tmp_path)
        exported = bot_project.data_dir / "project.xlsx"
        _write_gapped_plan(exported)
        target = gpt_project.data_dir / "project.xlsx"
        target.write_bytes(exported.read_bytes())

        from app.services.xlsx_v8_import import import_v8_xlsx

        await import_v8_xlsx(
            session,
            bot_project,
            exported,
            keep_fields=False,
            update_frames_voiceover=True,
        )
        for fr in bot_frames:
            await session.refresh(fr)
        _assert_gap_import(bot_frames)

        from app.services.chatgpt_xlsx import sync_project_xlsx

        await sync_project_xlsx(session, gpt_project, target)
        for fr in gpt_frames:
            await session.refresh(fr)
        _assert_gap_import(gpt_frames)

    await engine.dispose()
