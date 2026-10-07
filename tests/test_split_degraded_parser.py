"""split: which path produced the frames, pre-write validation of the model answer,
and the guard in front of replace_all_frames. Stubs only; no network."""
from __future__ import annotations

import contextlib
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from app.models import Project, ProjectStatus
from app.services import xlsx_step_runners as xsr
from app.services.xlsx_step_runners import (
    extract_frames_spec_from_gpt_reply,
    extract_frames_spec_with_source,
    split_answer_covers_voiceover,
)

VO = (
    "Сегодня мы разберём, как устроена настройка проекта с нуля. "
    "Сначала откроем папку и создадим файл конфигурации. "
    "Потом запустим первую проверку и посмотрим на результат. "
    "В конце сохраним изменения и закроем окно редактора."
)
REFUSAL = "Не могу выполнить эту просьбу - объясню - почему так получилось у меня сегодня"
DASH_OK = (
    "Сегодня мы разберём, как устроена настройка проекта с нуля.\n-\n"
    "Сначала откроем папку и создадим файл конфигурации.\n-\n"
    "Потом запустим первую проверку и посмотрим на результат.\n-\n"
    "В конце сохраним изменения и закроем окно редактора."
)


@pytest.fixture()
def vo_path(tmp_path: Path) -> Path:
    p = tmp_path / "voiceover.txt"
    p.write_text(VO, encoding="utf-8")
    return p


def test_source_json(vo_path):
    reply = json.dumps({"ops": [{"target": "replace_frames", "frames": [{"закадр": "первый кадр"}, {"закадр": "второй кадр"}]}]})
    spec, source, info = extract_frames_spec_with_source(reply, voiceover_path=vo_path)
    assert source == "json_replace_frames" and len(spec) == 2 and info["dash_rejected"] is False


def test_source_dash_blocks_covering_voiceover_accepted(vo_path):
    spec, source, info = extract_frames_spec_with_source(DASH_OK, voiceover_path=vo_path)
    assert source == "dash_blocks" and len(spec) == 4 and info["dash_rejected"] is False


def test_refusal_with_inline_dash_is_not_accepted_as_frames(vo_path):
    spec, source, info = extract_frames_spec_with_source(REFUSAL, voiceover_path=vo_path)
    assert info["dash_rejected"] is True
    assert source == "local_voiceover_fallback"
    joined = " ".join(f["закадр"] for f in spec)
    assert "Не могу" not in joined and joined.split() == VO.split()


def test_refusal_with_dash_lines_is_not_accepted(vo_path):
    reply = "Извините, не могу помочь.\n- я не вижу файла\n- попробуйте приложить его ещё раз позже"
    spec, source, _ = extract_frames_spec_with_source(reply, voiceover_path=vo_path)
    assert source == "local_voiceover_fallback"
    assert all("файла" not in f["закадр"] for f in spec)


def test_plain_garbage_uses_local_fallback_and_no_voiceover_gives_none(vo_path):
    spec, source, info = extract_frames_spec_with_source("мусор без разделителей", voiceover_path=vo_path)
    assert source == "local_voiceover_fallback" and len(spec) >= 2 and info["dash_rejected"] is False
    assert extract_frames_spec_with_source("мусор", voiceover_path=None) == ([], "none", {"dash_rejected": False})


def test_without_voiceover_dash_is_accepted_as_before():
    # Nothing to compare with: the previous behaviour is kept (documented limitation).
    spec, source, _ = extract_frames_spec_with_source(REFUSAL, voiceover_path=None)
    assert source == "dash_blocks" and len(spec) >= 2


def test_legacy_wrapper_matches_new_function(vo_path):
    for reply in (DASH_OK, REFUSAL, "x"):
        assert extract_frames_spec_from_gpt_reply(reply, voiceover_path=vo_path) == \
            extract_frames_spec_with_source(reply, voiceover_path=vo_path)[0]


def test_coverage_tolerance_cases():
    blocks = [s.strip() for s in VO.split(".") if s.strip()]
    assert split_answer_covers_voiceover(blocks, VO)
    assert not split_answer_covers_voiceover(["совсем другой текст про погоду"], VO)
    assert not split_answer_covers_voiceover(blocks[:2], VO)  # lost tail
    assert split_answer_covers_voiceover(["что угодно"], "")  # nothing to compare


@pytest.mark.asyncio
async def test_record_split_parse_carries_source_and_fallback_flag():
    from app.services.node_capabilities import record_split_parse
    from app.services.observation import _attempt

    ctx = SimpleNamespace(scope=SimpleNamespace(project_id=2), last_request_id="r", split_input_voiceover=VO)
    token = _attempt.set(ctx)
    try:
        spec = [{"закадр": s.strip() + "."} for s in VO.split(".") if s.strip()]
        with patch("app.services.observation.record_fact", new_callable=AsyncMock) as fact:
            await record_split_parse("junk", spec, frames_source="local_voiceover_fallback", dash_rejected=True)
            parse = fact.call_args_list[0].args[1]
        assert parse["frames_source"] == "local_voiceover_fallback"
        assert parse["fallback_used"] is True and parse["model_payload_usable"] is False
        assert parse["dash_rejected"] is True
        assert parse["payload_required"] is False and parse["contract"] == "native_split"
        with patch("app.services.observation.record_fact", new_callable=AsyncMock) as fact:
            await record_split_parse("junk", spec, frames_source="json_replace_frames")
            parse = fact.call_args_list[0].args[1]
        assert parse["fallback_used"] is False and parse["model_payload_usable"] is True
    finally:
        _attempt.reset(token)


def _patches(tmp_dir: Path, ask):
    async def _run_lock(_pid, _step, fn):
        return await fn()

    return (
        patch.object(xsr.cx, "tmp_gpt_dir", return_value=tmp_dir),
        patch.object(xsr.cx, "write_split_prompt_file", return_value=tmp_dir / "p.txt"),
        patch.object(xsr.cx, "chat_message", return_value="msg"),
        patch.object(xsr, "_ts", return_value="20260101_120000"),
        patch.object(xsr.xgf, "run_under_xlsx_lock", new=_run_lock),
        patch.object(xsr.xgf, "telegram_style_ask_with_files", ask),
    )


async def _run_split(tmp_path, monkeypatch, reply):
    from openpyxl import Workbook

    from app.services.xlsx_v8_import import SHEET_PLAN_V8

    proj_dir = tmp_path / "data" / "videos" / "p9"
    proj_dir.mkdir(parents=True)
    wb = Workbook(); wb.active.title = SHEET_PLAN_V8; wb.create_sheet("Общий план"); wb.save(proj_dir / "project.xlsx")
    (proj_dir / "voiceover.txt").write_text(VO, encoding="utf-8")
    tmp_dir = proj_dir / "tmp_gpt"; tmp_dir.mkdir(); (tmp_dir / "p.txt").write_text("p", encoding="utf-8")
    project = Project(id=9, topic="t", slug="p9")
    monkeypatch.setattr("app.models.settings.data_dir", tmp_path / "data")
    monkeypatch.setattr(type(project), "data_dir", property(lambda self: proj_dir))
    with contextlib.ExitStack() as st:
        for cm in _patches(tmp_dir, AsyncMock(return_value=reply)):
            st.enter_context(cm)
        return await xsr.run_split_xlsx(project)


@pytest.mark.asyncio
async def test_run_split_refusal_falls_back_locally_and_is_flagged(tmp_path, monkeypatch):
    result = await _run_split(tmp_path, monkeypatch, REFUSAL)
    assert result.split_source == "local_voiceover_fallback" and result.split_verified is True
    assert " ".join(f["закадр"] for f in result.frames_spec).split() == VO.split()


@pytest.mark.asyncio
async def test_run_split_happy_path_dash_unchanged(tmp_path, monkeypatch):
    result = await _run_split(tmp_path, monkeypatch, DASH_OK)
    assert result.split_source == "dash_blocks"
    assert [f["закадр"] for f in result.frames_spec] == [b.strip() for b in DASH_OK.split("\n-\n")]
    assert result.apply_ops == [{"target": "replace_frames", "frames": result.frames_spec}]


@pytest.mark.asyncio
async def test_split_frames_never_calls_apply_with_unverified_result():
    from app.orchestrator.steps import split_frames

    project = SimpleNamespace(id=1, status=ProjectStatus.splitting, meta={})
    session = MagicMock()
    execute = MagicMock(); execute.scalars.return_value.all.return_value = []
    session.execute = AsyncMock(return_value=execute)
    bad = SimpleNamespace(apply_ops=[{"target": "replace_frames", "frames": [{"закадр": "a"}, {"закадр": "b"}]}],
        frames_spec=[], split_source="dash_blocks", split_verified=False)
    apply_ops = AsyncMock()
    with patch.object(split_frames.xsr, "run_split_xlsx", AsyncMock(return_value=bad)), \
            patch("app.services.db_apply.apply_ops", apply_ops):
        with pytest.raises(RuntimeError, match="предпроверку"):
            await split_frames.run(session, project)
    apply_ops.assert_not_called()



# ---------------------------------------------------------------------------
# E1b (owner decision 2026-10-07): validate BEFORE write. In a managed attempt an answer
# that violates the split contract fails with SplitContractRejected before any DB write.
# ---------------------------------------------------------------------------

BAD_JSON = json.dumps({"ops": [{"target": "replace_frames", "frames": [
    {"закадр": "Сегодня мы разберём погоду."}, {"закадр": "Совсем другой текст."}]}]})


def _managed_ctx():
    return SimpleNamespace(scope=SimpleNamespace(project_id=9, mode="managed", observer_state="active"),
                           last_request_id="r", split_input_voiceover=None)


@contextlib.contextmanager
def _attempt_ctx(ctx):
    from app.services.observation import _attempt

    token = _attempt.set(ctx)
    try:
        with patch("app.services.observation.record_fact", new_callable=AsyncMock) as fact,                 patch("app.services.observation.emit", MagicMock()):
            yield fact
    finally:
        _attempt.reset(token)


def _split_project(tmp_path: Path, monkeypatch):
    from openpyxl import Workbook

    from app.services.xlsx_v8_import import SHEET_PLAN_V8

    proj_dir = tmp_path / "data" / "videos" / "p9"
    proj_dir.mkdir(parents=True)
    wb = Workbook(); wb.active.title = SHEET_PLAN_V8; wb.create_sheet("Общий план"); wb.save(proj_dir / "project.xlsx")
    (proj_dir / "voiceover.txt").write_text(VO, encoding="utf-8")
    tmp_dir = proj_dir / "tmp_gpt"; tmp_dir.mkdir(); (tmp_dir / "p.txt").write_text("p", encoding="utf-8")
    project = Project(id=9, topic="t", slug="p9", status=ProjectStatus.splitting, meta={})
    monkeypatch.setattr("app.models.settings.data_dir", tmp_path / "data")
    monkeypatch.setattr(type(project), "data_dir", property(lambda self: proj_dir))
    return project, tmp_dir


def _session_without_frames():
    session = MagicMock()
    execute = MagicMock(); execute.scalars.return_value.all.return_value = []
    session.execute = AsyncMock(return_value=execute)
    session.flush = AsyncMock(); session.commit = AsyncMock()
    return session


@pytest.mark.asyncio
async def test_managed_invalid_answer_is_rejected_before_any_write(tmp_path, monkeypatch):
    from app.orchestrator.steps import split_frames
    from app.services.node_capabilities import SplitContractRejected

    project, tmp_dir = _split_project(tmp_path, monkeypatch)
    session = _session_without_frames()
    apply_ops = AsyncMock()
    with contextlib.ExitStack() as st:
        for cm in _patches(tmp_dir, AsyncMock(return_value=BAD_JSON)):
            st.enter_context(cm)
        st.enter_context(patch("app.services.db_apply.apply_ops", apply_ops))
        fact = st.enter_context(_attempt_ctx(_managed_ctx()))
        with pytest.raises(SplitContractRejected) as ei:
            await split_frames.run(session, project)
    assert ei.value.failure_class == "contract_rejected"
    assert "voiceover_changed_or_missing" in ei.value.errors
    apply_ops.assert_not_called()                      # nothing written to frames/shots
    session.commit.assert_not_called()
    assert project.status is ProjectStatus.splitting and project.meta == {}
    # evidence is journaled before the failure: parse_result + failed runtime_check
    kinds = [c.args[0] for c in fact.call_args_list]
    assert "parse_result" in kinds and "runtime_check" in kinds
    check = [c.args[1] for c in fact.call_args_list if c.args[0] == "runtime_check"][0]
    assert check["status"] == "failed" and "voiceover_changed_or_missing" in check["errors"]


@pytest.mark.asyncio
async def test_managed_valid_answers_are_identical_to_native(tmp_path, monkeypatch):
    good_json = json.dumps({"ops": [{"target": "replace_frames", "frames": [
        {"закадр": s} for s in DASH_OK.split("\n-\n")]}]})
    for i, reply in enumerate((DASH_OK, REFUSAL, good_json)):
        native = await _run_split(tmp_path / f"n{i}", monkeypatch, reply)
        with _attempt_ctx(_managed_ctx()):
            managed = await _run_split(tmp_path / f"m{i}", monkeypatch, reply)
        assert managed.frames_spec == native.frames_spec
        assert managed.apply_ops == native.apply_ops
        assert managed.split_source == native.split_source and managed.split_verified is True


@pytest.mark.asyncio
async def test_unmanaged_behaviour_is_unchanged_for_the_same_answer(tmp_path, monkeypatch):
    result = await _run_split(tmp_path, monkeypatch, BAD_JSON)   # no attempt context: native path
    assert result.split_source == "json_replace_frames" and len(result.frames_spec) == 2
    with _attempt_ctx(SimpleNamespace(scope=SimpleNamespace(project_id=9, mode="observe", observer_state="active"),
                                      last_request_id="r", split_input_voiceover=None)):
        result = await _run_split(tmp_path / "obs", monkeypatch, BAD_JSON)   # observe mode: no gate
    assert result.split_source == "json_replace_frames"


@pytest.mark.asyncio
async def test_rejected_attempt_leaves_nothing_behind_and_a_later_attempt_succeeds(tmp_path, monkeypatch):
    """Restart safety: a rejected answer writes no state; the next attempt starts clean."""
    from app.orchestrator.steps import split_frames
    from app.services.node_capabilities import SplitContractRejected

    project, tmp_dir = _split_project(tmp_path, monkeypatch)
    before = (project.status, dict(project.meta), sorted(p.name for p in tmp_dir.parent.iterdir()))
    session = _session_without_frames()
    apply_ops = AsyncMock(return_value={"replace_frames": 4})
    for _ in range(2):  # the same rejection twice: deterministic, still no write
        with contextlib.ExitStack() as st:
            for cm in _patches(tmp_dir, AsyncMock(return_value=BAD_JSON)):
                st.enter_context(cm)
            st.enter_context(patch("app.services.db_apply.apply_ops", apply_ops))
            st.enter_context(_attempt_ctx(_managed_ctx()))
            with pytest.raises(SplitContractRejected):
                await split_frames.run(session, project)
    apply_ops.assert_not_called()
    assert (project.status, dict(project.meta), sorted(p.name for p in tmp_dir.parent.iterdir())) == before
    # a valid answer afterwards reaches the writer exactly once with the original ops
    with contextlib.ExitStack() as st:
        for cm in _patches(tmp_dir, AsyncMock(return_value=DASH_OK)):
            st.enter_context(cm)
        st.enter_context(patch("app.services.db_apply.apply_ops", apply_ops))
        st.enter_context(patch("app.services.db_apply.export_project_xlsx", MagicMock()))
        st.enter_context(patch("app.services.node_xlsx_snapshot.snapshot_and_bind_node_xlsx", AsyncMock()))
        st.enter_context(patch("app.services.storage_step_sync.sync_storage_after_step", AsyncMock()))
        st.enter_context(_attempt_ctx(_managed_ctx()))
        try:
            await split_frames.run(session, project)
        except RuntimeError as e:  # post-write checks need a real DB; the write gate is what matters
            assert "кадров" in str(e)
    apply_ops.assert_called_once()
    ops = apply_ops.call_args.args[2]
    assert ops == [{"target": "replace_frames", "frames": [{"закадр": b} for b in DASH_OK.split("\n-\n")]}]
