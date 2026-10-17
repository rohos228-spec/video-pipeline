"""Вызовы #342, без которых данные тихо теряются. Каждый сайт — отдельная проверка."""

from __future__ import annotations

import ast
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def _fn(path: str, name: str) -> str:
    src = (ROOT / path).read_text(encoding="utf-8")
    tree = ast.parse(src)
    fn = next(
        n
        for n in ast.walk(tree)
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name == name
    )
    return ast.get_source_segment(src, fn) or ""


def test_delete_montage_frame_remembers_deleted_on_both_paths() -> None:
    text = _fn("app/services/montage_board_frames.py", "delete_montage_frame")
    assert text.count("remember_deleted(project, number)") == 2


def test_delete_coverage_child_remembers_deleted() -> None:
    text = _fn("app/services/montage_coverage_ops.py", "delete_coverage_child")
    assert "remember_deleted(project, int(frame.number))" in text


def test_patch_project_pushes_display_fields() -> None:
    text = _fn("app/web/routers/projects.py", "patch_project")
    assert "push_display_fields_to_master" in text


def test_upload_xlsx_forgets_tombstones() -> None:
    text = _fn("app/web/routers/project_ops.py", "upload_xlsx")
    assert "forget_deleted(p)" in text


def test_gpt_operator_patch_syncs_project_row() -> None:
    text = _fn("app/web/routers/project_ops.py", "gpt_operator_patch")
    assert "sync_project_row_to_master_db" in text


def test_project_asr_words_reads_master_scope() -> None:
    text = _fn("app/web/routers/project_ops.py", "project_asr_words")
    assert "session_scope" in text


def test_montage_board_delete_frame_refuses_when_busy() -> None:
    text = _fn("app/web/routers/project_ops.py", "montage_board_delete_frame")
    assert "_refuse_if_project_busy(p)" in text


def test_load_frame_prompt_rows_uses_project_db_path() -> None:
    text = _fn("app/services/agent_harness.py", "_load_frame_prompt_rows")
    assert "_db_path(data_dir=data_dir, project_id=project_id)" in text
    assert "_frames_db_path" not in text


def test_enrich_run_asks_allows_qc_noop() -> None:
    text = _fn("app/orchestrator/steps/enrich_xlsx.py", "run")
    assert "allows_qc_noop(" in text


def test_remember_deleted_roundtrip() -> None:
    from types import SimpleNamespace

    from app.services.frame_tombstones import deleted_numbers, forget_deleted, remember_deleted

    project = SimpleNamespace(meta={})
    remember_deleted(project, 7)
    assert deleted_numbers(project) == {7}
    forget_deleted(project)
    assert deleted_numbers(project) == set()


# --- поведенческие проверки (мутация вызова больше не проходит) ---

import io
from contextlib import asynccontextmanager

import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.models import AsrWord, Base, Frame, Project, ProjectStatus
from app.services import db_v2
from app.settings import settings


async def _open_db(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    engine = create_async_engine(f"sqlite+aiosqlite:///{path}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
        await db_v2.migrate_db_v2_schema(conn)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    return engine, factory


def _scope_of(factory):
    @asynccontextmanager
    async def _scope():
        async with factory() as session:
            yield session

    return _scope


def _v8_bytes(voices: list[str]) -> bytes:
    from openpyxl import Workbook

    from app.services.xlsx_v8_import import ROW_VOICEOVER_V8, SHEET_GENERAL_V8, SHEET_PLAN_V8

    wb = Workbook()
    plan = wb.active
    plan.title = SHEET_PLAN_V8
    general = wb.create_sheet(SHEET_GENERAL_V8)
    general["A1"] = "pad" * 800
    for index, text in enumerate(voices):
        plan.cell(row=ROW_VOICEOVER_V8, column=3 + index, value=text)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    assert len(data) >= 1024
    return data


def _parent(project_id: int, number: int, uid: str, vo: str) -> Frame:
    return Frame(
        project_id=project_id,
        number=number,
        uuid=uid,
        voiceover_text=vo,
        sort_key=float(number * 10),
        status="planned",
        attrs={"camera_subdivide": {"role": "vo_parent", "parent_uuid": uid}},
    )


def _child(project_id: int, number: int, parent_uid: str, vo: str) -> Frame:
    return Frame(
        project_id=project_id,
        number=number,
        uuid=f"c{number:02d}" + "c" * 20,
        voiceover_text=vo,
        sort_key=float(number * 10),
        status="planned",
        attrs={
            "camera_subdivide": {
                "role": "shot",
                "parent_uuid": parent_uid,
                "coverage_kind": "child",
            }
        },
    )


async def _frame_numbers(session: AsyncSession, project_id: int) -> list[int]:
    rows = (
        await session.execute(select(Frame.number).where(Frame.project_id == project_id))
    ).scalars().all()
    return sorted(int(n) for n in rows)


async def test_patch_project_pushes_display_fields_to_master(tmp_path, monkeypatch) -> None:
    import app.db as app_db
    import app.project_db as project_db
    from app.web.routers.projects import patch_project

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    master_engine, master_factory = await _open_db(tmp_path / "master.db")
    proj_engine, proj_factory = await _open_db(tmp_path / "videos" / "slug" / "project.db")
    monkeypatch.setattr(app_db, "session_scope", _scope_of(master_factory))
    calls: list[list[str]] = []
    real_push = project_db.push_display_fields_to_master

    async def _counting_push(session, project, keys):
        calls.append(list(keys))
        return await real_push(session, project, keys)

    monkeypatch.setattr(project_db, "push_display_fields_to_master", _counting_push)
    async with master_factory() as master:
        master.add(
            Project(id=7, slug="slug", title="старое", topic="т", status=ProjectStatus.new, meta={})
        )
        await master.commit()
    async with proj_factory() as session:
        session.add(
            Project(id=7, slug="slug", title="старое", topic="т", status=ProjectStatus.new, meta={})
        )
        await session.commit()
        updated = await patch_project(7, {"title": "новое"}, session)
        assert updated.title == "новое"
    assert len(calls) == 1
    assert "title" in calls[0]
    async with master_factory() as master:
        row = await master.get(Project, 7)
        assert row is not None and row.title == "новое"
    await master_engine.dispose()
    await proj_engine.dispose()


async def test_delete_montage_frame_tombstones_both_paths(tmp_path, monkeypatch) -> None:
    from app.services.frame_tombstones import deleted_numbers
    from app.services.montage_board_frames import delete_montage_frame
    from app.services.xlsx_v8_import import import_v8_xlsx

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    engine, factory = await _open_db(tmp_path / "mf.db")
    xlsx = tmp_path / "resync.xlsx"
    xlsx.write_bytes(_v8_bytes(["ячейка", "кусок"]))
    async with factory() as session:
        project = Project(slug="montage-frames", topic="t", status=ProjectStatus.frames_ready, meta={})
        session.add(project)
        await session.flush()
        parent_uid = "aa" * 12
        parent = _parent(project.id, 1, parent_uid, "ячейка")
        child = _child(project.id, 2, parent_uid, "кусок")
        session.add_all([parent, child])
        await session.flush()
        child_id = int(child.id)
        parent_id = int(parent.id)

        child_result = await delete_montage_frame(session, project, child_id)
        assert child_result["deleted"] == 1
        assert 2 in deleted_numbers(project)

        parent_result = await delete_montage_frame(session, project, parent_id)
        assert parent_result["deleted"] == 1
        assert deleted_numbers(project) == {1, 2}
        await session.commit()

        await import_v8_xlsx(session, project, xlsx, keep_fields=True)
        await session.flush()
        assert await _frame_numbers(session, project.id) == []
    await engine.dispose()


async def test_delete_coverage_child_records_tombstone(tmp_path, monkeypatch) -> None:
    from app.services.frame_tombstones import deleted_numbers
    from app.services.montage_coverage_ops import delete_coverage_child

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    engine, factory = await _open_db(tmp_path / "cov.db")
    async with factory() as session:
        project = Project(slug="cov", topic="t", status=ProjectStatus.frames_ready, meta={})
        session.add(project)
        await session.flush()
        parent_uid = "bb" * 12
        parent = _parent(project.id, 4, parent_uid, "родитель")
        child = _child(project.id, 5, parent_uid, "дочь")
        session.add_all([parent, child])
        await session.flush()
        frames = [parent, child]
        await delete_coverage_child(session, project, child, frames)
        await session.flush()
        assert 5 in deleted_numbers(project)
        assert await _frame_numbers(session, project.id) == [4]
    await engine.dispose()


async def test_upload_xlsx_clears_tombstones(tmp_path, monkeypatch) -> None:
    from fastapi import UploadFile

    from app.services.frame_tombstones import deleted_numbers, remember_deleted
    from app.web.routers.project_ops import upload_xlsx

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    engine, factory = await _open_db(tmp_path / "up.db")
    async with factory() as session:
        project = Project(slug="upload-me", topic="t", status=ProjectStatus.frames_ready, meta={})
        session.add(project)
        await session.flush()
        remember_deleted(project, 1)
        await session.commit()
        assert deleted_numbers(project) == {1}
        assert await _frame_numbers(session, project.id) == []

        payload = _v8_bytes(["вернулся"])
        uploaded = UploadFile(filename="project.xlsx", file=io.BytesIO(payload))
        await upload_xlsx(project.id, uploaded, node_key=None, session=session)
        await session.refresh(project)
        assert deleted_numbers(project) == set()
        assert await _frame_numbers(session, project.id) == [1]
    await engine.dispose()


async def test_gpt_operator_patch_syncs_master(tmp_path, monkeypatch) -> None:
    import app.db as app_db
    from app.web.routers import project_ops

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    master_engine, master_factory = await _open_db(tmp_path / "state.db")
    proj_engine, proj_factory = await _open_db(tmp_path / "videos" / "op" / "project.db")
    monkeypatch.setattr(app_db, "session_scope", _scope_of(master_factory))
    sync_calls: list[int] = []
    real_sync = project_ops.sync_project_row_to_master_db

    async def _counting_sync(project):
        sync_calls.append(int(project.id))
        return await real_sync(project)

    monkeypatch.setattr(project_ops, "sync_project_row_to_master_db", _counting_sync)
    node = "n_excel"
    async with master_factory() as master:
        master.add(Project(id=3, slug="op", topic="t", status=ProjectStatus.new, meta={}))
        await master.commit()
    async with proj_factory() as session:
        session.add(Project(id=3, slug="op", topic="t", status=ProjectStatus.new, meta={}))
        await session.commit()
        result = await project_ops.gpt_operator_patch(
            3, node, {"role": "review", "outputMode": "text"}, session
        )
        assert result["ok"] is True
    assert sync_calls == [3]
    async with master_factory() as master:
        row = await master.get(Project, 3)
        cfg = (row.meta or {}).get("excel_gpt_nodes", {}).get(node, {})
        assert cfg.get("role") == "review"
    await master_engine.dispose()
    await proj_engine.dispose()


async def test_montage_delete_frame_refused_while_busy(tmp_path, monkeypatch) -> None:
    from fastapi import HTTPException

    from app.web.routers.project_ops import montage_board_delete_frame

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    engine, factory = await _open_db(tmp_path / "busy.db")
    async with factory() as session:
        project = Project(
            slug="busy",
            topic="t",
            status=ProjectStatus.generating_images,
            meta={},
        )
        session.add(project)
        frame = _parent(0, 1, "dd" * 12, "на месте")
        session.add(frame)
        await session.flush()
        frame.project_id = project.id
        await session.commit()
        frame_id = int(frame.id)
        with pytest.raises(HTTPException) as caught:
            await montage_board_delete_frame(project.id, frame_id, session)
        assert caught.value.status_code == 409
        assert await session.get(Frame, frame_id) is not None
    await engine.dispose()


async def test_project_asr_words_uses_own_session(tmp_path, monkeypatch) -> None:
    import app.db as app_db
    from app.web.routers.project_ops import project_asr_words

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    master_engine, master_factory = await _open_db(tmp_path / "state.db")
    proj_engine, proj_factory = await _open_db(tmp_path / "videos" / "asr" / "project.db")
    opened: list[str] = []
    real_scope = _scope_of(master_factory)

    @asynccontextmanager
    async def _counting_scope():
        opened.append("master")
        async with real_scope() as session:
            yield session

    monkeypatch.setattr(app_db, "session_scope", _counting_scope)
    async with master_factory() as master:
        master.add(Project(id=11, slug="asr", topic="t", status=ProjectStatus.new, meta={}))
        master.add(Project(id=12, slug="other", topic="t", status=ProjectStatus.new, meta={}))
        master.add(
            AsrWord(
                project_id=11,
                run_uuid="run-own",
                idx=0,
                word="своё",
                start_s=0.0,
                end_s=0.4,
                backend="whisper",
            )
        )
        master.add(
            AsrWord(
                project_id=12,
                run_uuid="run-other",
                idx=0,
                word="чужое",
                start_s=0.0,
                end_s=0.4,
                backend="whisper",
            )
        )
        await master.commit()
    async with proj_factory() as session:
        session.add(Project(id=11, slug="asr", topic="t", status=ProjectStatus.new, meta={}))
        await session.commit()
        payload = await project_asr_words(11, session)
    assert opened == ["master"]
    assert [row["word"] for row in payload["words"]] == ["своё"]
    assert payload["count"] == 1
    await master_engine.dispose()
    await proj_engine.dispose()


async def test_load_frame_prompt_rows_reads_project_db(tmp_path, monkeypatch) -> None:
    from app.services.agent_harness import _load_frame_prompt_rows

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    master = tmp_path / "state.db"
    monkeypatch.setattr(settings, "sqlite_path", master)
    dirs = {
        1: tmp_path / "videos" / "alpha",
        2: tmp_path / "videos" / "beta",
    }
    prompts = {1: "только-альфа", 2: "только-бета"}
    master_engine, master_factory = await _open_db(master)
    async with master_factory() as session:
        for pid, text in ((1, "мастер-альфа"), (2, "мастер-бета")):
            session.add(Project(id=pid, slug=f"p{pid}", topic="t", status=ProjectStatus.new))
            session.add(
                Frame(
                    project_id=pid,
                    number=1,
                    voiceover_text=text,
                    image_prompt=text,
                    animation_prompt="",
                    uuid=f"m{pid}",
                )
            )
        await session.commit()
    engines = []
    for pid, folder in dirs.items():
        folder.mkdir(parents=True)
        engine, factory = await _open_db(folder / "project.db")
        engines.append(engine)
        async with factory() as session:
            session.add(Project(id=pid, slug=f"p{pid}", topic="t", status=ProjectStatus.new))
            session.add(
                Frame(
                    project_id=pid,
                    number=1,
                    voiceover_text=prompts[pid],
                    image_prompt=prompts[pid],
                    animation_prompt="anim",
                    uuid=f"own{pid}",
                )
            )
            other = 2 if pid == 1 else 1
            session.add(
                Frame(
                    project_id=other,
                    number=9,
                    voiceover_text="чужой-ряд",
                    image_prompt="чужой-ряд",
                    animation_prompt="",
                    uuid=f"leak{pid}",
                )
            )
            await session.commit()

    for pid, folder in dirs.items():
        rows, err = _load_frame_prompt_rows(pid, data_dir=folder)
        assert err == ""
        assert [row[0] for row in rows] == [1]
        assert rows[0][2] == prompts[pid]
        assert "мастер" not in rows[0][2]
        assert "чужой" not in rows[0][2]
    await master_engine.dispose()
    for engine in engines:
        await engine.dispose()


async def test_enrich_qc_noop_accepted_when_allowed(tmp_path, monkeypatch) -> None:
    from app.orchestrator.steps.enrich_xlsx import run as enrich_run
    from app.services.gpt_api import GptChatResult

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr("app.services.gpt_api.gpt_api_enabled", lambda: True)
    calls: list[str] = []

    async def _chat(**kwargs):
        calls.append(kwargs.get("prompt") or "")
        return GptChatResult(
            text='{"ops":[]}',
            model="test",
            finish_reason="stop",
            raw={"choices": [{"finish_reason": "stop"}]},
        )

    monkeypatch.setattr("app.services.gpt_api.chat", _chat)

    async def _gate(*_args, **_kwargs):
        return None

    monkeypatch.setattr("app.services.agent_harness.harness_gate_or_raise", _gate)

    node = "qc_slot"
    engine, factory = await _open_db(tmp_path / "videos" / "qc" / "project.db")
    async with factory() as session:
        project = Project(
            slug="qc",
            topic="тема",
            status=ProjectStatus.enriching_1,
            auto_mode=False,
            meta={
                "active_excel_gpt_node_key": node,
                "excel_gpt_nodes": {
                    node: {
                        "role": "assist",
                        "outputMode": "project_file",
                        "transport": "api",
                        "checkMode": False,
                        "takeFromEdges": False,
                    }
                },
                "prompt_slot_variants": {node: {"main": "shots_qc_ru"}},
                "canvas_graph": {
                    "nodes": [
                        {
                            "id": node,
                            "type": "excel_gpt",
                            "data": {"groupId": "script_frames_qc"},
                        }
                    ],
                    "edges": [],
                },
            },
        )
        session.add(project)
        await session.flush()
        session.add(
            Frame(
                project_id=project.id,
                number=1,
                uuid="frame-qc-1",
                voiceover_text="Диктор открывает дверь и замолкает.",
                status="planned",
            )
        )
        await session.commit()
        await enrich_run(session, project, None)
        await session.refresh(project)
        assert project.status == ProjectStatus.enrich_1_ready
        assert len(calls) == 1
    await engine.dispose()
