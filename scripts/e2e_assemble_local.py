"""Local node `assemble` through the PRODUCTION route: worker -> observation -> diagnosis -> Decision -> ONE supervised
re-run under a ticket -> independent verification. Everything is local and FREE: no provider, no key, no network.

Sibling of scripts/e2e_media_production_route.py (paid Outsee nodes; imported, not edited: its request Gate, flow
helpers and report style are reused). `assemble` has no provider job and no paid request: its supervised recovery is
the only one that RE-RUNS the node (through the original worker, once, inside media_spend.resume_only) and verifies the
ticket from independent facts (app/services/local_node_recovery.py).

The flow:
  1. seed a small synthetic project: 3 frames, 3 tiny local clips (ffmpeg lavfi `color`, 160x284, 2 s each) as
     scene_video rows, 3 per-frame voice files and the full voice (ffmpeg lavfi `sine`, 1.5 s per frame, mp3) as an
     `audio` artifact in per_frame mode;
  2. observe-only scope, UI start of the node (POST /steps/assemble/run): the REAL step assemble.run runs the REAL
     ffmpeg assembly (app.services.assembly.assemble) and the REAL verify_artifact + track check. The forced failure:
     the assembly call is made to fail like ffmpeg ("ffmpeg exited with 1") on every call, so the in-node retry
     (exactly one) and the native retries (3 cycles of 3) run out and the project pauses;
  3. the failure is switched OFF (the "cause is gone"), diagnosis types a media incident whose failed attempt proved
     its OWN repeatable failure, actions are enabled, the kill switch is probed against the recovery, the Decision loop
     is enabled: Decision (local stub) -> retry_same -> ONE ticket -> the original worker re-runs assemble ONCE inside
     resume-only -> a real final.mp4 -> reconcile verifies from independent facts (file newer than the claim,
     verify_artifact again, digest equals the artifact row, approve_final of that file still PENDING, provider ledger
     unchanged, zero refused paid seams, no model request by the attempt, no part file);
  4. afterwards the stand re-reads the final with the application's own verifier (ffprobe: video + audio, duration) and
     checks that the approval is still with the human (pending, never approved by the supervisor).

ASR / subtitles. The per-frame voice route needs no ASR and subtitles are OFF (the default). Whether a Whisper is
available locally is probed WITHOUT network and reported: `faster_whisper` import + the names of the cached models in
the Hugging Face cache directory (nothing is loaded or downloaded). A lavfi sine tone has no words, so a real ASR could
not be exercised on it anyway: ASR and burned subtitles are listed as NOT proven.

REAL SPEECH mode (F07, opt-in, default OFF; the default behaviour above is unchanged): `--real-speech --voice-wav <wav>`.
The seeded project then uses the caller's 16-bit wav (a Windows SAPI voice, NOT a human) as the FULL voice artifact
(mode full_voice: the production assemble needs ASR and builds the timeline of the montage-v3 path from it), the frames
keep tiny local lavfi clips (vertical, long enough), subtitles are ENABLED in the project's assemble params and the REAL
faster-whisper `--asr-model` (offline, cached) runs through app.services.montage_asr (no stub). The stand additionally
proves: >= 8 of the 11 expected words recognised with monotonic timestamps inside the voice, the ASS file / ffmpeg filter
built by the production assemble come from those words, the final passes verify_artifact with the voice length, and the
subtitles are really burned on the picture (lower-third pixel diff against the same render without subtitles, ffmpeg +
numpy). Nothing leaves the machine: a socket guard records and refuses any non-loopback connect / name lookup.

Usage (repository root, project venv; --work-dir must be an empty dir OUTSIDE the repository):
    python scripts/e2e_assemble_local.py --work-dir <dir> --dry
    python scripts/e2e_assemble_local.py --work-dir <dir> --dry --real-speech --voice-wav <voice.wav> --asr-model medium
"""
from __future__ import annotations

import argparse
import asyncio
import difflib
import importlib.util
import ipaddress
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import time
import unicodedata
import uuid as _uuid
import wave
from pathlib import Path
from typing import Any

_HERE = Path(__file__).resolve().parent
ROOT = _HERE.parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
import e2e_guard as guard  # noqa: E402  (A19: refuses to touch live data)


def _load_script(name: str, alias: str):
    spec = importlib.util.spec_from_file_location(alias, _HERE / (name + ".py"))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


base = _load_script("e2e_media_production_route", "e2e_media_production_route_for_assemble")  # imported, not edited
qc = base.qc

PORT = 18808
NODE = "assemble"
STEP_CODE = "assemble"
TARGET = "media-assemble-A"
GROUP_ID = "media-g1"
MEDIA_PROVIDERS = ("outsee", "kie", "elevenlabs", "suno", "grsai")
FRAME_SECONDS = 1.5
CLIP_SECONDS = 2
COLORS = ("red", "green", "blue")
VOICEOVERS = ("Первый кадр: папка лежит на столе.", "Второй кадр: рядом стоит чашка.", "Третий кадр: раскрыт блокнот.")
FAULT = "ffmpeg exited with 1"      # what a failed ffmpeg run looks like to the node (retryable by the node itself)
REQUIRED_TICKET_CHECKS = ("artifact_row", "file_inside_project", "claimed_at_known", "produced_after_claim",
                          "verify_artifact_passed", "digest_matches_artifact", "approve_final_pending",
                          "journal_proof_matches", "ledger_unchanged", "no_paid_seam_hit", "no_model_requests",
                          "no_part_files")

# --- real speech mode (F07) ---
REAL_FRAME_TEXTS = ("Привет.", "Это проверка сборки видео.", "Сегодня мы покажем, как работает субтитр.")
REAL_VOICE_TEXT = " ".join(REAL_FRAME_TEXTS)          # what the SAPI wav says (11 words)
REAL_EXPECTED_WORDS = ("привет", "это", "проверка", "сборки", "видео", "сегодня", "мы", "покажем", "как", "работает", "субтитр")
REAL_MIN_WORDS = 8
REAL_COLORS = ("navy", "darkgreen", "maroon")        # dark flat colours: white subtitles with a black outline show clearly
REAL_CLIP_SIZE = "540x960"                           # vertical, big enough for the ASS font (PlayRes = clip size)
REAL_CLIP_SECONDS = 5
REAL_DURATION_TOLERANCE_S = 0.5
REAL_SUB_TOLERANCE_S = 0.6                           # a cue starts within this of the ASR word it was aligned to
LOWER_BAND = (0.55, 0.88)                            # rows (share of the height) where the subtitles are placed
UPPER_BAND = (0.05, 0.35)                            # rows that must NOT change (noise baseline of the re-encode)
PIXEL_THRESHOLD = 40                                 # a pixel "differs" when the gray values differ by more than this
MIN_LOWER_FRACTION = 0.01                            # >= 1% of the lower band must differ where a word is spoken
MAX_NOISE_FRACTION = 0.002                           # <= 0.2% may differ in the untouched upper band / in silence
DEFAULT_ASR_MODEL = "medium"


# ───────────────────────────── local media made by ffmpeg (lavfi), free ─────────────────────────────

def _ffmpeg() -> str:
    exe = shutil.which("ffmpeg")
    if not exe:
        raise RuntimeError("ffmpeg is required for the assemble stand (not on PATH)")
    return exe


def _run(cmd: list[str]) -> None:
    subprocess.run(cmd, check=True, timeout=120, capture_output=True)


def make_clip(path: Path, color: str, seconds: int = CLIP_SECONDS, size: str = "160x284") -> None:
    _run([_ffmpeg(), "-y", "-v", "error", "-f", "lavfi", "-i", f"color=c={color}:s={size}:r=24:d={seconds}",
          "-c:v", "libx264", "-pix_fmt", "yuv420p", "-preset", "veryfast", str(path)])


def make_voice(path: Path, frequency: int, seconds: float = FRAME_SECONDS) -> None:
    _run([_ffmpeg(), "-y", "-v", "error", "-f", "lavfi", "-i", f"sine=frequency={frequency}:duration={seconds}:sample_rate=44100",
          "-ac", "1", "-c:a", "libmp3lame", "-b:a", "64k", str(path)])


def concat_voice(parts: list[Path], out: Path) -> None:
    listing = out.parent / "voice_parts.txt"
    listing.write_text("\n".join(f"file '{p.as_posix()}'" for p in parts), encoding="utf-8")
    try:
        _run([_ffmpeg(), "-y", "-v", "error", "-f", "concat", "-safe", "0", "-i", str(listing), "-c", "copy", str(out)])
    finally:
        listing.unlink(missing_ok=True)


def asr_probe(real_speech: bool = False) -> dict:
    """Is a local Whisper available? Imports/lists only: nothing is loaded, nothing is downloaded, no network."""
    import importlib.util as iu
    hub = Path(os.environ.get("HF_HOME") or (Path.home() / ".cache" / "huggingface")) / "hub"
    try:
        cached = sorted(p.name for p in hub.glob("models--*") if "whisper" in p.name.lower())
    except OSError:
        cached = []
    return {"faster_whisper_installed": iu.find_spec("faster_whisper") is not None, "cached_whisper_models": cached,
            "used_by_this_stand": bool(real_speech),
            "why": ("--real-speech: full-voice route, the REAL faster-whisper runs offline on the SAPI wav, subtitles are on"
                    if real_speech else "per-frame voice route needs no ASR; subtitles are off; a sine tone has no words")}



# ───────────────────────────── real speech: pure helpers (unit-tested) ─────────────────────────────

def normalize_word(text: str) -> str:
    """Case, punctuation and yo/ye insensitive form of a word."""
    t = unicodedata.normalize("NFKC", str(text or "")).lower().replace("ё", "е")
    return "".join(ch for ch in t if ch.isalnum())


def words_equal(a: str, b: str) -> bool:
    """Equal after normalisation, or (for words of 5+ letters) a near-miss such as an ending (субтитр / субтитры)."""
    if not a or not b:
        return False
    if a == b:
        return True
    return min(len(a), len(b)) >= 5 and difflib.SequenceMatcher(None, a, b).ratio() >= 0.8


def align_words(left: list[str], right: list[str]) -> list[tuple[int, int]]:
    """In-order (LCS) pairs of indexes whose normalised words are equal; order is never crossed."""
    a, b = [normalize_word(x) for x in left], [normalize_word(x) for x in right]
    n, m = len(a), len(b)
    best = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            best[i][j] = best[i + 1][j + 1] + 1 if words_equal(a[i], b[j]) else max(best[i + 1][j], best[i][j + 1])
    pairs, i, j = [], 0, 0
    while i < n and j < m:
        if words_equal(a[i], b[j]) and best[i][j] == best[i + 1][j + 1] + 1:
            pairs.append((i, j))
            i, j = i + 1, j + 1
        elif best[i + 1][j] >= best[i][j + 1]:
            i += 1
        else:
            j += 1
    return pairs


def check_asr_words(words: list[dict], voice_s: float, expected=REAL_EXPECTED_WORDS, *, min_words: int = REAL_MIN_WORDS) -> dict:
    """The ASR result: how many expected words came back (in order), and whether the stamps are sane."""
    texts = [str(w.get("word", "")) for w in words]
    pairs = align_words(list(expected), texts)
    starts = [float(w.get("start", 0.0)) for w in words]
    ends = [float(w.get("end", 0.0)) for w in words]
    monotonic = all(e >= s - 1e-6 for s, e in zip(starts, ends)) and all(b >= a - 1e-6 for a, b in zip(starts, starts[1:]))
    inside = bool(words) and min(starts) >= -1e-6 and max(ends) <= float(voice_s) + 0.35
    return {"recognised": texts, "expected_count": len(expected), "matched": len(pairs), "min_words": min_words,
            "matched_expected": [expected[i] for i, _ in pairs], "monotonic": bool(monotonic), "inside_voice": bool(inside),
            "first_start": round(min(starts), 3) if starts else None, "last_end": round(max(ends), 3) if ends else None,
            "voice_s": round(float(voice_s), 3),
            "ok": len(pairs) >= min_words and bool(monotonic) and bool(inside)}


def _ass_seconds(text: str) -> float:
    h, m, s = str(text).strip().split(":")
    return int(h) * 3600 + int(m) * 60 + float(s)


def parse_ass_cues(text: str) -> list[tuple[float, float, str]]:
    """(start, end, text) of every Dialogue line of an ASS file; override tags are removed."""
    cues = []
    for line in str(text).splitlines():
        if not line.startswith("Dialogue:"):
            continue
        parts = line[len("Dialogue:"):].split(",", 9)
        if len(parts) < 10:
            continue
        body = re.sub(r"\{[^}]*\}", "", parts[9]).replace("\\N", " ").strip()
        cues.append((_ass_seconds(parts[1]), _ass_seconds(parts[2]), body))
    return cues


def check_subtitles_from_words(cues: list[tuple[float, float, str]], asr_words: list[dict],
                               script_words=REAL_EXPECTED_WORDS, *, tolerance: float = REAL_SUB_TOLERANCE_S) -> dict:
    """The cues of the production ASS file: the displayed words are the script words, and each cue starts where the ASR
    heard its FIRST word (a cue of two words shows both from the first word's start, so only first words are timed)."""
    flat = [(float(s), tok, k == 0) for s, _e, txt in cues for k, tok in enumerate(t for t in txt.split() if normalize_word(t))]
    shown = [normalize_word(t) for _s, t, _f in flat]
    script_ok = shown == [normalize_word(w) for w in script_words]
    pairs = align_words([t for _s, t, _f in flat], [str(w.get("word", "")) for w in asr_words])
    first = [(i, j) for i, j in pairs if flat[i][2]]
    offsets = [round(flat[i][0] - float(asr_words[j].get("start", 0.0)), 3) for i, j in first]
    within = sum(1 for o in offsets if abs(o) <= tolerance)
    return {"cue_count": len(cues), "cue_words": [t for _s, t, _f in flat], "script_words_shown": script_ok,
            "aligned_to_asr": len(pairs), "cue_first_words_aligned": len(first), "offsets_s": offsets,
            "within_tolerance": within, "tolerance_s": tolerance,
            "max_abs_offset_s": max((abs(o) for o in offsets), default=None),
            "ok": bool(cues) and script_ok and len(pairs) >= REAL_MIN_WORDS and len(first) >= max(1, len(cues) - 1)
            and within >= max(1, int(0.8 * len(first)))}


def pick_probe_times(cues: list[tuple[float, float, str]], asr_words: list[dict], voice_s: float) -> dict:
    """A time inside the longest spoken word with a cue on screen, and (if any) a silent time with no cue."""
    spoken = [w for w in asr_words if float(w.get("end", 0)) - float(w.get("start", 0)) > 0.05]

    def in_cue(t: float) -> bool:
        return any(s <= t <= e for s, e, _ in cues)

    word_t = None
    for w in sorted(spoken, key=lambda w: float(w["start"]) - float(w["end"])):
        t = round((float(w["start"]) + float(w["end"])) / 2, 3)
        if in_cue(t):
            word_t = {"t": t, "word": w.get("word")}
            break
    spans = sorted((float(w["start"]), float(w["end"])) for w in spoken)
    gaps, cursor = [], 0.0
    for s0, e0 in spans:
        gaps.append((cursor, s0))
        cursor = max(cursor, e0)
    gaps.append((cursor, float(voice_s) - 0.05))
    silent_t = None
    for g0, g1 in gaps:
        if g1 - g0 >= 0.3:
            t = round((g0 + g1) / 2, 3)
            if not in_cue(t):
                silent_t = t
                break
    return {"word": word_t, "silent": silent_t}


def band_difference(a, b, *, band: tuple[float, float], threshold: int = PIXEL_THRESHOLD) -> dict:
    """Two gray frames (numpy 2-D arrays): share of pixels in the row band that differ by more than the threshold."""
    import numpy as np
    h = a.shape[0]
    r0, r1 = int(h * band[0]), int(h * band[1])
    diff = np.abs(a[r0:r1].astype(np.int16) - b[r0:r1].astype(np.int16))
    return {"rows": [r0, r1], "fraction": round(float((diff > threshold).mean()), 5), "mean_abs": round(float(diff.mean()), 3)}


def check_subtitle_pixels(word_frame: dict | None, silent_frame: dict | None) -> dict:
    """word_frame / silent_frame: {"lower": band_difference, "upper": band_difference} of final vs the no-subtitles control."""
    if not word_frame:
        return {"ok": False, "reason": "no frame inside a spoken word"}
    lower, upper = word_frame["lower"]["fraction"], word_frame["upper"]["fraction"]
    silent_ok = True if not silent_frame else silent_frame["lower"]["fraction"] <= MAX_NOISE_FRACTION
    return {"ok": lower >= MIN_LOWER_FRACTION and upper <= MAX_NOISE_FRACTION and silent_ok,
            "lower_fraction_at_word": lower, "upper_fraction_at_word": upper,
            "lower_fraction_in_silence": None if not silent_frame else silent_frame["lower"]["fraction"],
            "silence_checked": bool(silent_frame),
            "thresholds": {"pixel": PIXEL_THRESHOLD, "min_lower_fraction": MIN_LOWER_FRACTION,
                           "max_noise_fraction": MAX_NOISE_FRACTION}}


def is_loopback_host(host) -> bool:
    if host is None:
        return True
    name = str(host).strip().strip("[]").lower()
    if name in ("", "localhost", "localhost.localdomain"):
        return True
    try:
        return ipaddress.ip_address(name.split("%")[0]).is_loopback
    except ValueError:
        return False


def install_network_guard(attempts: list) -> Any:
    """Record and refuse every non-loopback connect / name lookup in this process. Returns the restore function.

    Covers socket.connect/connect_ex, socket.getaddrinfo and the asyncio loop's sock_connect (the Windows proactor loop
    connects with ConnectEx, not with socket.connect). Best effort, honest limit: a C extension opening its own socket
    is not seen; nothing here configures any provider or key either."""
    try:
        own_name = socket.gethostname().lower()
    except OSError:
        own_name = ""
    originals = {"connect": socket.socket.connect, "connect_ex": socket.socket.connect_ex, "getaddrinfo": socket.getaddrinfo}

    def allowed(host) -> bool:
        return is_loopback_host(host) or bool(own_name and str(host).lower() == own_name)

    def refuse(kind, host):
        attempts.append({"kind": kind, "host": str(host)})
        return OSError(f"network guard: {kind} to non-loopback {host!r} refused")

    def connect(self, address):
        if isinstance(address, tuple) and address and not allowed(address[0]):
            raise refuse("connect", address[0])
        return originals["connect"](self, address)

    def connect_ex(self, address):
        if isinstance(address, tuple) and address and not allowed(address[0]):
            attempts.append({"kind": "connect_ex", "host": str(address[0])})
            return 10061
        return originals["connect_ex"](self, address)

    def getaddrinfo(host, *a, **k):
        if not allowed(host):
            attempts.append({"kind": "getaddrinfo", "host": str(host)})
            raise socket.gaierror(f"network guard: lookup of {host!r} refused")
        return originals["getaddrinfo"](host, *a, **k)

    socket.socket.connect, socket.socket.connect_ex, socket.getaddrinfo = connect, connect_ex, getaddrinfo
    patched_loops: list[tuple[type, Any]] = []
    try:
        loop_cls = type(asyncio.get_running_loop())
        orig_sock_connect = loop_cls.sock_connect

        async def sock_connect(self, sock, address):
            if isinstance(address, tuple) and address and not allowed(address[0]):
                raise refuse("sock_connect", address[0])
            return await orig_sock_connect(self, sock, address)

        loop_cls.sock_connect = sock_connect
        patched_loops.append((loop_cls, orig_sock_connect))
    except RuntimeError:
        pass

    def restore() -> None:
        socket.socket.connect, socket.socket.connect_ex, socket.getaddrinfo = (
            originals["connect"], originals["connect_ex"], originals["getaddrinfo"])
        for cls, fn in patched_loops:
            cls.sock_connect = fn

    return restore


def validate_voice_wav(path: Path) -> dict:
    """The caller's wav must be 16-bit PCM; returns its facts. Raises ValueError otherwise."""
    try:
        with wave.open(str(path), "rb") as w:
            info = {"channels": w.getnchannels(), "sample_rate": w.getframerate(), "sample_width": w.getsampwidth(),
                    "duration": round(w.getnframes() / float(w.getframerate() or 1), 3)}
    except (wave.Error, EOFError, OSError) as exc:
        raise ValueError(f"voice wav is not readable PCM: {exc}") from exc
    if info["sample_width"] != 2:
        raise ValueError(f"voice wav must be 16-bit, got {8 * info['sample_width']}-bit")
    if info["duration"] < 3.0:
        raise ValueError("voice wav is too short")
    return info


def wav_to_mp3(wav: Path, mp3: Path) -> None:
    """The voice artifact the pipeline expects (mono mp3 like voice_full.mp3)."""
    _run([_ffmpeg(), "-y", "-v", "error", "-i", str(wav), "-ac", "1", "-ar", "44100", "-c:a", "libmp3lame",
          "-b:a", "128k", str(mp3)])


def make_plan_xlsx(path: Path, texts=REAL_FRAME_TEXTS) -> None:
    """project.xlsx with the sheet «план»: R49 (voiceover) of frame N in column N+2 - what the full-voice path reads."""
    from openpyxl import Workbook
    from app.services.xlsx_v8_import import ROW_VOICEOVER_V8
    wb = Workbook()
    ws = wb.active
    ws.title = "план"
    for i, text in enumerate(texts, start=1):
        ws.cell(row=ROW_VOICEOVER_V8, column=i + 2, value=text)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def read_r15_labels(xlsx: Path, count: int) -> list[str]:
    from openpyxl import load_workbook
    from app.services.xlsx_v8_import import ROW_TIMECODE_V8
    wb = load_workbook(xlsx, data_only=True)
    try:
        ws = wb["план"]
        return [str(ws.cell(row=ROW_TIMECODE_V8, column=i + 2).value or "") for i in range(1, count + 1)]
    finally:
        wb.close()


def video_size(path: Path) -> tuple[int, int]:
    out = subprocess.run(["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=width,height",
                          "-of", "csv=p=0:s=x", str(path)], check=True, capture_output=True, text=True, timeout=60).stdout.strip()
    w, h = out.split("x")[:2]
    return int(w), int(h)


def gray_frame(path: Path, t: float):
    """One gray frame at time t as a numpy 2-D array (ffmpeg only; no OCR)."""
    import numpy as np
    w, h = video_size(path)
    raw = subprocess.run([_ffmpeg(), "-v", "error", "-ss", f"{t:.3f}", "-i", str(path), "-frames:v", "1",
                          "-vf", "format=gray", "-f", "rawvideo", "-"], check=True, capture_output=True, timeout=120).stdout
    if len(raw) < w * h:
        raise RuntimeError(f"no frame at {t}s in {path.name}")
    return np.frombuffer(raw[: w * h], dtype=np.uint8).reshape(h, w)


def frame_pair_difference(final: Path, control: Path, t: float) -> dict:
    a, b = gray_frame(final, t), gray_frame(control, t)
    if a.shape != b.shape:
        raise RuntimeError(f"final and control differ in size: {a.shape} vs {b.shape}")
    return {"t": t, "lower": band_difference(a, b, band=LOWER_BAND), "upper": band_difference(a, b, band=UPPER_BAND)}

# ───────────────────────────── checks (pure, unit-tested) ─────────────────────────────

def check(name: str, ok: bool, detail: Any = None, *, needs: str | None = None) -> dict:
    return qc.check(name, ok, detail, needs=needs)


def verdict(checks: list[dict]) -> tuple[bool, list[str]]:
    return qc.verdict(checks)


def _blocked(probes: dict, name: str, code: str) -> bool:
    return base._blocked(probes, name, code)


def reality_report(real_speech: bool = False) -> dict:
    if real_speech:
        return _reality_report_real_speech()
    return {
        "real": [
            "application, worker, observer, diagnosis, Decision supervisor, recovery ticket and verification in-process",
            "native step assemble.run: recover_before_assemble, the per-frame timeline, the REAL ffmpeg assembly "
            "(app.services.assembly.assemble: cut, concat, mux), verify_artifact (ffprobe) and the track check, the "
            "in-node retry, the atomic write, approve_final creation",
            "local_node_recovery (admission, resume-only scope, reconcile verification), media_spend.resume_only, spend_guard, kill switch",
            "ffprobe check of the final by artifact_verifier.verify_artifact after the run"],
        "not_real": [
            "seeded project: synthetic, built by this script (3 frames, 3 lavfi clips, a lavfi sine voice)",
            "the forced failure: injected locally (the assembly call raises 'ffmpeg exited with 1' until the stand switches it off)",
            "Decision model: local stub (nothing was sent to a provider)",
            "Telegram: off (the approve_final request is created in the database, no message was sent)",
            "ASR / Whisper and burned subtitles: not used (per-frame voice route, subtitles off)",
            "picture, voice and montage quality: not judged"]}


def _reality_report_real_speech() -> dict:
    base_report = reality_report()
    return {
        "real": [
            "application, worker, observer, diagnosis, Decision supervisor, recovery ticket and verification in-process",
            "native step assemble.run on the FULL-VOICE path: ASR (app.services.montage_asr -> faster-whisper, offline, "
            "no stub), R15 timeline written from the ASR words, montage-v3 slot encode + mux (real ffmpeg), the ASS file "
            "built by build_subtitle_cues_from_cells/make_simple_ass and burned with subtitles_vf_arg, verify_artifact, "
            "the track check, the in-node retry, the atomic write, approve_final creation",
            "local_node_recovery (admission, resume-only scope, reconcile verification), media_spend.resume_only, spend_guard, kill switch",
            "ffprobe check of the final by artifact_verifier.verify_artifact after the run",
            "burned subtitles measured on the picture: ffmpeg frame vs the same render without subtitles (numpy pixel diff)"],
        "not_real": [
            "the speech: a Windows SAPI synthetic Russian voice, NOT a human voice (no breath, accent, noise, room)",
            "the clips: synthetic flat-colour lavfi clips, not generated or shot footage",
            "seeded project: 3 frames, the plan sheet and the frame texts are written by this script",
            *[x for x in base_report["not_real"] if x.startswith(("the forced failure", "Decision model", "Telegram"))],
            "picture, voice and montage quality: not judged"]}


def reality_not_proven(real_speech: bool = False) -> list[str]:
    if real_speech:
        return ["ASR quality on real human voices, noise, accents, music under the voice, other languages or long files "
                "(one 8.7 s SAPI voice only; the words were known in advance)",
                "legibility and design of the subtitles on real footage (flat colour background, one font fallback of libass)",
                "assembly of real Outsee clips (a different size / codec / audio track than tiny lavfi clips)",
                "ASR on the GPU or with other models than the one named in the run",
                "approval: stays with the human by design; this stand never approves"]
    return ["ASR / Whisper timing and burned subtitles on the assembly path (no local faster_whisper in the venv unless "
            "asr_probe says otherwise; a sine tone has no words)",
            "assembly of real Outsee clips (a different size / codec / audio track than tiny lavfi clips)",
            "the montage-v3 (excel-r15) path: this stand uses the per-frame voice path",
            "approval: stays with the human by design; this stand never approves"]


def evaluate(snap: dict) -> list[dict]:
    counters = snap.get("counters", {})
    probes = snap.get("probes", {})
    incidents = [i for i in (snap.get("diagnoses") or {}).get("incidents", []) if i.get("node_key") == TARGET]
    media_inc = [i for i in incidents if i.get("type") == "media_failure"]
    records = (snap.get("decisions") or {}).get("records") or []
    recoveries = (snap.get("actions") or {}).get("recoveries") or []
    ticket = recoveries[0] if recoveries else {}
    precond = ticket.get("preconditions") or {}
    ver = ticket.get("verification_result") or {}
    ctxs = snap.get("decision_contexts") or []
    media_ctx = next((c.get("media_recovery") for c in ctxs if c.get("media_recovery")), None) or {}
    node_status = (snap.get("node_status") or {}).get(TARGET)
    attempts = int(snap.get("node_attempts") or 0)
    stepf = snap.get("step_failure") or {}
    before = snap.get("state_at_exhaustion") or {}
    final = snap.get("final_check") or {}
    approvals = snap.get("approvals") or []
    mine = [a for a in approvals if final.get("sha256") and a.get("sha256") == final.get("sha256")]
    real = snap.get("real_speech_mode") is True
    voice_s = float(((snap.get("real_speech") or {}).get("voice") or {}).get("duration") or 0.0)
    duration_window = ((voice_s - REAL_DURATION_TOLERANCE_S, voice_s + REAL_DURATION_TOLERANCE_S)
                       if real and voice_s else (3.5, 6.0))
    out = [
        check("failure_injected_and_native_retries_ran", counters.get("injected_failures", 0) >= 4 and attempts >= 3,
              {"injected": counters.get("injected_failures"), "node_attempts": attempts}),
        check("in_node_retry_is_exactly_one", counters.get("injected_failures", 0) >= 2
              and counters.get("injected_failures", 0) % 2 == 0, {"injected": counters.get("injected_failures")}),
        check("native_retries_exhausted_and_project_paused", bool(stepf.get("abandoned_at")) and before.get("node") == "failed",
              {"step_failure": {k: stepf.get(k) for k in ("total_fails", "abandoned_at", "abandoned_threshold")},
               "node_at_exhaustion": before.get("node")}),
        check("nothing_was_published_by_the_failed_phase", before.get("final_files") == [] and before.get("approvals") == 0
              and before.get("part_files") == [], before),
        check("media_incident_typed_by_diagnosis", bool(media_inc) and any(i.get("media") is True for i in media_inc),
              [(i.get("type"), i.get("status")) for i in incidents][-6:]),
        check("decision_sees_the_local_worker_rerun_as_the_only_offer",
              media_ctx.get("operation") == "retry_local_worker" and media_ctx.get("submits_new_paid_request") is False,
              media_ctx, needs="decision_v1._media_info and local_node_recovery.admit"),
        check("decision_done_with_retry_same", bool(records) and records[0].get("status") == "done"
              and precond.get("action_type") == "retry_same",
              {"decision": [(r.get("status"), r.get("model_steps")) for r in records], "action": precond.get("action_type")}),
        check("exactly_one_ticket_and_it_is_verified", len(recoveries) == 1 and ticket.get("status") == "verified"
              and ver.get("verified") is True, {"tickets": [(t.get("status")) for t in recoveries],
                                                "problems": ver.get("problems")}),
        check("ticket_checks_all_true_from_independent_facts",
              bool(ver) and ver.get("result_confirmed") is True
              and all(ver.get(k) is True for k in REQUIRED_TICKET_CHECKS),
              {k: ver.get(k) for k in REQUIRED_TICKET_CHECKS}, needs="local_node_recovery._verify_assemble"),
        check("recovery_attempt_ran_with_zero_paid_seam_hits",
              ver.get("no_paid_seam_hit") is True and (ver.get("local_details") or {}).get("paid_seam_hits") == 0,
              (ver.get("local_details") or {}).get("paid_seam_hits")),
        check("incident_closed_by_the_verified_recovery", any(
            i.get("incident_id") == ticket.get("incident_id") and i.get("status") == "resolved_confirmed"
            for i in incidents), [(i.get("type"), i.get("status")) for i in incidents if i.get("incident_id") == ticket.get("incident_id")]),
        check("final_file_passes_the_real_verifier_after_the_run",
              final.get("status") == "passed" and {"video", "audio"} <= set(final.get("streams") or [])
              and duration_window[0] <= (final.get("duration") or 0) <= duration_window[1],
              {k: final.get(k) for k in ("status", "reason", "streams", "duration", "width", "height", "bytes")},
              needs="artifact_verifier.verify_artifact (ffprobe)"),
        check("approve_final_for_this_file_is_pending_and_unapproved", len(mine) == 1 and mine[0].get("decision") == "pending"
              and not any(a.get("decision") in ("approved", "approve") for a in approvals), approvals),
        check("no_part_file_is_left", not snap.get("part_files"), snap.get("part_files")),
        check("kill_switch_stops_the_recovery", (probes.get("kill_switch_recovery") or {}).get("refused") is True
              and (probes.get("kill_switch_recovery") or {}).get("tickets_after") == 0, probes.get("kill_switch_recovery")),
        check("no_provider_ledger_and_no_spend", not snap.get("ledger_rows") and not snap.get("spend_rows_paid"),
              {"ledger": snap.get("ledger_rows"), "spend": snap.get("spend_rows_paid")}),
        check("nothing_left_the_machine", counters.get("refused_foreign", 0) == 0 and not snap.get("refused_hosts")
              and counters.get("unexpected_model", 0) == 0,
              {"refused": snap.get("refused_hosts"), "unexpected_model": counters.get("unexpected_model")}),
        check("asr_probe_is_reported", isinstance((snap.get("asr_probe") or {}).get("faster_whisper_installed"), bool),
              snap.get("asr_probe")),
    ]
    if real:
        out.extend(evaluate_real_speech(snap))
    return out


def evaluate_real_speech(snap: dict) -> list[dict]:
    """Extra checks of --real-speech. Pure over snap["real_speech"] (collected by _collect_real_speech)."""
    rs = snap.get("real_speech") or {}
    asr, subs, pix = rs.get("asr") or {}, rs.get("subtitles") or {}, rs.get("pixels") or {}
    burn, voice = rs.get("burn") or {}, rs.get("voice") or {}
    final = snap.get("final_check") or {}
    voice_s = float(voice.get("duration") or 0.0)
    dur = final.get("duration")
    r15 = [x for x in (rs.get("r15_labels") or []) if x.strip()]
    return [
        check("real_asr_ran_offline_without_a_stub", bool(asr.get("calls")) and asr.get("stubbed") is False
              and asr.get("model") == rs.get("asr_model") and asr.get("device") == "cpu" and bool(asr.get("words")),
              {k: asr.get(k) for k in ("calls", "stubbed", "model", "device", "compute_type", "seconds_each", "word_count")},
              needs="app.services.montage_asr.ensure_montage_words -> faster-whisper"),
        check("asr_returned_enough_of_the_expected_words", (asr.get("check") or {}).get("matched", 0) >= REAL_MIN_WORDS,
              {k: (asr.get("check") or {}).get(k) for k in ("matched", "expected_count", "recognised", "matched_expected")}),
        check("asr_timestamps_monotonic_and_inside_the_voice",
              bool((asr.get("check") or {}).get("monotonic")) and bool((asr.get("check") or {}).get("inside_voice")),
              {k: (asr.get("check") or {}).get(k) for k in ("monotonic", "inside_voice", "first_start", "last_end", "voice_s")}),
        check("timeline_r15_came_from_the_asr_words", len(r15) >= len(REAL_FRAME_TEXTS), rs.get("r15_labels"),
              needs="plan_timestamps.ensure_r15_from_asr"),
        check("voice_is_the_full_voice_artifact_not_per_frame", rs.get("audio_mode") == "full_voice"
              and rs.get("subtitles_enabled") is True, {"audio_mode": rs.get("audio_mode"),
                                                        "subtitles_enabled": rs.get("subtitles_enabled")}),
        check("subtitle_file_was_built_from_the_asr_words", bool(subs.get("ok")),
              {k: subs.get(k) for k in ("cue_count", "cue_words", "script_words_shown", "aligned_to_asr", "within_tolerance",
                                        "max_abs_offset_s")}, needs="subtitles.build_subtitle_cues_from_cells"),
        check("production_filter_burned_the_generated_ass", burn.get("vf_calls", 0) >= 1 and burn.get("burn_commands", 0) >= 1
              and burn.get("ass_files", 0) >= 1 and burn.get("filter") == "subtitles=subs.ass"
              and burn.get("burn_uses_vf_filter") is True, burn),
        check("final_duration_matches_the_voice", bool(voice_s) and dur is not None
              and abs(float(dur) - voice_s) <= REAL_DURATION_TOLERANCE_S,
              {"final_duration": dur, "voice_duration": voice_s, "tolerance_s": REAL_DURATION_TOLERANCE_S}),
        check("subtitles_are_really_on_the_picture", bool(pix.get("ok")), pix),
        check("no_socket_left_the_machine", rs.get("network_attempts") == [] and rs.get("network_guard") is True,
              {"guard": rs.get("network_guard"), "attempts": rs.get("network_attempts")}),
    ]


def final_verdict(ok: bool, real_speech: bool = False) -> str:
    if real_speech:
        return ("DRY OK (local assemble with REAL offline ASR and burned subtitles on a Windows SAPI voice over synthetic "
                "clips; real ffmpeg + verification + one supervised re-run; no provider involved; human-voice ASR quality "
                "NOT proven)" if ok else "NOT PASSED")
    return ("DRY OK (local assemble: real ffmpeg + real verification + one supervised re-run; no provider involved; "
            "ASR/subtitles NOT exercised)" if ok else "NOT PASSED")


# ───────────────────────────── runtime ─────────────────────────────

def _args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dry", action="store_true", required=True,
                   help="the only mode: FREE and local (no provider, no key, no network)")
    p.add_argument("--work-dir", help="scratch dir for data/, runtime.json, result.json (recreated)")
    p.add_argument("--port", type=int, default=PORT)
    p.add_argument("--request-limit", type=int, default=60)
    p.add_argument("--timeout-s", type=int, default=None, help="default 420 s; 1500 s with --real-speech (the ASR runs on every native attempt)")
    p.add_argument("--real-speech", action="store_true",
                   help="F07: full-voice route with the real offline ASR + burned subtitles on --voice-wav (default off)")
    p.add_argument("--voice-wav", help="16-bit wav that says: " + REAL_VOICE_TEXT + " (required with --real-speech)")
    p.add_argument("--asr-model", default=DEFAULT_ASR_MODEL, help="faster-whisper model name, cached offline (default medium)")
    p.add_argument("--settle-s", type=float, default=4.0)
    p.add_argument("--sleep-minutes", type=float, default=0.01, help="native step-failure sleep (patched for speed)")
    args = p.parse_args(argv)
    if not args.work_dir:
        p.error("--work-dir is required")
    if args.real_speech:
        if not args.voice_wav:
            p.error("--real-speech needs --voice-wav <16-bit wav>")
        if not Path(args.voice_wav).is_file():
            p.error(f"--voice-wav not found: {args.voice_wav}")
    elif args.voice_wav:
        p.error("--voice-wav is used only with --real-speech")
    if args.timeout_s is None:
        args.timeout_s = 1500 if args.real_speech else 420
    return args


async def _seed(args: argparse.Namespace) -> tuple[int, str, Path]:
    from app.db import session_scope
    from app.models import (Artifact, ArtifactKind, Frame, NodeRun, NodeRunStatus, Project, ProjectStatus, Workflow,
                            WorkflowRun, WorkflowRunStatus)

    real = bool(getattr(args, "real_speech", False))
    node = {"id": TARGET, "type": NODE, "position": {"x": 100.0, "y": 100.0},
            "data": {"label": "Native assemble", "groupId": GROUP_ID}}
    nodes, edges = [node], []
    slug = f"media-assemble-{_uuid.uuid4().hex[:6]}"
    if real:
        return await _seed_real_speech(args, node, slug)
    text = " ".join(VOICEOVERS)
    async with session_scope() as session:
        wf = Workflow(name=f"wf-{slug}", is_default=False, nodes=nodes, edges=edges)
        session.add(wf)
        await session.flush()
        meta = {"canvas_graph": {"workflow_id": wf.id, "nodes": nodes, "edges": edges}, "img_streams": 1}
        project = Project(slug=slug, topic="media assemble production route e2e", status=ProjectStatus.audio_ready,
                          auto_mode=False, general_plan=text, script_text=text, meta=meta)
        session.add(project)
        await session.flush()
        run = WorkflowRun(project_id=project.id, workflow_id=wf.id, status=WorkflowRunStatus.new,
                          nodes_snapshot=nodes, edges_snapshot=edges)
        session.add(run)
        await session.flush()
        session.add(NodeRun(workflow_run_id=run.id, node_key=TARGET, node_type=NODE, status=NodeRunStatus.pending))
        frames = []
        for i, voice in enumerate(VOICEOVERS, start=1):
            fr = Frame(project_id=project.id, number=i, uuid=(f"{i:02d}" * 12), sort_key=float(i), voiceover_text=voice,
                       image_prompt="local frame", animation_prompt="local clip", duration_seconds=FRAME_SECONDS, attrs={})
            session.add(fr)
            frames.append(fr)
        await session.flush()
        pid, data_dir = int(project.id), project.data_dir
        videos, audio_dir = data_dir / "videos", data_dir / "audio"
        videos.mkdir(parents=True, exist_ok=True)
        audio_dir.mkdir(parents=True, exist_ok=True)
        parts = []
        for i, (fr, color) in enumerate(zip(frames, COLORS), start=1):
            clip = videos / f"clip_{i:03d}_{_uuid.uuid4().hex[:8]}.mp4"
            make_clip(clip, color)
            session.add(Artifact(project_id=project.id, frame_id=fr.id, kind=ArtifactKind.scene_video,
                                 uuid=_uuid.uuid4().hex, path=str(clip), meta={"shot": 1}))
            part = audio_dir / f"frame_{i:03d}.mp3"
            make_voice(part, 330 + 110 * i)
            parts.append(part)
        voice_full = audio_dir / "voice_full.mp3"
        concat_voice(parts, voice_full)
        session.add(Artifact(project_id=project.id, kind=ArtifactKind.audio, uuid=_uuid.uuid4().hex,
                             path=str(voice_full), meta={"mode": "per_frame"}))
        await session.flush()
    data_dir.mkdir(parents=True, exist_ok=True)
    (data_dir / "voiceover.txt").write_text(text, encoding="utf-8")
    return pid, slug, data_dir


async def _seed_real_speech(args: argparse.Namespace, node: dict, slug: str) -> tuple[int, str, Path]:
    """F07 project: the caller's wav as the FULL voice (mode full_voice), 3 frames with their texts in the plan sheet,
    3 vertical lavfi clips, subtitles enabled. No per-frame voice files and no words: the real ASR must produce them."""
    from app.db import session_scope
    from app.models import (Artifact, ArtifactKind, Frame, NodeRun, NodeRunStatus, Project, ProjectStatus, Workflow,
                            WorkflowRun, WorkflowRunStatus)
    from app.services.media_probe import probe_duration

    wav_info = validate_voice_wav(Path(args.voice_wav))
    nodes, edges = [node], []
    text = REAL_VOICE_TEXT
    async with session_scope() as session:
        wf = Workflow(name=f"wf-{slug}", is_default=False, nodes=nodes, edges=edges)
        session.add(wf)
        await session.flush()
        meta = {"canvas_graph": {"workflow_id": wf.id, "nodes": nodes, "edges": edges}, "img_streams": 1,
                "node_step_params": {"assemble": {"subtitles_enabled": True}}}
        project = Project(slug=slug, topic="media assemble real speech e2e", status=ProjectStatus.audio_ready,
                          auto_mode=False, general_plan=text, script_text=text, meta=meta)
        session.add(project)
        await session.flush()
        run = WorkflowRun(project_id=project.id, workflow_id=wf.id, status=WorkflowRunStatus.new,
                          nodes_snapshot=nodes, edges_snapshot=edges)
        session.add(run)
        await session.flush()
        session.add(NodeRun(workflow_run_id=run.id, node_key=TARGET, node_type=NODE, status=NodeRunStatus.pending))
        frames = []
        for i, voice in enumerate(REAL_FRAME_TEXTS, start=1):
            fr = Frame(project_id=project.id, number=i, uuid=(f"{i:02d}" * 12), sort_key=float(i), voiceover_text=voice,
                       image_prompt="local frame", animation_prompt="local clip", duration_seconds=3.0, attrs={})
            session.add(fr)
            frames.append(fr)
        await session.flush()
        pid, data_dir = int(project.id), project.data_dir
        videos, audio_dir = data_dir / "videos", data_dir / "audio"
        videos.mkdir(parents=True, exist_ok=True)
        audio_dir.mkdir(parents=True, exist_ok=True)
        for i, (fr, color) in enumerate(zip(frames, REAL_COLORS), start=1):
            clip = videos / f"clip_{i:03d}_{_uuid.uuid4().hex[:8]}.mp4"
            make_clip(clip, color, seconds=REAL_CLIP_SECONDS, size=REAL_CLIP_SIZE)
            session.add(Artifact(project_id=project.id, frame_id=fr.id, kind=ArtifactKind.scene_video,
                                 uuid=_uuid.uuid4().hex, path=str(clip), meta={"shot": 1}))
        voice_full = audio_dir / "voice_full.mp3"
        wav_to_mp3(Path(args.voice_wav), voice_full)
        session.add(Artifact(project_id=project.id, kind=ArtifactKind.audio, uuid=_uuid.uuid4().hex,
                             path=str(voice_full), meta={"mode": "full_voice"}))
        await session.flush()
    data_dir.mkdir(parents=True, exist_ok=True)
    (data_dir / "voiceover.txt").write_text(text, encoding="utf-8")
    make_plan_xlsx(data_dir / "project.xlsx")
    args.voice_info = {**wav_info, "wav": str(args.voice_wav), "mp3": str(voice_full),
                       "mp3_duration": round(float(await probe_duration(voice_full)), 3)}
    return pid, slug, data_dir


def _approvals(project_db: Path) -> list[dict]:
    rows = base._rows(project_db, "SELECT kind, decision, payload FROM hitl_requests ORDER BY id")
    out = []
    for r in rows:
        try:
            payload = json.loads(r.get("payload") or "{}")
        except ValueError:
            payload = {}
        if str(r.get("kind")).endswith("approve_final"):
            out.append({"kind": str(r.get("kind")), "decision": str(r.get("decision")).split(".")[-1].lower(),
                        "sha256": payload.get("sha256")})
    return out


def _final_state(project_root: Path, project_db: Path) -> dict:
    final_dir = project_root / "final"
    return {"final_files": sorted(p.name for p in final_dir.glob("*.mp4") if not p.name.startswith(".")) if final_dir.is_dir() else [],
            "part_files": sorted(p.name for p in final_dir.glob(".*.part.mp4")) if final_dir.is_dir() else [],
            "approvals": len(_approvals(project_db)), "node": base._node_status(project_db).get(TARGET)}


async def main(argv: list[str] | None = None) -> int:
    args = _args(argv)
    if "app.settings" in sys.modules:
        print("refused: the application settings were loaded before this stand set its scratch environment "
              "(they would point at the live data folder); run the stand in its own process; nothing was started")
        return 3
    work = Path(args.work_dir).resolve()
    data = guard.reset_scratch_data(work)  # A19: refuses live data
    os.chdir(ROOT)
    runtime_path = work / "runtime.json"
    runtime_path.write_text(json.dumps(base._runtime_config(args.port, args.request_limit, None), indent=2), encoding="utf-8")
    os.environ.update(base.env_overrides(data, args.port, fake_base="http://127.0.0.1:9", gate_cap=0.45, dry=True))
    os.environ["ORCHESTRATOR_RUNTIME_CONFIG"] = str(runtime_path)
    real_speech = bool(args.real_speech)
    model_link = None
    if real_speech:
        try:
            model_link = link_cached_model(data, args.asr_model)
        except (FileNotFoundError, OSError) as exc:
            print(f"refused: {exc}; nothing was started", flush=True)
            return 3
    if real_speech:     # offline ASR: the cached model only, CPU int8, no provider; set BEFORE the settings are loaded
        os.environ.update({"HF_HUB_OFFLINE": "1", "TRANSFORMERS_OFFLINE": "1", "ASR_BACKEND": "whisper",
                           "WHISPER_MODEL": args.asr_model, "WHISPER_DEVICE": "cpu", "WHISPER_COMPUTE_TYPE": "int8"})

    import httpx
    from app.settings import settings
    if Path(str(settings.data_dir)).resolve() != Path(data).resolve():
        print(f"refused: the application settings point at {settings.data_dir}, not at the scratch data {data}; "
              "nothing was started or written", flush=True)
        return 3
    if real_speech and (settings.whisper_model, settings.whisper_device, settings.whisper_compute_type,
                        settings.asr_backend) != (args.asr_model, "cpu", "int8", "whisper"):
        print("refused: the ASR settings are not the requested offline CPU whisper "
              f"({settings.asr_backend}/{settings.whisper_model}/{settings.whisper_device}/{settings.whisper_compute_type}); "
              "nothing was started or written", flush=True)
        return 3
    settings.gpt_relay_token = ""
    settings.gpt_proxy_url = None
    settings.vibecode_api_key = qc.SYNTHETIC_KEY
    settings.vibecode_base_url = f"http://127.0.0.1:{args.port}/v1"
    settings.gpt_base_url = f"http://127.0.0.1:{args.port}"
    settings.kie_api_key = ""
    settings.elevenlabs_api_key = ""
    settings.outsee_api_key = ""
    import app.services.api_tracker_hook as tracker
    tracker._TRACKER_DB = work / "tracker.db"
    if tracker._TRACKER_DB.exists():
        tracker._TRACKER_DB.unlink()
    from app.services import step_failure_policy
    from app.services.decision_transport import current as decision_hook
    step_failure_policy.SLEEP_MINUTES = args.sleep_minutes
    step_failure_policy.XLSX_SHEET_FORMAT_SLEEP_MINUTES = args.sleep_minutes
    from app import main as app_main
    await app_main._init_db()
    rec: dict[str, Any] = {"asr": [], "ass": [], "vf": [], "burn": [], "net": [], "control": None}
    restore_net = install_network_guard(rec["net"]) if real_speech else (lambda: None)
    try:
        pid, slug, project_root = await _seed(args)
    except ValueError as exc:       # a wav that is not 16-bit PCM etc.
        restore_net()
        print(f"refused: {exc}; nothing was started", flush=True)
        return 3
    project_db = data / "videos" / slug / "project.db"
    state_db = data / "state.db"
    api = f"http://127.0.0.1:{args.port}/api/projects/{pid}"

    base.NODE, base.STEP_CODE, base.TARGET, base.GROUP_ID = NODE, STEP_CODE, TARGET, GROUP_ID
    gate = base.Gate(mode="dry", max_paid_posts=1, fake_port=0)
    gate.phase = "main"
    counters: dict[str, Any] = {"injected_failures": 0, "injected_hitl": 0, "injected_tail": 0, "decision_stub": 0}
    decision_contexts: list[dict] = []
    inject = {"on": True}
    original_send = httpx.AsyncClient.send

    async def send(client, request, **kwargs):
        host, port, path = request.url.host, request.url.port, request.url.path
        judged = gate.judge(request.method, host, port, path, is_decision=decision_hook.get() is not None)
        if judged.action == "refuse":
            raise RuntimeError(judged.reason)
        if judged.action == "stub_decision":
            body = json.loads(request.content.decode("utf-8") or "{}")
            ctx = qc.decision_context(qc.request_text(body))
            if ctx is None:
                raise RuntimeError("decision stub: no context in the request")
            counters["decision_stub"] += 1
            decision_contexts.append({"turn": counters["decision_stub"], "incident_type": ctx.get("incident_type"),
                                      "allowed_actions": ctx.get("allowed_actions"),
                                      "media_recovery": ctx.get("media_recovery")})
            answer = json.dumps(qc.decision_proposal(counters["decision_stub"], ctx, "retry_same"), ensure_ascii=False)
            return httpx.Response(200, headers={"content-type": "application/json"},
                                  content=qc.chat_json_body(answer, f"chatcmpl-{_uuid.uuid4().hex[:12]}"), request=request)
        return await original_send(client, request, **kwargs)

    from app.orchestrator.steps import assemble as asm
    real_assemble = asm.assemble

    async def forced_failure(*a, **k):
        if inject["on"]:
            counters["injected_failures"] += 1
            raise RuntimeError(FAULT)
        return await real_assemble(*a, **k)

    real_variant2 = asm.run_variant2

    async def forced_failure_v2(project, frame_numbers, voice, out, *a, **k):
        """The full-voice path assembles with run_variant2 (montage-v3), not with assemble(): same forced failure, and a
        copy of its output (the picture BEFORE the subtitle burn) is kept as the no-subtitles control render."""
        if inject["on"]:
            counters["injected_failures"] += 1
            raise RuntimeError(FAULT)
        got = await real_variant2(project, frame_numbers, voice, out, *a, **k)
        control = work / "control_no_subtitles.mp4"
        shutil.copy2(out, control)
        rec["control"] = str(control)
        return got

    asm.assemble = forced_failure
    undo_real = _install_real_speech_recorders(asm, rec, forced_failure_v2) if real_speech else (lambda: None)
    httpx.AsyncClient.send = send

    result: dict[str, Any] = {"mode": "dry", "node": NODE, "target": TARGET, "project_id": pid, "slug": slug,
                              "args": {k: v for k, v in vars(args).items()}, "asr_probe": asr_probe(real_speech)}
    probes: dict[str, Any] = {}
    task = asyncio.create_task(app_main.main())
    execution = None
    try:
        await qc._wait_api(api)
        from app.services import access_guard
        operator = {"x-vp-operator": access_guard.get_operator_token()}
        print(f"preflight: node={NODE} mode=dry local-only asr={result['asr_probe']}", flush=True)
        async with httpx.AsyncClient(timeout=30, trust_env=False, headers=operator,
                                     limits=httpx.Limits(max_keepalive_connections=0)) as client:
            prefix = api + "/observations"
            r = await client.post(prefix, json={"node_keys": [TARGET]})
            print("observation:", r.status_code, flush=True)
            r.raise_for_status()
            execution = r.json()["execution_id"]
            result["execution_id"] = execution
            r = await client.post(api + f"/steps/{STEP_CODE}/run", params={"node_key": TARGET})
            print("start assemble:", r.status_code, (r.text[:300] if r.status_code >= 300 else "ok"), flush=True)
            r.raise_for_status()
            await base._wait_exhausted(args, project_db, counters, gate, None)
            result["state_at_exhaustion"] = _final_state(project_root, project_db)
            inject["on"] = False            # the cause is gone: the supervised re-run may now succeed
            incident = await base._wait_incident(client, prefix + f"/{execution}", args.timeout_s)
            result["incident_id"] = incident["incident_id"] if incident else None
            if not incident:
                raise RuntimeError("no current media incident of the assemble node after the native retries were exhausted")
            r = await client.post(prefix + f"/{execution}/actions/enable", json={
                "group_id": GROUP_ID, "node_keys": [TARGET], "rationale": "assemble production route e2e: supervised re-run",
                "request_limit": min(args.request_limit, 60)})
            print("actions/enable:", r.status_code, (r.text[:300] if r.status_code >= 300 else "ok"), flush=True)
            r.raise_for_status()
            await base._probe_kill_switch_recovery(client, api, prefix + f"/{execution}", incident["incident_id"], probes)
            r = await client.post(prefix + f"/{execution}/decisions/control", json={
                "enabled": True, "rationale": "dry decision loop for the local assemble re-run"})
            print("decisions/control:", r.status_code, flush=True)
            r.raise_for_status()
            await base._poll_recovery(client, prefix + f"/{execution}", args, project_db, counters, gate, result)
        result["final_check"] = await _verify_final(project_root)
        if real_speech:
            result["real_speech"] = await asyncio.to_thread(_collect_real_speech, args, rec, project_root, work, project_db)
    except RuntimeError as exc:
        result["error"] = str(exc)
        print("STOP:", exc)
    finally:
        task.cancel()
        try:
            await task
        except BaseException:
            pass
        httpx.AsyncClient.send = original_send
        asm.assemble = real_assemble
        undo_real()
        restore_net()
        if model_link is not None:
            unlink_cached_model(model_link)
    await asyncio.sleep(1)
    if real_speech:
        result.setdefault("real_speech", {})["network_attempts"] = list(rec["net"])
        result["real_speech"]["network_guard"] = True
        result["real_speech_mode"] = True
    return _finish(args, result, gate, counters, probes, decision_contexts, state_db, project_db, project_root, execution, work)


def model_cache_dir_name(model_name: str) -> str:
    return "models--Systran--faster-whisper-" + str(model_name).strip()


def find_cached_model(model_name: str, env: dict | None = None) -> Path | None:
    """The offline copy of a faster-whisper model in the user's Hugging Face cache (nothing is downloaded or loaded)."""
    env = os.environ if env is None else env
    hubs = [Path(env[k]) for k in ("HF_HUB_CACHE", "HUGGINGFACE_HUB_CACHE") if env.get(k)]
    if env.get("HF_HOME"):
        hubs.append(Path(env["HF_HOME"]) / "hub")
    hubs.append(Path.home() / ".cache" / "huggingface" / "hub")
    for hub in hubs:
        cand = hub / model_cache_dir_name(model_name)
        if (cand / "snapshots").is_dir() and any((cand / "snapshots").glob("*/model.bin")):
            return cand
    return None


def link_cached_model(data: Path, model_name: str, env: dict | None = None) -> Path:
    """The app redirects the Hugging Face cache into <data>/.cache/huggingface/hub (nvidia_asr_env), so the scratch data
    would look empty offline. Make the user's cached model visible there (a directory junction / symlink: no copy, no
    download). Returns the link path; raises FileNotFoundError when the model is not cached."""
    source = find_cached_model(model_name, env)
    if source is None:
        raise FileNotFoundError(f"faster-whisper model {model_name!r} is not in the local Hugging Face cache (offline run)")
    hub = Path(data) / ".cache" / "huggingface" / "hub"
    hub.mkdir(parents=True, exist_ok=True)
    link = hub / source.name
    if link.exists() or link.is_symlink():
        unlink_cached_model(link)
    if os.name == "nt":
        done = subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(source)], capture_output=True, text=True)
        if done.returncode != 0:
            raise OSError(f"cannot link the cached model: {done.stdout} {done.stderr}".strip())
    else:
        os.symlink(source, link, target_is_directory=True)
    return link


def unlink_cached_model(link: Path) -> None:
    """Remove ONLY the link (never the cached model behind it)."""
    link = Path(link)
    try:
        is_link = link.is_symlink() or (hasattr(os.path, "isjunction") and os.path.isjunction(link))
        if is_link:
            if os.name == "nt":
                os.rmdir(link)
            else:
                link.unlink()
        elif link.exists():
            raise OSError(f"{link} is not a link: refusing to remove it")
    except FileNotFoundError:
        pass


def _install_real_speech_recorders(asm, rec: dict, forced_variant2):
    """Observe (never replace) the production seams of the full-voice path. The ASR is the REAL one: the wrapper calls
    through and keeps the result. Returns the function that puts everything back."""
    from dataclasses import asdict
    from app.services import montage_asr
    from app.settings import settings

    real_transcribe, real_ass, real_vf = montage_asr.transcribe_words, asm.make_simple_ass, asm.subtitles_vf_arg
    real_exec, real_v2 = asyncio.create_subprocess_exec, asm.run_variant2
    stubbed = getattr(real_transcribe, "__module__", "") != "app.services.asr" or real_transcribe.__name__ != "transcribe_words"

    def transcribe(path, *a, **k):
        t0 = time.monotonic()
        words = real_transcribe(path, *a, **k)
        rec["asr"].append({"audio": str(path), "seconds": round(time.monotonic() - t0, 2), "stubbed": stubbed,
                           "model": settings.whisper_model, "device": settings.whisper_device,
                           "compute_type": settings.whisper_compute_type, "kwargs": {kk: str(v) for kk, v in k.items()},
                           "words": [asdict(w) for w in words]})
        return words

    def make_ass(entries, path, **kw):
        out = real_ass(entries, path, **kw)
        rec["ass"].append({"path": str(path), "entries": [list(e) for e in entries], "kwargs": kw,
                           "text": Path(path).read_text(encoding="utf-8")})
        return out

    def vf_arg(*a, **k):
        value = real_vf(*a, **k)
        rec["vf"].append(value)
        return value

    async def exec_(*cmd, **kw):
        if any(str(c).startswith("subtitles=") for c in cmd):
            rec["burn"].append([str(c) for c in cmd])
        return await real_exec(*cmd, **kw)

    montage_asr.transcribe_words, asm.make_simple_ass, asm.subtitles_vf_arg = transcribe, make_ass, vf_arg
    asyncio.create_subprocess_exec, asm.run_variant2 = exec_, forced_variant2

    def undo() -> None:
        montage_asr.transcribe_words, asm.make_simple_ass, asm.subtitles_vf_arg = real_transcribe, real_ass, real_vf
        asyncio.create_subprocess_exec, asm.run_variant2 = real_exec, real_v2

    return undo


def _collect_real_speech(args, rec: dict, project_root: Path, work: Path, project_db: Path | None = None) -> dict:
    """Evidence of --real-speech from what the production path itself produced (words, ASS file, filter, final)."""
    voice = dict(getattr(args, "voice_info", {}) or {})
    voice["duration"] = voice.get("mp3_duration") or voice.get("duration")
    voice_s = float(voice.get("duration") or 0.0)
    calls = rec["asr"]
    words = calls[-1]["words"] if calls else []
    asr = {"calls": len(calls), "stubbed": bool(calls) and any(c["stubbed"] for c in calls) or not calls,
           "model": calls[-1]["model"] if calls else None, "device": calls[-1]["device"] if calls else None,
           "compute_type": calls[-1]["compute_type"] if calls else None,
           "seconds_each": [c["seconds"] for c in calls], "word_count": len(words), "words": words,
           "check": check_asr_words(words, voice_s)}
    ass = rec["ass"][-1] if rec["ass"] else None
    cues = parse_ass_cues(ass["text"]) if ass else []
    subs = check_subtitles_from_words(cues, words)
    subs["cues"] = [[round(s, 2), round(e, 2), t] for s, e, t in cues]
    burn_cmd = rec["burn"][-1] if rec["burn"] else []
    burn = {"vf_calls": len(rec["vf"]), "filter": rec["vf"][-1] if rec["vf"] else None, "burn_commands": len(rec["burn"]),
            "ass_files": len(rec["ass"]), "ass_path": ass["path"] if ass else None,
            "burn_uses_vf_filter": bool(burn_cmd) and "-vf" in burn_cmd}
    finals = sorted(p for p in (project_root / "final").glob("*.mp4") if not p.name.startswith(".")) \
        if (project_root / "final").is_dir() else []
    pixels: dict[str, Any] = {"ok": False, "reason": "no final or no control render"}
    if finals and rec.get("control") and Path(rec["control"]).is_file() and words:
        probe_t = pick_probe_times(cues, words, voice_s)
        word_frame = frame_pair_difference(finals[-1], Path(rec["control"]), probe_t["word"]["t"]) if probe_t["word"] else None
        silent_frame = frame_pair_difference(finals[-1], Path(rec["control"]), probe_t["silent"]) if probe_t["silent"] else None
        pixels = {**check_subtitle_pixels(word_frame, silent_frame), "probe": probe_t,
                  "word_frame": word_frame, "silent_frame": silent_frame, "control": rec["control"]}
    try:
        r15 = read_r15_labels(project_root / "project.xlsx", len(REAL_FRAME_TEXTS))
    except Exception as exc:  # noqa: BLE001
        r15 = [f"unreadable: {type(exc).__name__}"]
    audio_mode = None
    if project_db is not None:
        for row in base._rows(project_db, "SELECT meta FROM artifacts WHERE kind = 'audio' ORDER BY id DESC LIMIT 1"):
            try:
                audio_mode = json.loads(row.get("meta") or "{}").get("mode")
            except ValueError:
                audio_mode = None
    return {"voice": voice, "asr_model": args.asr_model, "asr": asr, "subtitles": subs, "burn": burn, "pixels": pixels,
            "r15_labels": r15, "audio_mode": audio_mode, "subtitles_enabled": bool(ass), "control_render": rec.get("control")}


async def _verify_final(project_root: Path) -> dict:
    from app.services.artifact_verifier import verify_artifact
    final_dir = project_root / "final"
    files = sorted(p for p in final_dir.glob("*.mp4") if not p.name.startswith(".")) if final_dir.is_dir() else []
    if not files:
        return {"status": "failed", "reason": "no_final_file"}
    got = await verify_artifact(files[-1], root=project_root, kind="assembly")
    return {k: v for k, v in got.items() if k in ("status", "reason", "bytes", "sha256", "format", "width", "height",
                                                 "duration", "streams")} | {"file": files[-1].name}


def _finish(args, result, gate, counters, probes, decision_contexts, state_db, project_db, project_root, execution, work) -> int:
    records = base.read_records(state_db)
    events = qc._events(state_db, execution)
    attempts = sum(1 for e in events if e["kind"] == "node_start" and e.get("node_key") == TARGET)
    snap: dict[str, Any] = {
        **result,
        "counters": {**counters, "paid_posts": gate.paid_posts, "unexpected_model": gate.unexpected_model,
                     "refused_foreign": len([r for r in gate.refused if r["kind"] in ("dry_foreign", "foreign")])},
        "probes": probes, "ledger_rows": records["ledger_rows"],
        "spend_rows_paid": [r for r in records["spend_rows"] if str(r.get("provider")).lower() in MEDIA_PROVIDERS],  # not the (stub) Decision request rows
        "node_status": base._node_status(project_db), "node_attempts": attempts, "step_failure": qc._step_failure(project_db),
        "decision_contexts": decision_contexts, "approvals": _approvals(project_db),
        "part_files": _final_state(project_root, project_db)["part_files"],
        "refused_hosts": sorted({r["host"] for r in gate.refused if r["kind"] in ("dry_foreign", "foreign")}),
    }
    checks = evaluate(snap)
    ok, failed = verdict(checks)
    if result.get("error"):
        ok = False
        failed = ["flow_error"] + failed
        checks.insert(0, check("flow_completed_without_error", False, str(result["error"])))
    real_mode = bool(getattr(args, "real_speech", False))
    reality = reality_report(real_mode)
    summary = {"mode": "dry", "node": NODE, "verdict": final_verdict(ok, real_mode), "failed_checks": failed,
               "node_attempts": attempts, "asr_probe": result.get("asr_probe"), "not_proven": reality_not_proven(real_mode)}
    slim = dict(snap)
    slim.update(checks=checks, ok=ok, failed=failed, reality=reality, summary=summary)
    (work / "result.json").write_text(json.dumps(slim, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print("\n=== RESULT ===")
    print(f"node: {NODE} | mode: dry (local, free) | node status: {snap['node_status']} | native attempts: {attempts}")
    print("counters:", json.dumps(snap["counters"], default=str))
    for c in checks:
        print(("PASS " if c["ok"] else "FAIL ") + c["name"] + ("" if c["ok"] or not c["needs"] else f"   [needs: {c['needs']}]"))
        if not c["ok"]:
            print("     detail:", json.dumps(c["detail"], ensure_ascii=False, default=str)[:400])
    print("real:", reality["real"])
    print("not real:", reality["not_real"])
    print("not proven:", summary["not_proven"])
    print("asr probe:", result.get("asr_probe"))
    rs = snap.get("real_speech") or {}
    if real_mode and rs:
        chk, pix = (rs.get("asr") or {}).get("check") or {}, rs.get("pixels") or {}
        print(f"real speech: asr calls={(rs.get('asr') or {}).get('calls')} model={(rs.get('asr') or {}).get('model')} "
              f"matched={chk.get('matched')}/{chk.get('expected_count')} recognised={chk.get('recognised')}")
        print(f"real speech: subtitles cues={(rs.get('subtitles') or {}).get('cue_count')} "
              f"offsets={(rs.get('subtitles') or {}).get('offsets_s')} | pixels lower={pix.get('lower_fraction_at_word')} "
              f"upper={pix.get('upper_fraction_at_word')} silence={pix.get('lower_fraction_in_silence')} | "
              f"voice={rs.get('voice', {}).get('duration')}s final={(snap.get('final_check') or {}).get('duration')}s")
    print("verdict:", summary["verdict"])
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
